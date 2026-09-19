#!/usr/bin/env python3
"""
The KPI Profile Workbook, and the per-project tracker workbook.

The profile workbook is the human face of profile.yaml: one tab per section, one row per
setting, and every row carrying a plain-language description of what it is for. It is
generated from schemas/profile.schema.json and read back into the same YAML, so the thing a
person edits and the thing the tools read can never drift apart.

Why bother, when the YAML already exists: because the person who needs to know which chat
space matters, or where the plan lives, is usually not the person who wrote the YAML. The
Tools tab is the one people actually open - it is a map of the project's truth.

Commands:
    workbook.py build   --profile profile.yaml --out "KPI Profile Workbook.xlsx"
    workbook.py read    --xlsx "KPI Profile Workbook.xlsx" --out profile.yaml
    workbook.py tracker --results results.json --kif run.kif.json --out tracker.xlsx
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import sys
from pathlib import Path
from typing import Any

HERE = Path(__file__).resolve().parent
PLUGIN_ROOT = HERE.parent
SCHEMA = PLUGIN_ROOT / "schemas" / "profile.schema.json"

sys.path.insert(0, str(HERE))
from profile_lib import describe_overrides, unflatten  # noqa: E402

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    print("openpyxl is needed for the workbook. Run: pip3 install openpyxl", file=sys.stderr)
    raise SystemExit(2)

# One look everywhere, so every team's workbook is recognisable as the same thing.
HEAD = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
INPUT = PatternFill("solid", fgColor="FFF2CC")   # yellow = you edit this
LOCKED = PatternFill("solid", fgColor="EDEDED")  # grey  = do not edit
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
HINT_FONT = Font(italic=True, size=9, color="666666")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)


def _load_any(path: Path) -> Any:
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() in (".yaml", ".yml"):
        import yaml  # type: ignore
        return yaml.safe_load(text)
    return json.loads(text)


def _dump_yaml(data: Any, path: Path) -> None:
    import yaml  # type: ignore
    path.write_text(
        yaml.safe_dump(data, sort_keys=False, allow_unicode=True, width=100), encoding="utf-8"
    )


def _flatten(value: Any) -> str:
    """Lists become one-per-line so a person can edit them without learning YAML."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, list):
        return "\n".join(_flatten(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _unflatten(text: str, spec: dict) -> Any:
    """The inverse, guided by the schema so 'Yes' becomes True only where a boolean belongs."""
    text = (text or "").strip()
    # A schema type can be a list ("integer" or null), so normalise before matching -
    # otherwise a PMS project id comes back as the string "361" and every later comparison
    # against PMS quietly fails.
    t = spec.get("type")
    types = set(t) if isinstance(t, list) else {t}
    if "boolean" in types:
        # An empty cell means "not set", not False. Otherwise every optional flag comes back
        # explicitly false and the round trip reports differences that are not differences.
        if not text:
            return None
        return text.lower() in ("yes", "true", "1", "on")
    if "array" in types:
        return [line.strip() for line in text.splitlines() if line.strip()] if text else []
    if "integer" in types and text:
        try:
            return int(float(text))
        except ValueError:
            return None
    if "number" in types and text:
        try:
            return float(text)
        except ValueError:
            return None
    if "object" in types and text:
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            return {}
    return text or None


def _enum_of(spec: dict) -> list[str] | None:
    if "enum" in spec:
        return [str(v) for v in spec["enum"] if v is not None]
    if spec.get("type") == "boolean":
        return ["Yes", "No"]
    return None


# --------------------------------------------------------------------------------------
# build: schema + profile -> xlsx
# --------------------------------------------------------------------------------------


def _style_header(ws, row: int, headers: list[str], widths: list[int]) -> None:
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font, c.border, c.alignment = HEAD, HEAD_FONT, BORDER, WRAP
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def build(profile_path: Path | None, out: Path) -> None:
    schema = _load_any(SCHEMA)
    profile = _load_any(profile_path) if profile_path and profile_path.exists() else {}
    wb = Workbook()
    wb.remove(wb.active)

    # --- Read Me -------------------------------------------------------------------
    ws = wb.create_sheet("Read Me")
    ws.column_dimensions["A"].width = 110
    lines = [
        ("KPI Profile Workbook", "title"),
        ("", ""),
        ("This workbook is the configuration for KPI Copilot. It is the same information as profile.yaml, "
         "shown so a person can read it. Edit either one; the tools convert between them.", ""),
        ("", ""),
        ("Yellow cells are yours to change. Grey cells are not edited here: they are read from PMS or fixed "
         "for the company. Targets are the common case - a threshold is configurable per project in PMS, so if "
         "yours needs a different bar, set it there and the next run reads it. What does not vary is the "
         "counting behind each number, because that is what lets two projects be compared at all.", ""),
        ("", ""),
        ("Tabs", "h"),
        ("Owner - who this profile belongs to.", ""),
        ("Organization - PMS, the KPI set and the note format. Targets live in PMS, per project.", ""),
        ("Tools - where everything lives. The tab people actually use: tracker, chat spaces, plan, sheets.", ""),
        ("Tracker - which adapter reads your board, and the fields it should look at.", ""),
        ("Conventions - how your team names tickets, defects and the things that are not deliverables.", ""),
        ("Workflow - what your states mean. The most important tab: it turns 'a column called In Test' into "
         "a number management can compare across projects.", ""),
        ("Sources - where the plan, the estimates and the timeline live. All optional.", ""),
        ("Periods - how you slice a project for PMS.", ""),
        ("Counting Rules - the few counting choices a team is allowed to make.", ""),
        ("Custom Instructions - your house style, always/never rules, glossary and rule overrides.", ""),
        ("Output - what the run produces, and how far it may go on its own.", ""),
        ("Accounts - your client accounts. Most of what varies between two projects actually varies "
         "between two clients, so say it once here.", ""),
        ("Projects - one row per PMS project. A project inherits its account and only says what differs.", ""),
        ("Overrides - every way an account or project departs from the defaults, on one page.", ""),
        ("Prerequisites - the readiness checklist, refreshed by scripts/preflight.py.", ""),
        ("", ""),
        ("After editing", "h"),
        ("python3 scripts/workbook.py read --xlsx \"KPI Profile Workbook.xlsx\" --out profile.yaml", "code"),
        ("python3 scripts/preflight.py --profile profile.yaml", "code"),
    ]
    r = 1
    for text, kind in lines:
        c = ws.cell(row=r, column=1, value=text)
        c.alignment = WRAP
        if kind == "title":
            c.font = TITLE_FONT
        elif kind == "h":
            c.font = Font(bold=True, size=11, color="1F3864")
        elif kind == "code":
            c.font = Font(name="Menlo", size=9)
            c.fill = LOCKED
        r += 1

    # --- one tab per schema section -------------------------------------------------
    for key, spec in schema.get("properties", {}).items():
        xw = spec.get("x-workbook") or {}
        tab = xw.get("tab")
        if not tab:
            continue
        ws = wb.create_sheet(tab[:31])
        ws.cell(row=1, column=1, value=xw.get("title", tab)).font = TITLE_FONT
        if xw.get("readonly_hint"):
            c = ws.cell(row=2, column=1, value=xw["readonly_hint"])
            c.font = HINT_FONT
            c.alignment = WRAP

        if xw.get("as_table"):
            _build_table_tab(ws, key, spec, xw, profile.get(key) or [])
        else:
            _build_settings_tab(ws, key, spec, profile.get(key) or {})

    # --- Overrides -------------------------------------------------------------------
    # Flattened so a person can see, on one page, every way a project or account departs
    # from the defaults. Two projects that drifted apart for a reason nobody remembers show
    # up here and nowhere else.
    over_rows = describe_overrides(profile)
    ws = wb.create_sheet("Overrides")
    ws.cell(row=1, column=1, value="What each account and project changes").font = TITLE_FONT
    c = ws.cell(row=2, column=1,
                value="The top level of this workbook is your default way of working. An account says what is "
                      "true for every project on that client; a project says what is true for just itself. "
                      "Most rows here belong on an account - if the same override appears on several projects, "
                      "move it up.")
    c.font, c.alignment = HINT_FONT, WRAP
    ws.row_dimensions[2].height = 42
    _style_header(ws, 4, ["#", "Scope", "Account / project", "Setting", "Value"],
                  [5, 12, 24, 40, 70])
    for i, row in enumerate(over_rows, start=1):
        for ci, v in enumerate([i, row["scope"], row["id"], row["setting"], row["value"]], start=1):
            cell = ws.cell(row=4 + i, column=ci, value=v)
            cell.border, cell.alignment = BORDER, WRAP
            if ci == 5:
                cell.fill = INPUT
    if not over_rows:
        ws.cell(row=5, column=2,
                value="None. Every project uses the defaults above.").font = HINT_FONT
    ws.cell(row=3, column=9, value="_overrides").font = Font(size=8, color="FFFFFF")

    # --- Prerequisites --------------------------------------------------------------
    ws = wb.create_sheet("Prerequisites")
    ws.cell(row=1, column=1, value="Readiness checklist").font = TITLE_FONT
    c = ws.cell(row=2, column=1,
                value="Refreshed by: python3 scripts/preflight.py --profile profile.yaml. "
                      "Do not edit by hand - confirmations are recorded with the command, so they carry a date.")
    c.font, c.alignment = HINT_FONT, WRAP
    _style_header(ws, 4, ["#", "Group", "Check", "State", "Why it matters", "How to fix", "Who owns it", "Checked"],
                  [5, 16, 42, 14, 50, 50, 14, 20])
    state = _read_preflight(profile_path)
    for i, chk in enumerate(state, start=1):
        row = 4 + i
        vals = [i, chk.get("group"), chk.get("label"),
                {"pass": "Ready", "fail": "Blocked", "manual": "Confirm", "warn": "Limited", "skip": "N/A"}
                .get(chk.get("state"), chk.get("state")),
                chk.get("why"), chk.get("fix"), chk.get("owner"), (chk.get("checked_at") or "")[:10]]
        for ci, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=ci, value=v)
            cell.border, cell.alignment = BORDER, WRAP
            if ci == 4:
                cell.fill = {"Ready": PatternFill("solid", fgColor="D9EAD3"),
                             "Blocked": PatternFill("solid", fgColor="F4CCCC"),
                             "Limited": PatternFill("solid", fgColor="FFF2CC"),
                             "Confirm": PatternFill("solid", fgColor="FFF2CC")}.get(v, LOCKED)
    if not state:
        ws.cell(row=5, column=3, value="Not run yet. Run preflight.py and rebuild this workbook.").font = HINT_FONT

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out}")


def _read_preflight(profile_path: Path | None) -> list[dict]:
    if not profile_path:
        return []
    p = profile_path.parent / "preflight.json"
    if not p.exists():
        return []
    try:
        return json.loads(p.read_text(encoding="utf-8")).get("checks", [])
    except Exception:  # noqa: BLE001
        return []


def _build_settings_tab(ws, section: str, spec: dict, values: dict) -> None:
    _style_header(ws, 4, ["Setting", "Value", "What it means", "Allowed"], [34, 42, 62, 26])
    row = 5
    deferred_tables: list[tuple[str, dict]] = []
    for prop, pspec in (spec.get("properties") or {}).items():
        pxw = pspec.get("x-workbook") or {}
        label = pxw.get("label", prop.replace("_", " ").capitalize())
        meaning = pxw.get("hint") or pspec.get("description", "")
        if pspec.get("type") == "array" and (pxw.get("as_table") or (pspec.get("items") or {}).get("properties")):
            # An array of objects inside a settings tab (rule overrides) gets its own small
            # table under the settings, rather than one unreadable JSON cell.
            deferred_tables.append((prop, pspec))
            continue
        if pspec.get("type") == "object" and pspec.get("properties"):
            # A nested block (delivered_when, plan, ...) gets one row per leaf so nobody has
            # to hand-write JSON in a spreadsheet cell.
            ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=11)
            ws.cell(row=row, column=3, value=meaning).alignment = WRAP
            ws.cell(row=row, column=3).font = HINT_FONT
            row += 1
            for sub, sspec in pspec["properties"].items():
                sval = (values.get(prop) or {}).get(sub)
                row = _write_setting(ws, row, f"{prop}.{sub}", f"    {sub.replace('_', ' ')}",
                                     sval, sspec, sspec.get("description", ""))
            continue
        row = _write_setting(ws, row, prop, label, values.get(prop), pspec, meaning)

    for prop, pspec in deferred_tables:
        pxw = pspec.get("x-workbook") or {}
        cols = pxw.get("columns") or list(((pspec.get("items") or {}).get("properties") or {}).keys())
        item_props = ((pspec.get("items") or {}).get("properties")) or {}
        row += 2
        t = ws.cell(row=row, column=1, value=prop.replace("_", " ").title())
        t.font = Font(bold=True, size=12, color="1F3864")
        row += 1
        d = ws.cell(row=row, column=1, value=pspec.get("description", ""))
        d.font, d.alignment = HINT_FONT, WRAP
        ws.row_dimensions[row].height = 46
        row += 1
        for ci, c in enumerate(cols, start=1):
            h = ws.cell(row=row, column=ci, value=c.replace("_", " ").title())
            h.fill, h.font, h.border, h.alignment = HEAD, HEAD_FONT, BORDER, WRAP
        marker = ws.cell(row=row, column=9, value=f"_subtable:{prop}:{','.join(cols)}")
        marker.font = Font(size=8, color="CCCCCC")
        row += 1
        for ci, c in enumerate(cols, start=1):
            m = ws.cell(row=row, column=ci, value=(item_props.get(c) or {}).get("description", ""))
            m.font, m.alignment, m.border = HINT_FONT, WRAP, BORDER
        ws.row_dimensions[row].height = 40
        row += 1
        existing = values.get(prop) or []
        for item in list(existing) + [{}] * 6:
            for ci, c in enumerate(cols, start=1):
                cell = ws.cell(row=row, column=ci, value=_flatten(item.get(c)))
                cell.fill, cell.border, cell.alignment = INPUT, BORDER, WRAP
            ws.row_dimensions[row].height = 28
            row += 1


def _write_setting(ws, row: int, key: str, label: str, value: Any, spec: dict, meaning: str) -> int:
    ws.cell(row=row, column=1, value=label).alignment = WRAP
    ws.cell(row=row, column=1).border = BORDER
    v = ws.cell(row=row, column=2, value=_flatten(value if value is not None else spec.get("default")))
    v.fill, v.border, v.alignment = INPUT, BORDER, WRAP
    m = ws.cell(row=row, column=3, value=meaning)
    m.border, m.alignment, m.font = BORDER, WRAP, Font(size=9, color="444444")
    enum = _enum_of(spec)
    a = ws.cell(row=row, column=4, value=" | ".join(enum) if enum else "")
    a.border, a.alignment, a.font = BORDER, WRAP, HINT_FONT
    if enum and len(" ".join(enum)) < 240:
        dv = DataValidation(type="list", formula1='"' + ",".join(enum) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(v)
    # The machine key, kept so 'read' can find its way back without guessing from the label.
    k = ws.cell(row=row, column=9, value=key)
    k.font = Font(size=8, color="CCCCCC")
    ws.column_dimensions["I"].hidden = True
    ws.row_dimensions[row].height = 30
    return row + 1


def _build_table_tab(ws, section: str, spec: dict, xw: dict, rows: list[dict]) -> None:
    cols = xw.get("columns") or []
    item_props = ((spec.get("items") or {}).get("properties")) or {}
    headers = [c.replace("_", " ").title() for c in cols]
    _style_header(ws, 4, headers, [16, 16, 24, 34, 54, 26, 18, 18, 18, 30][: len(cols)])
    # Column meanings go on row 3 so the table itself stays clean.
    for i, c in enumerate(cols, start=1):
        d = (item_props.get(c) or {}).get("description", "")
        cell = ws.cell(row=3, column=i, value=d)
        cell.font, cell.alignment = HINT_FONT, WRAP
    ws.row_dimensions[3].height = 42
    for ri, item in enumerate(rows, start=1):
        for ci, c in enumerate(cols, start=1):
            cell = ws.cell(row=4 + ri, column=ci, value=_flatten(item.get(c)))
            cell.fill, cell.border, cell.alignment = INPUT, BORDER, WRAP
        ws.row_dimensions[4 + ri].height = 30
    for extra in range(len(rows) + 1, len(rows) + 9):  # blank rows, ready to fill
        for ci in range(1, len(cols) + 1):
            cell = ws.cell(row=4 + extra, column=ci)
            cell.fill, cell.border = INPUT, BORDER
    ws.cell(row=2, column=1, value="_table:" + section).font = Font(size=8, color="FFFFFF")


# --------------------------------------------------------------------------------------
# read: xlsx -> profile
# --------------------------------------------------------------------------------------


def read(xlsx: Path, out: Path) -> None:
    schema = _load_any(SCHEMA)
    wb = load_workbook(xlsx, data_only=True)
    profile: dict = {"profile_version": "1.0"}

    for key, spec in schema.get("properties", {}).items():
        xw = spec.get("x-workbook") or {}
        tab = xw.get("tab")
        if not tab or tab[:31] not in wb.sheetnames:
            continue
        ws = wb[tab[:31]]
        if xw.get("as_table"):
            rows = _read_table(ws, spec, xw)
            # An empty tab means the section is absent, not present-and-empty. A profile with
            # no accounts should read back with no accounts key, or the round trip reports a
            # difference that is not one.
            if rows:
                profile[key] = rows
        else:
            profile[key] = _read_settings(ws, spec)

    # Fold the Overrides tab back onto the accounts and projects it belongs to.
    if "Overrides" in wb.sheetnames:
        ws = wb["Overrides"]
        rows = []
        for r in ws.iter_rows(min_row=5, max_col=5):
            scope, rid, setting, value = (r[1].value, r[2].value, r[3].value, r[4].value)
            if rid and setting:
                rows.append({"scope": scope or "project", "id": str(rid),
                             "setting": str(setting), "value": value})
        folded = unflatten(rows, schema)
        for (scope, rid), block in folded.items():
            bucket = "accounts" if scope == "account" else "projects"
            target = next((x for x in profile.get(bucket) or [] if x.get("id") == rid), None)
            if target is not None:
                target["overrides"] = block

    _dump_yaml(profile, out)
    print(f"Wrote {out}")


def _read_settings(ws, spec: dict) -> dict:
    props = spec.get("properties") or {}
    result: dict = {}
    sub_start: tuple[str, list[str], int] | None = None
    for row in ws.iter_rows(min_row=5, max_col=9):
        key = row[8].value
        if key and str(key).startswith("_subtable:"):
            _, prop, cols = str(key).split(":", 2)
            sub_start = (prop, cols.split(","), row[0].row + 2)  # +1 header meanings, +1 first data row
            continue
        if not key:
            continue
        raw = row[1].value
        if "." in str(key):
            parent, child = str(key).split(".", 1)
            pspec = ((props.get(parent) or {}).get("properties") or {}).get(child, {})
            val = _unflatten(str(raw) if raw is not None else "", pspec)
            if val not in (None, "", []):
                result.setdefault(parent, {})[child] = val
        else:
            val = _unflatten(str(raw) if raw is not None else "", props.get(str(key), {}))
            if val not in (None, "", []):
                result[str(key)] = val

    if sub_start:
        prop, cols, first = sub_start
        item_props = ((props.get(prop) or {}).get("items") or {}).get("properties") or {}
        rows = []
        for r in ws.iter_rows(min_row=first, max_col=len(cols)):
            item = {}
            for ci, c in enumerate(cols):
                raw = r[ci].value if ci < len(r) else None
                v = _unflatten(str(raw) if raw is not None else "", item_props.get(c, {}))
                if v not in (None, "", []):
                    item[c] = v
            if item:
                rows.append(item)
        if rows:
            result[prop] = rows
    return result


def _read_table(ws, spec: dict, xw: dict) -> list[dict]:
    cols = xw.get("columns") or []
    item_props = ((spec.get("items") or {}).get("properties")) or {}
    out = []
    for row in ws.iter_rows(min_row=5, max_col=len(cols)):
        values = {}
        for ci, c in enumerate(cols):
            raw = row[ci].value if ci < len(row) else None
            v = _unflatten(str(raw) if raw is not None else "", item_props.get(c, {}))
            if v not in (None, "", []):
                values[c] = v
        if values.get("id") or values.get("name"):
            out.append(values)
    return out


# --------------------------------------------------------------------------------------
# tracker workbook: the per-run audit trail
# --------------------------------------------------------------------------------------


# --------------------------------------------------------------------------------------
# The tracker workbook: the per-run audit trail
#
# Colour is the whole visual grammar, and it is worth stating plainly because the obvious
# alternative - banding every other row - is decoration that tells you nothing:
#
#   yellow  yours. Change it and `workbook.py review` folds the change back in.
#   white   read from your tracker or your profile. True, but not yours to change here.
#   grey    computed. Editing it achieves nothing; the next run rebuilds it.
#
# The rest is the ordinary craft of a hand-built tracker: 16px rows so thirty fit on a
# screen instead of six, real dates so date columns sort as dates, dropdowns on the columns
# with a fixed set of answers, red on the cells you want to notice, and the formatting
# carried on past the last row so anything typed underneath still looks right.
# --------------------------------------------------------------------------------------

IN, CALC, DATA = "input", "computed", "data"          # column roles

FILL_IN = PatternFill("solid", fgColor="FFF2CC")      # yellow: yours
FILL_CALC = PatternFill("solid", fgColor="F2F2F2")    # grey: computed
FILL_DATA = PatternFill("solid", fgColor="FFFFFF")    # white: read, not derived
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_BAND = PatternFill("solid", fgColor="2F5597")    # section strip on Dashboard / Config
FILL_SOFT = PatternFill("solid", fgColor="D9E2F3")
GREEN = PatternFill("solid", fgColor="D9EAD3")
RED = PatternFill("solid", fgColor="F4CCCC")
AMBER = PatternFill("solid", fgColor="FFE599")

BASE = Font(name="Arial", size=10)
BASE_B = Font(name="Arial", size=10, bold=True)
HEADF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BANDF = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TITLEF = Font(name="Arial", size=14, bold=True, color="1F3864")
SUBF = Font(name="Arial", size=11, bold=True, color="1F3864")
NOTEF = Font(name="Arial", size=9, italic=True, color="666666")
LINKF = Font(name="Arial", size=10, color="1155CC", underline="single")
MONOF = Font(name="Menlo", size=10)
ITALF = Font(name="Arial", size=10, italic=True)

TOP = Alignment(vertical="top", wrap_text=True)
TOPC = Alignment(vertical="top", horizontal="center", wrap_text=True)
HEADA = Alignment(vertical="bottom", horizontal="left", wrap_text=True)

ROW_H = 16        # a normal data row, the height a person would leave it at
HEAD_H = 32
SPARE = 120       # rows of live formatting past the data, for anything typed by hand
DATEFMT = "yyyy-mm-dd"

KEYROW = 3        # hidden machine names
HDR = 4           # visible header
FIRST = 5         # first data row

ROLE_FILL = {IN: FILL_IN, CALC: FILL_CALC, DATA: FILL_DATA}

TAB_COLOUR = {"Read Me": "1F3864", "Dashboard": "1F3864", "Config": "2F5597",
              "Periods": "2F5597", "Task Register": "548235", "Defect Register": "548235",
              "KPI Summary": "BF8F00", "PMS Push Log": "808080", "Gaps": "808080"}

# Columns with a fixed set of answers get a dropdown, so a typo cannot quietly change a KPI.
# The lists come from the KIF schema rather than from a copy kept here, because a copy is a
# thing that goes stale: add a value to the contract and every sheet offers it the same day.
CHOICES_SECTION = {"Task Register": "tasks", "Defect Register": "defects", "Periods": "periods"}
CHOICES_ALIAS = {"item_type": "type"}          # a friendlier header over the same field
CHOICES_EXTRA = {"client_expected": "Yes,No", "team_committed": "Yes,No"}   # the two gates


def _choices(section: str) -> dict[str, str]:
    """Dropdown lists for one KIF section, read straight off the schema."""
    schema = _load_any(PLUGIN_ROOT / "schemas" / "kif.schema.json")
    props = (((schema.get("properties") or {}).get(section) or {})
             .get("items", {}).get("properties") or {})
    out = dict(CHOICES_EXTRA)
    for field, spec in props.items():
        values = [v for v in (spec.get("enum") or []) if v is not None]
        types = spec.get("type")
        if values:
            out[field] = ",".join(str(v) for v in values)
        elif types == "boolean" or (isinstance(types, list) and "boolean" in types):
            out[field] = "Yes,No"
    for alias, field in CHOICES_ALIAS.items():
        if field in out:
            out[alias] = out[field]
    return out

# Which columns a person may edit, per tab. `review` reads exactly these back; the IN role
# below paints exactly these yellow, so the colour cannot promise what the tool will refuse.
EDITABLE = {
    "Task Register": {"item_type", "type", "planned", "status", "delivered", "story_points",
                      "hours_dev", "hours_qa", "understood", "understood_why",
                      "understood_evidence", "client_expected", "client_date", "met_client_date",
                      "client_date_evidence", "team_committed", "commit_date", "met_commitment",
                      "commitment_evidence", "reopened", "rework_evidence", "exclude_reason",
                      "remarks"},
    "Defect Register": {"kind", "phase", "pre_existing", "rejected", "rejection_reason",
                        "evidence", "final_status", "remarks"},
    "Periods": {"client_date", "commit_date", "client_check", "handover_date", "team_hours",
                "notes", "plan_text"},
    "KPI Summary": {"why", "set_value_by_hand", "set_note_by_hand", "why_set_by_hand"},
}

# (key, header, width, role, is_date)
COLS_TASKS = [
    ("period", "Period", 15, DATA, 0),
    ("key", "Ticket", 12, DATA, 0),
    ("link", "Link", 7, DATA, 0),
    ("title", "Title", 44, DATA, 0),
    ("item_type", "Item Type", 10, IN, 0),
    ("planned", "Planned (initial scope)?", 11, IN, 0),
    ("status", "Status", 12, IN, 0),
    ("assignee", "Assignee", 17, DATA, 0),
    ("created", "Created", 11, DATA, 1),
    ("delivered", "Delivered", 11, IN, 1),
    ("closed", "Closed", 11, DATA, 1),
    ("hours_dev", "Dev Hours", 9, IN, 0),
    ("hours_qa", "QA Hours (counted)", 10, IN, 0),
    ("hours_total", "Est. Hours (dev + QA)", 10, CALC, 0),
    ("story_points", "Story Points", 9, IN, 0),
    ("hours_source", "Hours Source", 20, DATA, 0),
    ("understood", "Understood w/o Client?", 11, IN, 0),
    ("understood_why", "Comprehension Note", 34, IN, 0),
    ("understood_evidence", "Comprehension Evidence", 13, IN, 0),
    ("client_expected", "Client-Expected?", 10, IN, 0),
    ("client_date", "Client Expected Date", 13, IN, 1),
    ("met_client_date", "Met Client Date?", 11, IN, 0),
    ("client_date_evidence", "Client Date Evidence", 13, IN, 0),
    ("team_committed", "Team Committed?", 10, IN, 0),
    ("commit_date", "Commitment Date", 13, IN, 1),
    ("met_commitment", "Met Commitment?", 11, IN, 0),
    ("commitment_evidence", "Commitment Evidence", 13, IN, 0),
    ("reopened", "Reopened after closing?", 11, IN, 0),
    ("rework_evidence", "Rework Evidence", 30, IN, 0),
    ("exclude_reason", "Why excluded", 24, IN, 0),
    ("remarks", "Remarks and sources", 38, IN, 0),
]

COLS_DEFECTS = [
    ("period", "Period", 15, DATA, 0),
    ("key", "Ticket", 12, DATA, 0),
    ("link", "Link", 7, DATA, 0),
    ("title", "Title", 48, DATA, 0),
    ("kind", "Kind", 11, IN, 0),
    ("reported_by", "Reported By", 18, DATA, 0),
    ("reported_on", "Reported On", 12, DATA, 1),
    ("phase", "Phase", 13, IN, 0),
    ("pre_existing", "Pre-existing?", 10, IN, 0),
    ("rejected", "Rejected?", 9, IN, 0),
    ("rejection_reason", "Rejection Reason", 32, IN, 0),
    ("evidence", "Evidence", 30, IN, 0),
    ("final_status", "Final Status", 15, IN, 0),
    ("counts", "Counts in Defect Rate?", 11, CALC, 0),
    ("against_task", "Against Task", 13, DATA, 0),
    ("remarks", "Remarks and sources", 36, IN, 0),
]

COLS_PERIODS = [
    ("name", "Period Name (max 25)", 22, DATA, 0),
    ("start", "Start", 11, DATA, 1),
    ("end", "End", 11, DATA, 1),
    ("description", "PMS Description (auto)", 26, CALC, 0),
    ("client_date", "Client Expected Date", 14, IN, 1),
    ("commit_date", "Commitment Date", 14, IN, 1),
    ("client_check", "Client Date Check", 13, IN, 0),
    ("handover_date", "Handed Over On", 13, IN, 1),
    ("team_hours", "Team-level Effort (h)", 12, IN, 0),
    ("notes", "Notes", 56, IN, 0),
    ("plan_text", "Original Plan", 44, IN, 0),
    ("pms_period_id", "PMS Period ID", 12, DATA, 0),
]

COLS_SUMMARY = [
    ("period", "Period", 19, DATA, 0),
    ("kpi", "KPI", 21, DATA, 0),
    ("pms_id", "PMS KPI ID", 8, CALC, 0),
    ("numerator", "Numerator", 9, CALC, 0),
    ("denominator", "Denominator", 10, CALC, 0),
    ("value", "Value", 9, CALC, 0),
    ("target", "Target", 8, CALC, 0),
    ("min_max", "Min / Max", 8, CALC, 0),
    ("status", "Status", 12, CALC, 0),
    ("needs_reason", "Needs a reason?", 10, CALC, 0),
    ("target_set_by", "Target set by", 20, CALC, 0),
    ("note_facts", "What the numbers say (auto)", 50, CALC, 0),
    ("why", "Why / context (you write; links on words)", 52, IN, 0),
    ("note_sent_to_pms", "Note sent to PMS (auto, no links)", 62, CALC, 0),
    ("set_value_by_hand", "Set value by hand", 11, IN, 0),
    ("set_note_by_hand", "Set note by hand", 40, IN, 0),
    ("why_set_by_hand", "Why set by hand", 40, IN, 0),
]

COLS_PUSHLOG = [
    ("pushed_on", "Pushed On", 17, DATA, 0),
    ("period", "Period", 18, DATA, 0),
    ("pms_period_id", "PMS Period ID", 12, DATA, 0),
    ("action", "Action", 10, DATA, 0),
    ("changed", "What changed", 70, DATA, 0),
    ("by", "Note or pushed by", 18, DATA, 0),
    ("result", "Result", 26, DATA, 0),
]


def _yn(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    return str(v)


def _is_link(v) -> bool:
    return isinstance(v, str) and v.startswith("http")


def _as_date(value):
    """A real date if it looks like one, so the column sorts and filters as dates."""
    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    if isinstance(value, str) and len(value) >= 10:
        try:
            return _dt.datetime.strptime(value[:10], "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def _cell(value):
    """Numbers as numbers, booleans as Yes/No, everything else as text."""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else round(float(value), 2)
    return _flatten(value)


def _sheet(wb, title: str, heading: str, subtitle: str = "", gutter: int = 4):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TAB_COLOUR.get(title, "808080")
    ws.column_dimensions["A"].width = gutter
    ws.cell(row=1, column=2, value=heading).font = TITLEF
    ws.row_dimensions[1].height = 22
    if subtitle:
        c = ws.cell(row=2, column=2, value=subtitle)
        c.font, c.alignment = NOTEF, Alignment(vertical="center", wrap_text=False)
    return ws


def _band(ws, row: int, label: str, last_col: int) -> None:
    """A section strip, so a long tab reads as a few blocks rather than one wall."""
    for col in range(2, last_col + 1):
        ws.cell(row=row, column=col).fill = FILL_BAND
    ws.cell(row=row, column=2, value=label).font = BANDF
    ws.row_dimensions[row].height = 18


def _headers(ws, cols, row: int) -> None:
    c = ws.cell(row=row, column=1, value="#")
    c.fill, c.font, c.border, c.alignment = FILL_HEAD, HEADF, BORDER, HEADA
    for i, (key, header, width, _role, _d) in enumerate(cols, start=2):
        ws.cell(row=KEYROW, column=i, value=key).font = Font(size=8, color="FFFFFF")
        c = ws.cell(row=row, column=i, value=header)
        c.fill, c.font, c.border, c.alignment = FILL_HEAD, HEADF, BORDER, HEADA
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[KEYROW].hidden = True
    ws.row_dimensions[row].height = HEAD_H


def _register(wb, title: str, subtitle: str, cols, rows: list[dict],
              freeze: str, row_height: int = ROW_H, spare: int = SPARE):
    """One list tab: coloured by role, sized to be read without dragging anything."""
    ws = _sheet(wb, title, f"{title}  ·  {len(rows)} rows", subtitle)
    _headers(ws, cols, HDR)
    ws.freeze_panes = freeze

    last = FIRST + len(rows) + spare - 1
    for i in range(1, len(rows) + spare + 1):
        r = FIRST + i - 1
        row = rows[i - 1] if i <= len(rows) else {}
        n = ws.cell(row=r, column=1, value=i if i <= len(rows) else None)
        n.fill, n.font, n.border, n.alignment = FILL_CALC, BASE, BORDER, TOPC
        for ci, (key, _hd, _w, role, is_date) in enumerate(cols, start=2):
            value = row.get(key)
            cell = ws.cell(row=r, column=ci)
            d = _as_date(value) if is_date else None
            if d is not None:
                cell.value, cell.number_format, cell.font = d, DATEFMT, BASE
            elif _is_link(value):
                cell.value, cell.hyperlink, cell.font = "Open", value, LINKF
            else:
                cell.value, cell.font = _cell(value), BASE
            cell.fill, cell.border, cell.alignment = ROLE_FILL[role], BORDER, TOP
        ws.row_dimensions[r].height = row_height

    choices = _choices(CHOICES_SECTION[title]) if title in CHOICES_SECTION else {}
    for ci, (key, _hd, _w, role, _d) in enumerate(cols, start=2):
        if role is IN and key in choices and len(choices[key]) < 250:
            dv = DataValidation(type="list", formula1=f'"{choices[key]}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(ci)}{FIRST}:{get_column_letter(ci)}{last}")

    if rows:
        ws.auto_filter.ref = f"A{HDR}:{get_column_letter(len(cols) + 1)}{FIRST + len(rows) - 1}"
    ws.print_title_rows = f"{HDR}:{HDR}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws, last


def _flag(ws, cols, key: str, equals: str, fill, last: int) -> None:
    """Keep a rule alive on the column rather than colouring cells once: an edit made in
    Excel then shows the same colour this run would have given it."""
    if key not in [c[0] for c in cols]:
        return
    ci = get_column_letter(2 + [c[0] for c in cols].index(key))
    ws.conditional_formatting.add(
        f"{ci}{FIRST}:{ci}{last}",
        CellIsRule(operator="equal", formula=[f'"{equals}"'], fill=fill))


def tracker(results_path: Path, kif_path: Path, out: Path, reasons_path: Path | None = None,
            profile_path: Path | None = None) -> None:
    """Build the working file: nine tabs laid out the way a hand-built tracker is, so
    anybody who has read one project's tracker can read anybody else's."""
    results = _load_any(results_path)
    kif = _load_any(kif_path)
    reasons = _load_any(reasons_path) if reasons_path and reasons_path.exists() else {}
    profile = _load_any(profile_path) if profile_path and profile_path.exists() else {}
    proj = kif.get("project") or {}
    periods = kif.get("periods") or []
    wb = Workbook()
    wb.remove(wb.active)
    wb._named_styles["Normal"].font = Font(name="Arial", size=10)

    reg_path = PLUGIN_ROOT / "schemas" / "kpi_registry.json"
    reg = _load_any(reg_path if reg_path.exists()
                    else PLUGIN_ROOT / "schemas" / "kpi_registry.default.json")
    reg_by_name = {k["name"]: k for k in reg.get("kpis", [])}

    # ------------------------------------------------------------------ Read Me
    ws = _sheet(wb, "Read Me", f"KPI Tracker  ·  {proj.get('name', '')}",
                "Built and refreshed by KPI Copilot. Edit the yellow cells, run the review "
                "step, and every number and note follows from them.")
    ws.column_dimensions["B"].width = 112
    lines = [
        ("", ""),
        ("Colour tells you what a cell is", "h"),
        ("Yellow — yours. Change it, then fold it back with  workbook.py review.", "in"),
        ("White — read from your tracker or your profile. True, but not yours to change here.", "data"),
        ("Grey — computed. Editing it achieves nothing; the next run rebuilds it.", "calc"),
        ("", ""),
        ("Tabs", "h"),
        ("Dashboard — totals per period, the nine KPIs with colour, and how many are met.", ""),
        ("Config — project details, counting rules, and each KPI with the PMS formula behind it.", ""),
        ("Periods — one row per KPI period with its dates, notes and PMS period id.", ""),
        ("Task Register — every planned task and addition, with the evidence behind each judgement.", ""),
        ("Defect Register — every report, including the ones not counted and why.", ""),
        ("KPI Summary — the nine KPIs per period, the facts, your reason, and the note sent to PMS.", ""),
        ("PMS Push Log — what was sent to PMS and when.", ""),
        ("Gaps — anything this run could not measure, and why.", ""),
        ("", ""),
        ("Things worth knowing", "h"),
        ("Dates: a date on the task wins, then the period date, then the project date.", ""),
        ("Client date check: Handover = the period's build reached the client by the date. "
         "Delivery = the item itself was delivered by the date.", ""),
        ("Delivery Commitment counts only the items the team committed to. Items with no "
         "commitment are left out and named — it measures promises kept, not work finished.", ""),
        ("Escaped Defect Rate divides by defects, not by delivered items.", ""),
        ("If a computed figure is wrong and you cannot fix the input, put the right one in "
         "'Set value by hand' on KPI Summary with a reason. Both figures are then shown.", ""),
        ("Links sit on the words. The note sent to PMS never contains a link.", ""),
        ("Formatting carries on past the last row, so anything you add by hand still fits.", ""),
        ("", ""),
        ("Note format sent to PMS", "h"),
        ("what is measured || the numbers || what was left out || why", "code"),
        ("", ""),
        ("Period names", "h"),
        ("Initial Scope, Additional Requests 1, 2 …, Milestone 1, 2 …, Full Project. "
         "PMS allows 25 characters.", ""),
    ]
    r = 3
    for text, kind in lines:
        c = ws.cell(row=r, column=2, value=text)
        c.alignment = TOP
        c.font = {"h": SUBF, "code": MONOF}.get(kind, BASE)
        if kind in ("in", "data", "calc"):
            c.fill = {"in": FILL_IN, "data": FILL_DATA, "calc": FILL_CALC}[kind]
            c.border = BORDER
        elif kind == "code":
            c.fill, c.border = FILL_SOFT, BORDER
        ws.row_dimensions[r].height = 26 if kind in ("in", "data", "calc", "code") else ROW_H
        r += 1

    # ------------------------------------------------------------------ Dashboard
    ws = _sheet(wb, "Dashboard", f"KPI Dashboard  ·  {proj.get('name', '')}",
                f"PMS project {proj.get('pms_project_id') or 'not set'}   ·   "
                f"velocity in {proj.get('velocity_unit', 'Estimated Hours')}   ·   "
                f"data read {(kif.get('generated') or {}).get('at', '')[:10]}")

    dash = [("Period", 20), ("Delivered", 10), ("Not delivered", 12), ("Planned", 9),
            ("Additional", 10), ("Velocity", 11), ("Issues reported", 12), ("Bugs counted", 11),
            ("Rejected", 9), ("Found after handover", 13), ("Client date", 12),
            ("Commit date", 12),
            ("Handed over", 12)]
    _band(ws, 4, "WORK AND ISSUES BY PERIOD", len(dash) + 1)
    hr = 5
    c = ws.cell(row=hr, column=1, value="#")
    c.fill, c.font, c.border, c.alignment = FILL_HEAD, HEADF, BORDER, HEADA
    ws.column_dimensions["A"].width = 4
    for i, (h, w) in enumerate(dash, start=2):
        c = ws.cell(row=hr, column=i, value=h)
        c.fill, c.font, c.border, c.alignment = FILL_HEAD, HEADF, BORDER, HEADA
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hr].height = HEAD_H

    unit = "pts" if proj.get("velocity_unit") == "Story Points" else "h"
    tot = [0] * 6
    r = hr + 1
    for i, per in enumerate(results["periods"], start=1):
        f = per["facts"]
        by_name = {m["name"]: m for m in per["measures"]}
        vel = by_name.get("Velocity", {})
        esc = by_name.get("Escaped Defect Rate", {})
        rej = by_name.get("Defect Rejection Rate", {})
        vals = [i, per["period"], f["delivered"], f["not_delivered"], f["planned"],
                f["additional_requests"],
                f"{_cell(vel.get('value'))} {unit}" if vel.get("value") is not None else "—",
                f["reports"], f["defects_counted"], rej.get("numerator") or 0,
                esc.get("numerator") or 0, f.get("client_date") or "—",
                f.get("commit_date") or "—", f.get("handover_date") or "Not yet"]
        for ci, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=ci, value=v)
            cell.font, cell.border, cell.alignment = BASE, BORDER, TOP
            cell.fill = FILL_CALC if ci == 1 else FILL_DATA
        for j, k in enumerate(("delivered", "not_delivered", "planned", "additional_requests",
                               "reports", "defects_counted")):
            tot[j] += f[k] or 0
        ws.row_dimensions[r].height = ROW_H
        r += 1
    c = ws.cell(row=r, column=2, value="Total")
    c.font, c.border, c.fill = BASE_B, BORDER, FILL_SOFT
    for ci, v in zip((3, 4, 5, 6, 8, 9), tot):
        cell = ws.cell(row=r, column=ci, value=v)
        cell.font, cell.border, cell.fill, cell.alignment = BASE_B, BORDER, FILL_SOFT, TOPC

    # The KPI matrix: one row per KPI, one column per period.
    r += 2
    ncols = 3 + max(len(results["periods"]), 1)
    _band(ws, r, "KPI RESULTS BY PERIOD", ncols)
    r += 1
    for lbl, col in (("KPI", 2), ("Target", 3)):
        c = ws.cell(row=r, column=col, value=lbl)
        c.fill, c.font, c.border, c.alignment = FILL_HEAD, HEADF, BORDER, HEADA
    for j, per in enumerate(results["periods"]):
        c = ws.cell(row=r, column=4 + j, value=per["period"])
        c.fill, c.font, c.border, c.alignment = FILL_HEAD, HEADF, BORDER, HEADA
        ws.column_dimensions[get_column_letter(4 + j)].width = 19
    ws.row_dimensions[r].height = HEAD_H
    head_r = r

    names = [m["name"] for m in results["periods"][0]["measures"]] if results["periods"] else []
    met = notmet = notmeas = noreason = 0
    for k, name in enumerate(names, start=1):
        rr = head_r + k
        c = ws.cell(row=rr, column=2, value=name)
        c.font, c.border, c.fill, c.alignment = BASE, BORDER, FILL_DATA, TOP
        first = next((m for m in results["periods"][0]["measures"] if m["name"] == name), {})
        tgt = "" if first.get("threshold") is None else (
            ("at least " if first.get("direction") == "higher-is-better" else "at most ")
            + _flatten(first["threshold"]) + ("%" if first.get("unit") == "%" else ""))
        c = ws.cell(row=rr, column=3, value=tgt)
        c.font, c.border, c.fill, c.alignment = BASE, BORDER, FILL_CALC, TOP
        for j, per in enumerate(results["periods"]):
            m = next((x for x in per["measures"] if x["name"] == name), {})
            c = ws.cell(row=rr, column=4 + j,
                        value=_cell(m.get("value")) if m.get("value") is not None else "—")
            c.border, c.alignment = BORDER, TOPC
            c.font = ITALF if m.get("overridden") else BASE
            c.fill = {"Met": GREEN, "Not met": RED}.get(m.get("status"), FILL_CALC)
            st = m.get("status")
            met += st == "Met"
            notmet += st == "Not met"
            notmeas += st == "Not measured"
            if st == "Not met" and not (reasons.get(per["period"]) or {}).get(name):
                noreason += 1
        ws.row_dimensions[rr].height = ROW_H

    r = head_r + len(names) + 2
    _band(ws, r, "STATUS ACROSS ALL PERIODS", 9)
    for j, (label, val, fill) in enumerate([("Met", met, GREEN), ("Not met", notmet, RED),
                                            ("Not measured", notmeas, FILL_CALC),
                                            ("Reasons missing", noreason, AMBER)]):
        c = ws.cell(row=r + 1, column=2 + j * 2, value=label)
        c.font, c.alignment = BASE_B, TOP
        v = ws.cell(row=r + 1, column=3 + j * 2, value=val)
        v.fill, v.border, v.font, v.alignment = fill, BORDER, BASE_B, TOPC
    ws.row_dimensions[r + 1].height = ROW_H

    # Quick links, straight from the profile's tool registry: the map of the project.
    tools = [t for t in (profile.get("tools") or [])
             if str(t.get("url_or_id", "")).startswith("http")]
    if tools:
        r += 3
        _band(ws, r, "WHERE EVERYTHING LIVES", 9)
        for j, t in enumerate(tools[:15]):
            cell = ws.cell(row=r + 1 + j // 3, column=2 + (j % 3) * 2,
                           value=f"{t.get('name')} ↗")
            cell.hyperlink, cell.font, cell.alignment = t["url_or_id"], LINKF, TOP
            ws.cell(row=r + 1 + j // 3, column=3 + (j % 3) * 2,
                    value=(t.get("what_for") or t.get("description") or "")).font = NOTEF
            ws.row_dimensions[r + 1 + j // 3].height = ROW_H

    # ------------------------------------------------------------------ Config
    ws = _sheet(wb, "Config", "Configuration",
                "Yellow is a setting you own. Targets are read from PMS and shown here for "
                "reference — change them on the project in PMS, not in this sheet.")
    for col, w in (("B", 31), ("C", 42), ("D", 62), ("E", 9), ("F", 11), ("G", 50), ("H", 46)):
        ws.column_dimensions[col].width = w

    def kv(row: int, label: str, value, hint: str = "", mine: bool = False) -> None:
        a = ws.cell(row=row, column=2, value=label)
        a.font, a.border, a.alignment = BASE, BORDER, TOP
        c = ws.cell(row=row, column=3, value=_cell(value))
        c.border, c.alignment, c.font = BORDER, TOP, BASE
        c.fill = FILL_IN if mine else FILL_DATA
        if hint:
            h = ws.cell(row=row, column=4, value=hint)
            h.font, h.alignment = NOTEF, TOP
        ws.row_dimensions[row].height = ROW_H

    out_cfg = profile.get("output") or {}
    src = profile.get("sources") or {}
    pol = profile.get("policy") or {}
    base = (profile.get("organization") or {}).get("pms_base_url", "")

    _band(ws, 4, "PROJECT", 8)
    facts = [("Project name", proj.get("name"), ""),
             ("Client / account", (profile.get("_account") or {}).get("name")
              or (profile.get("owner") or {}).get("account"), ""),
             ("PMS project ID", proj.get("pms_project_id"), ""),
             ("PMS KPI page", f"{base}/all-projects/{proj.get('pms_project_id')}/kpis"
              if (base and proj.get("pms_project_id")) else "", "Where these numbers land."),
             ("Issue tracker", proj.get("tracker"), ""),
             ("Tracker board", proj.get("tracker_url"), ""),
             ("Period type", proj.get("period_type"), ""),
             ("Plan / estimate sources", proj.get("sources_text"), ""),
             ("Data read on", (kif.get("generated") or {}).get("at", "")[:16].replace("T", " "), ""),
             ("Adapter", (kif.get("generated") or {}).get("adapter"), ""),
             ("Delivery of results", out_cfg.get("mode", "review"),
              "review = you approve each push. auto-push = sent without asking.")]
    for i, (k, v, h) in enumerate(facts):
        kv(5 + i, k, v, h)

    r = 5 + len(facts) + 1
    _band(ws, r, "COUNTING RULES  ·  these decide what the numbers mean", 8)
    crules = [("Velocity unit", proj.get("velocity_unit"),
               "What Velocity adds up for delivered items."),
              ("Velocity counts", src.get("hours_basis", "dev"),
               "dev = development hours only. dev+qa adds each item's QA hours once QA is done."),
              ("Count observations as defects?", _yn(pol.get("count_observations", False)),
               "No = only Bug rows count in Defect Rate."),
              ("Count improvements as defects?", _yn(pol.get("count_improvements", False)), ""),
              ("Count pre-existing defects?", _yn(pol.get("count_pre_existing", False)),
               "Pre-existing = already in the product before this work."),
              ("Count post-release defects in Defect Rate?", _yn(pol.get("count_post_release", True)),
               "They always count in Escaped Defect Rate."),
              ("When plan and tracker disagree on hours, trust", src.get("hours_first", "tracker"), ""),
              ("A commitment is met on", ((profile.get("workflow") or {}).get("commitment") or {})
               .get("met_when", "delivery"), "Printed in the Delivery Commitment note.")]
    for i, (k, v, h) in enumerate(crules):
        kv(r + 1 + i, k, v, h, mine=True)

    r += len(crules) + 2
    _band(ws, r, "KPI TARGETS  ·  read from PMS; amber means this project sets its own", 8)
    hdr_cols = ["KPI", "Unit", "PMS KPI ID", "Type", "Target", "Formula (PMS definition)",
                "How this sheet counts it"]
    for i, h in enumerate(hdr_cols, start=2):
        c = ws.cell(row=r + 1, column=i, value=h)
        c.fill, c.font, c.border, c.alignment = FILL_HEAD, HEADF, BORDER, HEADA
    ws.row_dimensions[r + 1].height = HEAD_H
    firsts = results["periods"][0]["measures"] if results["periods"] else []
    for i, m in enumerate(firsts):
        rr = r + 2 + i
        meta = reg_by_name.get(m["name"], {})
        own = "for project" in (m.get("threshold_source") or "")
        vals = [m["name"], m.get("unit"), m.get("pms_id"),
                "Min" if m.get("direction") == "higher-is-better" else "Max",
                m.get("threshold"), meta.get("formula_pms", ""), meta.get("basis", "")]
        for ci, v in enumerate(vals, start=2):
            c = ws.cell(row=rr, column=ci, value=_cell(v))
            c.border, c.alignment, c.font = BORDER, TOP, BASE
            c.fill = AMBER if (ci == 6 and own) else FILL_CALC
        ws.row_dimensions[rr].height = 38

    r += len(firsts) + 3
    _band(ws, r, "PROJECT DATES  ·  used when a period or a task has no date of its own", 8)
    dates = [("Client expected date", proj.get("client_date"),
              "What the client expects: the planned handover, or a client deadline."),
             ("Team commitment date", proj.get("commit_date"),
              "What the team negotiated and promised."),
             ("Client date check", proj.get("client_check"),
              "Handover = met when the build reached the client. Delivery = when the item was."),
             ("Why these dates", proj.get("dates_why"),
              "Short reason, with the source on the words.")]
    for i, (k, v, h) in enumerate(dates):
        kv(r + 1 + i, k, v, h, mine=True)

    # ------------------------------------------------------------------ Periods
    _register(wb, "Periods",
              "One row per KPI period. Leave a date empty to fall back to the project date "
              "on Config.",
              COLS_PERIODS,
              [dict(p, description=next((x["description"] for x in results["periods"]
                                         if x["period"] == p.get("name")), ""))
               for p in periods], freeze="B5", spare=20)

    # ------------------------------------------------------------------ Task Register
    tasks = [dict(t, item_type=t.get("type"), planned=_yn(t.get("planned")),
                  hours_total=((t.get("hours_dev") or 0) + (t.get("hours_qa") or 0)) or None,
                  client_expected="Yes" if (t.get("met_client_date")
                                            or t.get("client_date")) else "",
                  team_committed="Yes" if (t.get("commit_date")
                                           or t.get("met_commitment")) else "")
             for t in kif["tasks"]]
    ws_t, last_t = _register(
        wb, "Task Register",
        "Every planned task and addition, with the evidence behind each judgement. "
        "Setting a gate to No takes the item out of that KPI.",
        COLS_TASKS, tasks, freeze="E5")
    for key, bad in (("met_commitment", "No"), ("met_client_date", "No"), ("reopened", "Yes")):
        _flag(ws_t, COLS_TASKS, key, bad, RED, last_t)
    for key in ("met_commitment", "met_client_date"):
        _flag(ws_t, COLS_TASKS, key, "Yes", GREEN, last_t)
        _flag(ws_t, COLS_TASKS, key, "Pending", AMBER, last_t)
    _flag(ws_t, COLS_TASKS, "understood", "No", AMBER, last_t)

    # ------------------------------------------------------------------ Defect Register
    counted = set()
    for per in results["periods"]:
        for m in per["measures"]:
            if m["name"] == "Defect Rate":
                counted |= set(m.get("counted_keys") or [])
    defects = [dict(d, counts="Yes" if d.get("key") in counted else "No") for d in kif["defects"]]
    ws_d, last_d = _register(
        wb, "Defect Register",
        "Every report, including the ones not counted and why. A rejected report still "
        "belongs here — it is the numerator of the rejection rate.",
        COLS_DEFECTS, defects, freeze="E5")
    _flag(ws_d, COLS_DEFECTS, "rejected", "Yes", RED, last_d)
    _flag(ws_d, COLS_DEFECTS, "phase", "Post-release", RED, last_d)
    _flag(ws_d, COLS_DEFECTS, "counts", "Yes", AMBER, last_d)

    # ------------------------------------------------------------------ KPI Summary
    summary_rows = []
    for per in results["periods"]:
        for m in per["measures"]:
            why = ((reasons.get(per["period"]) or {}).get(m["name"]) or "")
            shown = m.get("computed_value") if m.get("overridden") else m["value"]
            facts3 = (m.get("note_parts") or ["", "", ""])[:3]
            summary_rows.append({
                "period": per["period"], "kpi": m["name"], "pms_id": m.get("pms_id"),
                "numerator": m.get("numerator"), "denominator": m.get("denominator"),
                "value": shown, "target": m.get("threshold"),
                "min_max": "Min" if m.get("direction") == "higher-is-better" else "Max",
                "target_set_by": m.get("threshold_source", ""), "status": m["status"],
                "needs_reason": "Yes" if (m["status"] == "Not met" and not why) else "No",
                "note_facts": " || ".join(x for x in facts3 if x),
                "why": why, "note_sent_to_pms": m["note"],
                "set_value_by_hand": m["value"] if m.get("overridden") else "",
                "set_note_by_hand": m["note"] if (m.get("overridden")
                                                  and m.get("computed_note") != m["note"]) else "",
                "why_set_by_hand": m.get("override_reason", "")})
    ws_s, last_s = _register(
        wb, "KPI Summary",
        "Grey is rebuilt every run. Yellow is yours: 'Why' becomes the fourth part of the "
        "note; 'Set value by hand' overrides a figure and needs a reason beside it.",
        COLS_SUMMARY, summary_rows, freeze="C5", row_height=44, spare=12)
    _flag(ws_s, COLS_SUMMARY, "status", "Met", GREEN, last_s)
    _flag(ws_s, COLS_SUMMARY, "status", "Not met", RED, last_s)
    _flag(ws_s, COLS_SUMMARY, "status", "Not measured", AMBER, last_s)
    _flag(ws_s, COLS_SUMMARY, "needs_reason", "Yes", RED, last_s)

    # ------------------------------------------------------------------ PMS Push Log
    log_path = results_path.parent / "push_log.json"
    log_rows: list[dict] = []
    if log_path.exists():
        try:
            log = _load_any(log_path)
            for res in (log.get("results") or []):
                log_rows.append({
                    "pushed_on": log.get("at", "")[:16].replace("T", " "),
                    "period": res.get("period"), "pms_period_id": res.get("periodId"),
                    "action": "Update",
                    "changed": "; ".join(f"{k['name']}: {_cell(k['value'])}"
                                         for k in res.get("kpis") or []),
                    "by": "", "result": res.get("result")})
        except Exception:  # noqa: BLE001
            pass
    _register(wb, "PMS Push Log",
              "One line each time values are sent to PMS. Written automatically after a push.",
              COLS_PUSHLOG, log_rows, freeze="B5", spare=30)

    # ------------------------------------------------------------------ Gaps
    ws = _sheet(wb, "Gaps", "What this run could not measure",
                "An empty tab is a good sign. Each line says what would close the gap.")
    gaps_cols = [("period", "Period", 20, DATA, 0), ("kpi", "KPI", 21, DATA, 0),
                 ("gap", "Reason, and what would fix it", 92, DATA, 0)]
    _headers(ws, gaps_cols, HDR)
    ws.freeze_panes = "A5"
    r, n = FIRST, 0
    for per in results["periods"]:
        for m in per["measures"]:
            for g in m.get("gaps") or []:
                n += 1
                for ci, v in enumerate([n, per["period"], m["name"], g], start=1):
                    c = ws.cell(row=r, column=ci, value=v)
                    c.border, c.alignment, c.font = BORDER, TOP, BASE
                    c.fill = FILL_CALC if ci == 1 else FILL_DATA
                ws.row_dimensions[r].height = ROW_H
                r += 1
    if n == 0:
        c = ws.cell(row=FIRST, column=2,
                    value="Nothing. Every KPI was measurable from the sources available.")
        c.font = NOTEF

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out} ({len(wb.sheetnames)} tabs)")


# --------------------------------------------------------------------------------------
# review: read a person's edits back out of the tracker workbook
# --------------------------------------------------------------------------------------


# Some workbook columns are a friendlier view of a KIF field, or two of them. Mapping them
# back explicitly beats letting the names drift apart silently.
WB_TO_KIF = {
    "Task Register": {"item_type": "type"},
    "Defect Register": {},
    "Periods": {},
}


# Columns that are a view of something else, or are computed. The direct loop skips them:
# the fold handles the first kind and nothing should write the second.
VIEW_COLUMNS = {"client_expected", "team_committed", "hours_total", "counts", "description",
                "pushed_on"}


def _fold_task_row(row: dict, hit: dict, note) -> set[str]:
    """The two columns that are a gate rather than a field: whether an item was
    client-expected, and whether the team committed to it at all."""
    # "Client-Expected?" and "Team Committed?" are gates. Turning one off takes the item out
    # of that KPI entirely, which is a real and useful edit - an item nobody promised should
    # not be counted against the team's reliability.
    locked: set[str] = set()
    for gate, judged, dated in (("client_expected", "met_client_date", "client_date"),
                                ("team_committed", "met_commitment", "commit_date")):
        val = str(row.get(gate) or "").strip().lower()
        if not val:
            continue
        if val in ("no", "false"):
            if hit.get(judged) is not None or hit.get(dated):
                hit[judged], hit[dated] = None, None
                note(f"{row.get('key')}: {gate} set to No, so it is out of that KPI")
        elif val in ("yes", "true") and hit.get(judged) is None:
            hit[judged] = "Pending"
            note(f"{row.get('key')}: {gate} set to Yes, now Pending until judged")
        if val in ("no", "false"):
            # The gate wins: do not let the raw judged/date columns put it back.
            locked.update({judged, dated})
    return locked


def _read_register(ws, keyrow: int, first_data_row: int) -> list[dict]:
    """Rows keyed by the hidden machine names, so a reordered or hidden column still lands in
    the right field."""
    cols = {}
    for cell in ws[keyrow]:
        if cell.value:
            cols[cell.column] = str(cell.value)
    out = []
    for row in ws.iter_rows(min_row=first_data_row):
        # A link cell reads as the word on it, so take the target instead - otherwise a
        # round trip quietly replaces every URL with "Open".
        values = {cols[c.column]: (c.hyperlink.target if c.hyperlink else c.value)
                  for c in row if c.column in cols}
        if any(v not in (None, "") for v in values.values()):
            out.append(values)
    return out


def _norm(v: Any) -> str:
    """One spelling for comparison. A whole number is '100' whether it arrived as 100, 100.0
    or the string '100.0' - otherwise an untouched cell reads as an edit."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)):
        return str(int(v)) if float(v).is_integer() else f"{float(v):.2f}".rstrip("0").rstrip(".")
    text = str(v).strip()
    try:
        f = float(text)
    except ValueError:
        return text
    return str(int(f)) if f.is_integer() else f"{f:.2f}".rstrip("0").rstrip(".")


def review(tracker_path: Path, kif_path: Path, out_kif: Path, out_reasons: Path,
           reasons_path: Path | None = None, out_manual: Path | None = None,
           manual_path: Path | None = None, by: str = "",
           results_path: Path | None = None) -> int:
    """Fold a person's edits in the tracker workbook back into the KIF and the reasons file.

    The point is that the sheet is a working surface, not a report you read and then retype
    somewhere else. Edit a Yes/No, a date, an hour figure or the 'Why' text, run this, rerun
    the engine, and the numbers and the notes follow. What you cannot do is edit the computed
    note itself - that would make the sheet and PMS say different things about the same
    number, so those edits are reported and dropped, with the reason.
    """
    wb = load_workbook(tracker_path, data_only=True)
    kif = _load_any(kif_path)
    reasons = _load_any(reasons_path) if reasons_path and reasons_path.exists() else {}
    manual = _load_any(manual_path) if manual_path and manual_path.exists() else {}
    # What the sheet said when it was built, so an edit to a computed cell can be spotted and
    # pointed at the column that would actually have worked.
    computed: dict[tuple[str, str], dict] = {}
    if results_path and results_path.exists():
        for per in (_load_any(results_path).get("periods") or []):
            for m in per.get("measures") or []:
                computed[(per["period"], m["name"])] = m

    changes: list[str] = []
    refused: list[str] = []

    def apply(rows: list[dict], target: list[dict], key_fields: tuple[str, ...], tab: str) -> None:
        editable = EDITABLE[tab]
        index = {tuple(_norm(t.get(k)) for k in key_fields): t for t in target}
        for row in rows:
            hit = index.get(tuple(_norm(row.get(k)) for k in key_fields))
            if hit is None:
                refused.append(f"{tab}: no row in the extract matches "
                               f"{' / '.join(_norm(row.get(k)) for k in key_fields)}. "
                               f"Rows cannot be added here - add them in the tracker or the profile.")
                continue
            handled: set[str] = set()
            if tab == "Task Register":
                handled = _fold_task_row(row, hit, changes.append)
            mapping = WB_TO_KIF.get(tab, {})
            for field, new in row.items():
                field = mapping.get(field, field)
                if field in VIEW_COLUMNS or field in handled:
                    continue          # a view of another field, folded above, or computed
                if field not in editable and mapping.get(field, field) not in editable:
                    if _norm(new) != _norm(hit.get(field)):
                        refused.append(f"{tab} {_norm(row.get(key_fields[-1]))}: '{field}' is computed, "
                                       f"so the edit was ignored. Change the input and rerun.")
                    continue
                old = hit.get(field)
                if _norm(new) == _norm(old):
                    continue
                if field == "type" and str(new).strip() not in ("Task", "CR", "Scope", "Excluded"):
                    refused.append(f"{tab} {_norm(row.get(key_fields[-1]))}: item type "
                                   f"'{new}' is not Task, CR, Scope or Excluded.")
                    continue
                if field in ("hours_dev", "hours_qa", "team_hours", "story_points"):
                    try:
                        new = float(new) if _norm(new) else None
                    except (TypeError, ValueError):
                        refused.append(f"{tab} {_norm(row.get(key_fields[-1]))}: '{new}' is not a number.")
                        continue
                elif field == "planned":
                    new = _norm(new).lower() in ("yes", "true")
                else:
                    new = _norm(new) or None
                changes.append(f"{tab:<16} {_norm(row.get(key_fields[-1])):<22} {field:<22} "
                               f"{_norm(old) or '(blank)'} -> {_norm(new) or '(blank)'}")
                hit[field] = new

    if "Task Register" in wb.sheetnames:
        apply(_read_register(wb["Task Register"], KEYROW, FIRST), kif["tasks"], ("period", "key"), "Task Register")
    if "Defect Register" in wb.sheetnames:
        apply(_read_register(wb["Defect Register"], KEYROW, FIRST), kif["defects"], ("period", "key"), "Defect Register")
    if "Periods" in wb.sheetnames:
        apply(_read_register(wb["Periods"], KEYROW, FIRST), kif["periods"], ("name",), "Periods")

    # The 'Why' column, which is the only free text a person owns.
    if "KPI Summary" in wb.sheetnames:
        ws = wb["KPI Summary"]
        for row in _read_register(ws, KEYROW, FIRST):
            period, kpi, why = _norm(row.get("period")), _norm(row.get("kpi")), _norm(row.get("why"))
            if not (period and kpi):
                continue
            old = _norm((reasons.get(period) or {}).get(kpi))
            if why != old:
                changes.append(f"{'KPI Summary':<16} {period + ' / ' + kpi:<22} {'why':<22} "
                               f"{(old[:40] + '...') if len(old) > 40 else (old or '(blank)')} -> "
                               f"{(why[:40] + '...') if len(why) > 40 else (why or '(blank)')}")
                reasons.setdefault(period, {})[kpi] = why
            # A hand-set value or note. Allowed, and recorded with its reason - the engine
            # refuses one without a reason, because it prints these above the numbers.
            # Somebody typed over a grey cell. It would be regenerated on the next run, so
            # say what they should have done rather than losing the edit in silence.
            was = computed.get((period, kpi))
            if was is not None:
                shown = was.get("computed_value") if was.get("overridden") else was.get("value")
                expected = _norm(shown) if shown is not None else ""
                if _norm(row.get("value")) != expected:
                    refused.append(
                        f"KPI Summary {period} / {kpi}: 'Value' is computed, so the edit was not kept. "
                        f"If the figure really is wrong, put it in 'Set value by hand' with a reason, "
                        f"or fix the register rows and rerun.")
                if _norm(row.get("note_sent_to_pms")) != _norm(was.get("note")):
                    refused.append(
                        f"KPI Summary {period} / {kpi}: 'Note sent to PMS' is assembled from the numbers "
                        f"and your 'Why', so the edit was not kept. Put wording in 'Why', or the whole "
                        f"note in 'Set note by hand' with a reason.")

            set_value = _norm(row.get("set_value_by_hand"))
            set_note = _norm(row.get("set_note_by_hand"))
            set_why = _norm(row.get("why_set_by_hand"))
            if set_value or set_note:
                if not set_why:
                    refused.append(f"KPI Summary {period} / {kpi}: a hand-set value needs a reason in "
                                   f"'Why set by hand'. It is printed with the numbers, so a reader can "
                                   f"see a person set it.")
                else:
                    entry: dict = {"why": set_why}
                    if set_value:
                        entry["value"] = set_value
                    if set_note:
                        entry["note"] = set_note
                    if by:
                        entry["by"] = by
                    manual.setdefault(period, {})[kpi] = entry
                    changes.append(f"{'KPI Summary':<16} {period + ' / ' + kpi:<22} "
                                   f"{'set by hand':<22} {set_value or 'note only'}")
            elif (manual.get(period) or {}).pop(kpi, None) is not None:
                changes.append(f"{'KPI Summary':<16} {period + ' / ' + kpi:<22} "
                               f"{'hand-set value':<22} cleared")

    out_kif.parent.mkdir(parents=True, exist_ok=True)
    out_kif.write_text(json.dumps(kif, indent=2), encoding="utf-8")
    _dump_yaml(reasons, out_reasons)
    if out_manual is not None:
        _dump_yaml({k: v for k, v in manual.items() if v}, out_manual)

    if changes:
        print(f"{len(changes)} edit(s) read back from {tracker_path.name}:\n")
        for c in changes:
            print("  " + c)
    else:
        print(f"No edits found in {tracker_path.name}.")
    if refused:
        print(f"\n{len(refused)} edit(s) not applied:\n")
        for c in dict.fromkeys(refused):
            print("  " + c)
    tail = f", {out_manual}" if out_manual is not None else ""
    print(f"\nWrote {out_kif}, {out_reasons}{tail}.")
    print("Now rerun the engine against the patched extract; the notes and values follow from it,\n"
          "which is what keeps the sheet and PMS saying the same thing.")
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Build and read the KPI Copilot workbooks.")
    sub = ap.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("build", help="Profile -> KPI Profile Workbook")
    b.add_argument("--profile", type=Path)
    b.add_argument("--out", type=Path, required=True)

    r = sub.add_parser("read", help="KPI Profile Workbook -> profile.yaml")
    r.add_argument("--xlsx", type=Path, required=True)
    r.add_argument("--out", type=Path, required=True)

    t = sub.add_parser("tracker", help="Results + KIF -> per-project tracker workbook")
    t.add_argument("--results", type=Path, required=True)
    t.add_argument("--kif", type=Path, required=True)
    t.add_argument("--profile", type=Path)
    t.add_argument("--reasons", type=Path, help="Pre-fill the Why column.")
    t.add_argument("--out", type=Path, required=True)

    v2 = sub.add_parser("review", help="Read a person's edits back out of a tracker workbook")
    v2.add_argument("--tracker", type=Path, required=True)
    v2.add_argument("--kif", type=Path, required=True)
    v2.add_argument("--results", type=Path, help="The run's results.json, so edits to computed cells are spotted.")
    v2.add_argument("--reasons", type=Path)
    v2.add_argument("--manual", type=Path)
    v2.add_argument("--out-kif", type=Path, required=True)
    v2.add_argument("--out-reasons", type=Path, required=True)
    v2.add_argument("--out-manual", type=Path)
    v2.add_argument("--by", default="", help="Who made the edits, recorded with any hand-set value.")

    a = ap.parse_args(argv)
    if a.cmd == "build":
        build(a.profile, a.out)
    elif a.cmd == "read":
        read(a.xlsx, a.out)
    elif a.cmd == "review":
        return review(a.tracker, a.kif, a.out_kif, a.out_reasons, a.reasons,
                      a.out_manual, a.manual, a.by, a.results)
    else:
        tracker(a.results, a.kif, a.out, a.reasons, a.profile)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
