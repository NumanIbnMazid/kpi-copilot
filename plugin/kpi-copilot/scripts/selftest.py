#!/usr/bin/env python3
"""
Self-test: run the whole pipeline against the bundled examples and check the answers.

This exists because the expensive mistakes in a KPI tool are not crashes. They are a number
that comes out plausible and wrong - a green "Met" bought by a missing handover date, a
denominator that quietly includes work nobody could judge, a note that says "1 observations".
Those survive code review and die in a management meeting.

So the assertions here are mostly about honesty rather than arithmetic.

    python3 scripts/selftest.py
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml  # type: ignore

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
EX = ROOT / "examples"

passed, failed = 0, 0


def check(name: str, condition: bool, detail: str = "") -> None:
    global passed, failed
    if condition:
        passed += 1
        print(f"  ok    {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}" + (f"\n          {detail}" if detail else ""))


def run(args: list[str], env: dict | None = None) -> subprocess.CompletedProcess:
    import os
    return subprocess.run([sys.executable, *args], capture_output=True, text=True, cwd=ROOT,
                          env={**os.environ, **(env or {})})


def compute(kif: Path, profile: Path, reasons: Path | None, out: Path) -> dict:
    args = ["scripts/kpi_engine.py", "--kif", str(kif), "--profile", str(profile), "--out", str(out)]
    if reasons:
        args += ["--reasons", str(reasons)]
    r = run(args)
    if r.returncode != 0:
        raise SystemExit(f"engine failed: {r.stderr}")
    return json.loads(out.read_text())


def measure(doc: dict, period: str, name: str) -> dict:
    per = next(p for p in doc["periods"] if p["period"] == period)
    return next(m for m in per["measures"] if m["name"] == name)


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="kpi-selftest-"))

    print("\nSchemas and profiles")
    for p in (EX / "northwind-q3" / "profile.yaml", EX / "acme-jira" / "profile.yaml"):
        r = run(["scripts/profile_tool.py", "validate", "--profile", str(p)])
        check(f"{p.parent.name} profile validates", r.returncode == 0, r.stdout)

    print("\nAdapter: csv (a Jira export, no history)")
    r = run(["adapters/csv/extract.py", "--profile", str(EX / "acme-jira" / "profile.yaml"),
             "--project", "acme-identity", "--out", str(tmp / "acme.kif.json")])
    check("csv adapter runs", r.returncode == 0, r.stderr)
    r = run(["scripts/validate_kif.py", "--kif", str(tmp / "acme.kif.json")])
    check("csv output is valid KIF", r.returncode == 0, r.stdout)
    acme_kif = json.loads((tmp / "acme.kif.json").read_text())
    check("csv does not claim status history it cannot see",
          "status_history" not in (acme_kif["generated"]["capabilities"]))
    check("csv splits sprints into periods", len(acme_kif["periods"]) == 2,
          str([p["name"] for p in acme_kif["periods"]]))
    check("exclusion patterns removed the two admin cards",
          sum(1 for t in acme_kif["tasks"] if t["type"] == "Excluded") == 2)
    check("every excluded row says why",
          all(t.get("exclude_reason") for t in acme_kif["tasks"] if t["type"] == "Excluded"))
    check("the change-request label was recognised",
          any(t["type"] == "CR" for t in acme_kif["tasks"]))

    print("\nEngine: honesty")
    acme = compute(tmp / "acme.kif.json", EX / "acme-jira" / "profile.yaml", None, tmp / "acme.results.json")
    m = measure(acme, "Sprint 14", "Rework Rate")
    check("rework is Not measured when no row records a reopen", m["status"] == "Not measured",
          f"got {m['status']} = {m['value']}")
    m = measure(acme, "Sprint 14", "Escaped Defect Rate")
    check("escaped defects are Not measured before handover", m["status"] == "Not measured",
          f"got {m['status']} = {m['value']}")
    for kpi in ("Task Comprehension", "Client Expectation", "Delivery Commitment"):
        m = measure(acme, "Sprint 14", kpi)
        check(f"{kpi} is Not measured with no evidence", m["status"] == "Not measured")

    print("\nEngine: arithmetic")
    sav = compute(EX / "northwind-q3" / "run.kif.json", EX / "northwind-q3" / "profile.yaml",
                  EX / "northwind-q3" / "reasons.yaml", tmp / "sav.results.json")
    m = measure(sav, "Initial Scope", "Task Comprehension")
    check("task comprehension 5/7 = 71.43%", m["value"] == 71.43 and m["numerator"] == 5 and m["denominator"] == 7,
          f"got {m['numerator']}/{m['denominator']} = {m['value']}")
    m = measure(sav, "Initial Scope", "Defect Rate")
    check("defect rate counts 3 of 6 reports on 7 delivered items",
          m["numerator"] == 3 and m["denominator"] == 7, f"got {m['numerator']}/{m['denominator']}")
    m = measure(sav, "Initial Scope", "Rework Rate")
    check("rework counts the reopen and not the first-round QA failure", m["numerator"] == 1,
          f"got numerator {m['numerator']}")
    check("the first-round QA failure is named in the note",
          "first time" in m["note"], m["note"])
    m = measure(sav, "Additional Requests 1", "Client Expectation")
    check("a pending item stays out of the denominator", m["denominator"] == 4, f"got {m['denominator']}")
    check("the pending item is named in the note", "not due until" in m["note"], m["note"])
    m = measure(sav, "Additional Requests 1", "CR Rate")
    check("a CR-only cycle falls back to the project's planned scope",
          "planned items" in m["note"], m["note"])

    print("\nDelivery Commitment means promises kept, not work finished")
    # PMS: "Team-negotiated commitments delivered on time / total team-committed tasks.
    # Insight: reliability of the team." The denominator is commitments, not deliverables.
    mix = json.loads((EX / "northwind-q3" / "run.kif.json").read_text())
    for t in mix["tasks"]:
        if t.get("key") in ("TKT-3479", "TKT-3590"):
            t["met_commitment"] = None
            t.pop("commit_date", None)
    (tmp / "mix.kif.json").write_text(json.dumps(mix))
    mixed = compute(tmp / "mix.kif.json", EX / "northwind-q3" / "profile.yaml", None, tmp / "mix.results.json")
    m = measure(mixed, "Initial Scope", "Delivery Commitment")
    check("an item with no commitment is out of the denominator", m["denominator"] == 5,
          f"got {m['denominator']}, expected 5 of 7 deliverables")
    check("...and the review sheet explains the missing commitment dates",
          any("No commitment date is recorded for 2 other items" in x for x in m["review_items"])
          and "No commitment date is recorded" not in m["note"])
    check("the heading is about commitments, not about QA",
          "recorded commitment dates" in m["note"] and "QA" not in m["note"], m["note"])

    none_committed = json.loads((EX / "northwind-q3" / "run.kif.json").read_text())
    for t in none_committed["tasks"]:
        t["met_commitment"] = None
        t.pop("commit_date", None)
    none_committed["project"]["commit_date"] = None
    for per in none_committed["periods"]:
        per["commit_date"] = None
    (tmp / "nc.kif.json").write_text(json.dumps(none_committed))
    nc = compute(tmp / "nc.kif.json", EX / "northwind-q3" / "profile.yaml", None, tmp / "nc.results.json")
    m = measure(nc, "Initial Scope", "Delivery Commitment")
    check("no commitments at all is Not measured, not 0%", m["status"] == "Not measured",
          f"{m['status']} = {m['value']}")
    check("...and records the reason as a review item",
          any("nothing to measure reliability against" in x for x in m["review_items"]) and not m["note"])

    allcfg = yaml.safe_load((EX / "northwind-q3" / "profile.yaml").read_text())
    allcfg["workflow"]["commitment"] = {"scope": "all-deliverables", "met_when": "handover"}
    (tmp / "allcfg.yaml").write_text(yaml.safe_dump(allcfg))
    alld = compute(tmp / "mix.kif.json", tmp / "allcfg.yaml", None, tmp / "all.results.json")
    m = measure(alld, "Initial Scope", "Delivery Commitment")
    check("a team that commits to the whole scope can say so",
          m["denominator"] == 5 and "made no commitment" not in m["note"], m["note"])
    check("what counts as meeting a commitment is configurable, and said in words",
          "counted on the handover to the client" in m["note"], m["note"])

    print("\nEscaped Defect Rate is defects over defects, not over items")
    # PMS: "Defects found after release / Total defects (before + after release)."
    m = measure(sav, "Initial Scope", "Escaped Defect Rate")
    check("the denominator is valid defects, not delivered items",
          m["denominator"] == 5, f"got {m['denominator']}, expected 5 non-rejected reports of 6")
    check("rejected reports are excluded from both counts", "rejected report is excluded from both counts" in m["note"], m["note"])
    m2 = measure(sav, "Additional Requests 1", "Escaped Defect Rate")
    check("no handover still means Not measured", m2["status"] == "Not measured", m2["status"])

    print("\nThe workbook matches the live tracker's shape")
    r = run(["scripts/workbook.py", "tracker", "--results", str(tmp / "sav.results.json"),
             "--kif", str(EX / "northwind-q3" / "run.kif.json"),
             "--reasons", str(EX / "northwind-q3" / "reasons.yaml"),
             "--profile", str(EX / "northwind-q3" / "profile.yaml"),
             "--out", str(tmp / "full.xlsx")])
    check("the tracker workbook builds", r.returncode == 0, r.stderr)
    from openpyxl import load_workbook as _lwb
    fwb = _lwb(tmp / "full.xlsx")
    check("it has the eight tabs the live trackers have, plus Gaps",
          fwb.sheetnames == ["Read Me", "Dashboard", "Config", "Periods", "Task Register",
                             "Defect Register", "KPI Summary", "PMS Push Log", "Gaps"],
          str(fwb.sheetnames))
    sys.path.insert(0, str(ROOT / "scripts"))
    from workbook import HDR as WB_HDR  # noqa: E402
    tr = fwb["Task Register"]
    hdrs = {c.value for c in tr[WB_HDR] if c.value}
    check("the register carries the gate columns the KPIs depend on",
          {"Client-Expected?", "Team Committed?", "Met Commitment?"} <= hdrs,
          str(sorted(hdrs))[:200])
    check("...and the evidence beside each judgement",
          {"Comprehension Note", "Comprehension Evidence", "Client Date Evidence",
           "Commitment Evidence", "Rework Evidence"} <= hdrs, str(sorted(hdrs))[:220])
    check("columns are sized, not left at default",
          (tr.column_dimensions["E"].width or 0) > 20, str(tr.column_dimensions["E"].width))
    check("headers freeze so a wide register stays readable",
          tr.freeze_panes == f"E{WB_HDR + 1}", str(tr.freeze_panes))
    check("Config explains each KPI with the PMS formula",
          any("Formula (PMS definition)" == c.value
              for row in fwb["Config"].iter_rows() for c in row), "not found")
    dash = fwb["Dashboard"]
    check("the dashboard totals work and issues per period",
          any(c.value == "Found after handover" for row in dash.iter_rows() for c in row))

    print("\nEngine: the English")
    notes = [m["note"] for p in sav["periods"] for m in p["measures"]] + \
            [m["note"] for p in acme["periods"] for m in p["measures"]]
    joined = " ".join(notes)
    for bad in ("1 observations", "1 improvements", "1 items", "1 bugs", "1 requirements were",
                "of 1 items", "all 1 items", "the 1 items"):
        check(f"no '{bad}'", bad not in joined,
              next((n for n in notes if bad in n), ""))
    check("no links reach a PMS note", "http" not in joined)
    for bad in ("1 reports were", "1 tasks", "1 commitments", "1 issues", "The other 1 ", "All 1 ", "1 requests on",
                "1 items"):
        check(f"no '{bad}'", bad not in joined, next((n for n in notes if bad in n), ""))
    check("every part of every note is a whole sentence",
          all(part.strip().endswith(".") for n in notes if n for part in n.split(" || ")),
          next((part for n in notes for part in n.split(" || ") if not part.strip().endswith(".")), ""))
    check("no note opens with a heading: the first part says the number", all(any(ch.isdigit() for ch in n.split(" || ")[0])
          or "othing" in n or "No " in n or "None " in n or "handover has not been recorded" in n or "not been handed over" in n or "Everything" in n for n in notes if n),
          next((n for n in notes if n and not any(ch.isdigit() for ch in n.split(" || ")[0])), ""))
    import yaml as _y
    fp = _y.safe_load((EX / "northwind-q3" / "profile.yaml").read_text())
    fp.setdefault("organization", {})["note_style"] = "fragments"
    (tmp / "fragments.yaml").write_text(_y.safe_dump(fp, sort_keys=False))
    frag = compute(EX / "northwind-q3" / "run.kif.json", tmp / "fragments.yaml", None, tmp / "fragments.json")
    a = {(p["period"], m["name"]): m["value"] for p in sav["periods"] for m in p["measures"]}
    b = {(p["period"], m["name"]): m["value"] for p in frag["periods"] for m in p["measures"]}
    check("the older fragment style is still there for anyone who prefers it",
          measure(frag, "Additional Requests 1", "Velocity")["note"].startswith("Work finished in this cycle ||"))
    check("...and the style changes the wording only, never a figure", a == b)

    print("\nCustom instructions")
    ov = {o["rule"]: o for o in sav.get("custom_overrides", [])}
    check("a supported override is applied", ov.get("defect_phase", {}).get("status") == "applied",
          json.dumps(ov.get("defect_phase", {})))
    m = measure(sav, "Initial Scope", "Escaped Defect Rate")
    check("the override moved the number", m["numerator"] == 1, f"got {m['numerator']}")

    illegal = json.loads(json.dumps(json.loads((EX / "northwind-q3" / "run.kif.json").read_text())))
    prof = yaml.safe_load((EX / "northwind-q3" / "profile.yaml").read_text())
    prof.setdefault("custom_instructions", {}).setdefault("rule_overrides", []).append(
        {"rule": "threshold_defect_rate", "value": "25", "why": "legacy surface"})
    (tmp / "illegal.yaml").write_text(yaml.safe_dump(prof))
    (tmp / "illegal.kif.json").write_text(json.dumps(illegal))
    doc = compute(tmp / "illegal.kif.json", tmp / "illegal.yaml", None, tmp / "illegal.results.json")
    bad = next((o for o in doc["custom_overrides"] if o["rule"] == "threshold_defect_rate"), {})
    check("a local threshold is refused, and points at PMS", "in PMS" in bad.get("status", ""),
          json.dumps(bad))
    m = measure(doc, "Initial Scope", "Defect Rate")
    check("the refused override did not change the target", m["threshold"] == 15, f"got {m['threshold']}")
    r = run(["scripts/profile_tool.py", "validate", "--profile", str(tmp / "illegal.yaml")])
    check("validation also refuses it, and says where to set it",
          r.returncode != 0 and "on this project in PMS" in r.stdout, r.stdout)

    print("\nPer-project targets from PMS")
    reg = json.loads((ROOT / "schemas" / "kpi_registry.default.json").read_text())
    reg["source"], reg["synced_at"] = "pms", "2026-09-18T15:00:00+00:00"
    reg["projects"] = {"101": {"defect_rate": {"threshold": 25, "direction": "lower-is-better"},
                               "rework_rate": {"threshold": 20, "direction": "lower-is-better"}}}
    (tmp / "reg_project.json").write_text(json.dumps(reg, indent=2))
    r = run(["scripts/kpi_engine.py", "--kif", str(EX / "northwind-q3" / "run.kif.json"),
             "--profile", str(EX / "northwind-q3" / "profile.yaml"),
             "--registry", str(tmp / "reg_project.json"), "--out", str(tmp / "proj.results.json")])
    check("engine runs against a per-project registry", r.returncode == 0, r.stderr)
    proj = json.loads((tmp / "proj.results.json").read_text())
    m = measure(proj, "Initial Scope", "Defect Rate")
    check("a project's own target from PMS is used", m["threshold"] == 25, f"got {m['threshold']}")
    check("and it says the target came from that project",
          "project 101" in m.get("threshold_source", ""), m.get("threshold_source"))
    m = measure(proj, "Initial Scope", "Rework Rate")
    check("a looser project target can flip a status to Met",
          m["threshold"] == 20 and m["status"] == "Met", f"{m['threshold']} / {m['status']}")
    m = measure(proj, "Initial Scope", "Task Comprehension")
    check("a KPI the project does not override keeps the PMS default",
          m["threshold"] == 80 and "default" in m.get("threshold_source", ""),
          f"{m['threshold']} / {m.get('threshold_source')}")
    dflt = measure(sav, "Initial Scope", "Defect Rate")
    check("the bundled fallback says it is not from PMS",
          "fallback" in dflt.get("threshold_source", ""), dflt.get("threshold_source"))

    print("\nSeveral accounts, several trackers, one profile")
    sys.path.insert(0, str(ROOT / "scripts"))
    import profile_lib as pl  # noqa: E402

    multi = pl.load(EX / "multi-account" / "profile.yaml")
    r = run(["scripts/profile_tool.py", "validate", "--profile", str(EX / "multi-account" / "profile.yaml")])
    check("a multi-account profile validates", r.returncode == 0, r.stdout)

    sav_p, _ = pl.resolve(multi, "q3-release")
    acme, _ = pl.resolve(multi, "acme-identity")
    check("two projects on one profile resolve to different trackers",
          sav_p["tracker"]["adapter"] == "asana" and acme["tracker"]["adapter"] == "csv",
          f"{sav_p['tracker']['adapter']} / {acme['tracker']['adapter']}")
    check("...different chat tools",
          sav_p["sources"]["evidence_channels"] == ["chat-northwind-devqa", "chat-northwind-mgmt"]
          and acme["sources"]["evidence_channels"] == ["slack-acme"],
          f"{sav_p['sources']['evidence_channels']} / {acme['sources']['evidence_channels']}")
    check("...different output modes",
          sav_p["output"]["mode"] == "assisted-push" and acme["output"]["mode"] == "review-only",
          f"{sav_p['output']['mode']} / {acme['output']['mode']}")
    check("...different period models",
          sav_p["periods"]["model"] == "Delivery cycle" and acme["periods"]["model"] == "Sprint")
    check("...different defect conventions",
          sav_p["conventions"]["defect_by"] == "title-pattern" and acme["conventions"]["defect_by"] == "issue-type")

    gateway, _ = pl.resolve(multi, "gateway")
    check("a project inherits its account without repeating it",
          gateway["tracker"]["adapter"] == "asana" and gateway["conventions"]["key_pattern"] == r"TKT-\d+")
    check("a project override sits on top of its account",
          (gateway["sources"].get("timeline") or {}).get("tool_id") == "sheet-gateway-timeline")
    check("its sibling does not get that override",
          (sav_p["sources"].get("timeline") or {}).get("tool_id") is None)

    check("tools are scoped, so Acme cannot see Northwind's chat",
          "chat-northwind-devqa" not in {t["id"] for t in acme["tools"]}
          and "slack-acme" in {t["id"] for t in acme["tools"]},
          str(sorted(t["id"] for t in acme["tools"])))
    check("shared tools are visible to everyone",
          "pms" in {t["id"] for t in sav_p["tools"]} and "pms" in {t["id"] for t in acme["tools"]})

    deep = pl.deep_merge({"a": {"x": 1, "y": 2}, "list": [1, 2]}, {"a": {"y": 9}, "list": [3]})
    check("dicts merge deeply, lists replace",
          deep == {"a": {"x": 1, "y": 9}, "list": [3]}, str(deep))

    r = run(["adapters/csv/extract.py", "--profile", str(EX / "multi-account" / "profile.yaml"),
             "--project", "acme-identity", "--out", str(tmp / "multi.kif.json")])
    check("an adapter runs against the resolved project", r.returncode == 0, r.stderr)
    r = run(["scripts/preflight.py", "--profile", str(EX / "multi-account" / "profile.yaml"),
             "--state", str(tmp / "pf-multi.json"), "--json"])
    ids = {c["id"] for c in json.loads(r.stdout)["checks"]}
    check("preflight checks every project's adapter",
          {"adapter-present-q3-release", "adapter-present-acme-identity"} <= ids,
          str(sorted(i for i in ids if i.startswith("adapter"))))

    bad = pl.load(EX / "multi-account" / "profile.yaml")
    bad["projects"][2]["overrides"] = {"sources": {"evidence_channels": ["chat-northwind-devqa"]}}
    (tmp / "crosstalk.yaml").write_text(yaml.safe_dump(bad))
    r = run(["scripts/profile_tool.py", "validate", "--profile", str(tmp / "crosstalk.yaml")])
    check("a project reaching into another account's tools is refused",
          r.returncode != 0 and "not a tool this project can see" in r.stdout, r.stdout)

    r = run(["scripts/workbook.py", "build", "--profile", str(EX / "multi-account" / "profile.yaml"),
             "--out", str(tmp / "multi.xlsx")])
    check("a multi-account workbook builds", r.returncode == 0, r.stderr)
    r = run(["scripts/workbook.py", "read", "--xlsx", str(tmp / "multi.xlsx"), "--out", str(tmp / "multi.back.yaml")])
    back = pl.load(tmp / "multi.back.yaml")
    check("overrides survive the workbook round trip",
          back["accounts"][0]["overrides"]["conventions"]["client_names"]
          == multi["accounts"][0]["overrides"]["conventions"]["client_names"],
          str(back["accounts"][0]["overrides"]["conventions"].get("client_names")))
    check("a sentence containing a comma is not turned into a list",
          isinstance(back["accounts"][0]["overrides"]["workflow"]["closed_when"]["completed_flag_means"], str),
          repr(back["accounts"][0]["overrides"]["workflow"]["closed_when"].get("completed_flag_means")))

    print("\nCapturing settings from conversation")
    import shutil
    shutil.copy(EX / "northwind-q3" / "profile.yaml", tmp / "rm.yaml")
    r = run(["scripts/remember.py", "--profile", str(tmp / "rm.yaml"),
             "--add", "custom_instructions.never=Never push during a release freeze",
             "--why", "asked in chat", "--by", "Tester"])
    check("a setting said in chat can be written down", r.returncode == 0, r.stderr)
    got = pl.load(tmp / "rm.yaml")
    check("...and it is actually in the profile",
          "Never push during a release freeze" in got["custom_instructions"]["never"])
    check("...with the sentence it came from",
          any(e.get("why") == "asked in chat" for e in got.get("learned") or []),
          json.dumps(got.get("learned"))[:200])

    r = run(["scripts/remember.py", "--profile", str(tmp / "rm.yaml"),
             "--set", "output.mode=review-only"])
    check("a change with no reason is refused", r.returncode != 0 and "--why is required" in r.stderr)

    r = run(["scripts/remember.py", "--profile", str(tmp / "rm.yaml"),
             "--set", "output.mode=whenever", "--why", "testing"])
    check("a change that breaks the profile is rolled back", r.returncode != 0, r.stdout)
    still = pl.load(tmp / "rm.yaml")
    check("...leaving the profile loadable and unchanged",
          still["output"]["mode"] == "assisted-push", still["output"]["mode"])

    r = run(["scripts/remember.py", "--profile", str(tmp / "rm.yaml"),
             "--scope", "account:nope", "--set", "output.mode=review-only", "--why", "x"])
    check("an unknown account is refused by name", r.returncode != 0 and "Known" in r.stderr)

    r = run(["scripts/remember.py", "--profile", str(tmp / "rm.yaml"),
             "--tool", "id=slack-ops,kind=chat,name=#ops,used_for=evidence|handover",
             "--why", "named in chat"])
    check("a tool named in chat can be added", r.returncode == 0, r.stderr)
    got = pl.load(tmp / "rm.yaml")
    t = next((x for x in got["tools"] if x["id"] == "slack-ops"), None)
    check("...with its list fields parsed", t and t["used_for"] == ["evidence", "handover"], str(t))

    print("\nPush guards")
    run(["scripts/kpi_engine.py", "--kif", str(tmp / "acme.kif.json"),
         "--profile", str(EX / "acme-jira" / "profile.yaml"), "--payloads", str(tmp / "acme.payloads.json")])
    r = run(["scripts/pms_push.py", "--payloads", str(tmp / "acme.payloads.json"),
             "--profile", str(EX / "acme-jira" / "profile.yaml"), "--apply"])
    check("review-only refuses --apply", r.returncode == 2 and "Refusing to write" in r.stderr, r.stderr)
    payloads = json.loads((tmp / "acme.payloads.json").read_text())
    sent = [k["name"] for p in payloads for k in p["kpis"]]
    check("unmeasured KPIs are not sent to PMS", "Rework Rate" not in sent, str(sent))
    check("unmeasured KPIs are listed as skipped, with a reason",
          all(s.get("reason") for p in payloads for s in p["skipped"]))
    check("...and carry a plain note for PMS saying why there is no value",
          all(s.get("note") and "||" not in s["note"] and "http" not in s["note"]
              for p in payloads for s in p["skipped"] if not s.get("preserve_existing")),
          str([s.get("note") for p in payloads for s in p["skipped"]]))

    print("\nWorkbook")
    r = run(["scripts/workbook.py", "build", "--profile", str(EX / "northwind-q3" / "profile.yaml"),
             "--out", str(tmp / "profile.xlsx")])
    check("profile workbook builds", r.returncode == 0, r.stderr)
    r = run(["scripts/workbook.py", "read", "--xlsx", str(tmp / "profile.xlsx"), "--out", str(tmp / "back.yaml")])
    check("profile workbook reads back", r.returncode == 0, r.stderr)
    a = yaml.safe_load((EX / "northwind-q3" / "profile.yaml").read_text())
    b = yaml.safe_load((tmp / "back.yaml").read_text())

    def lost(orig, back, path=""):
        """Nothing the profile SAID may change or disappear. The workbook also fills in schema
        defaults the source left implicit, which is the point of showing them to a human, so
        keys that were absent and came back are not losses."""
        out = []
        if isinstance(orig, dict):
            if not isinstance(back, dict):
                return [(path, orig, back)]
            for k, v in orig.items():
                out += lost(v, back.get(k), f"{path}.{k}")
        elif orig != back:
            out.append((path, orig, back))
        return out

    missing = lost(a, b)
    check("workbook round trip loses nothing that was written",
          not missing, "; ".join(f"{p}: {o!r} -> {n!r}" for p, o, n in missing[:4]))
    local_profile = json.loads(json.dumps(a))
    local_profile["organization"]["pms_base_url"] = ""
    local_profile["output"]["mode"] = "review-only"
    (tmp / "local-profile.yaml").write_text(yaml.safe_dump(local_profile))
    built = run(["scripts/workbook.py", "build", "--profile", str(tmp / "local-profile.yaml"),
                 "--out", str(tmp / "local-profile.xlsx")])
    imported = run(["scripts/workbook.py", "read", "--xlsx", str(tmp / "local-profile.xlsx"),
                    "--out", str(tmp / "local-back.yaml")])
    checked = run(["scripts/profile_tool.py", "validate", "--profile", str(tmp / "local-back.yaml")])
    restored = yaml.safe_load((tmp / "local-back.yaml").read_text())
    check("a local-only profile stays valid after workbook import with a blank PMS address",
          built.returncode == imported.returncode == checked.returncode == 0
          and restored["organization"].get("pms_base_url") == ""
          and not lost(local_profile, restored), checked.stdout + checked.stderr)
    r = run(["scripts/workbook.py", "tracker", "--results", str(tmp / "sav.results.json"),
             "--kif", str(EX / "northwind-q3" / "run.kif.json"), "--out", str(tmp / "tracker.xlsx")])
    check("tracker workbook builds", r.returncode == 0, r.stderr)

    print("\nEditing in the sheet, and what reaches PMS")
    from openpyxl import load_workbook as _lw  # noqa: E402

    # Ask the builder where its rows are rather than hard-coding them here, so a change of
    # layout shows up as a failing assertion rather than a KeyError in the test itself.
    sys.path.insert(0, str(ROOT / "scripts"))
    from workbook import FIRST as WB_FIRST, KEYROW as WB_KEYROW  # noqa: E402
    r = run(["scripts/workbook.py", "tracker", "--results", str(tmp / "sav.results.json"),
             "--kif", str(EX / "northwind-q3" / "run.kif.json"),
             "--reasons", str(EX / "northwind-q3" / "reasons.yaml"), "--out", str(tmp / "t.xlsx")])
    check("tracker workbook builds with a Why column", r.returncode == 0, r.stderr)

    wbk = _lw(tmp / "t.xlsx")
    ws = wbk["KPI Summary"]
    keys = {c.value: c.column for c in ws[WB_KEYROW] if c.value}
    check("the sheet offers a place to set a value by hand",
          {"why", "set_value_by_hand", "why_set_by_hand"} <= set(keys), str(sorted(keys)))
    check("the Why column is pre-filled from reasons.yaml",
          any(r_[keys["why"] - 1].value for r_ in ws.iter_rows(min_row=WB_FIRST)))

    # Edit: a judgement, an hour figure, a defect, a Why, and a value set by hand.
    tr = wbk["Task Register"]
    tk = {c.value: c.column for c in tr[WB_KEYROW] if c.value}
    for row in tr.iter_rows(min_row=WB_FIRST):
        if row[tk["key"] - 1].value == "TKT-3175":
            row[tk["understood"] - 1].value = "Yes"
            row[tk["hours_dev"] - 1].value = 30
    dr = wbk["Defect Register"]
    dk = {c.value: c.column for c in dr[WB_KEYROW] if c.value}
    for row in dr.iter_rows(min_row=WB_FIRST):
        if row[dk["key"] - 1].value == "Bug 7":
            row[dk["rejected"] - 1].value = "Yes"
            row[dk["rejection_reason"] - 1].value = "by design"
    for row in ws.iter_rows(min_row=WB_FIRST):
        if row[keys["period"] - 1].value == "Initial Scope":
            if row[keys["kpi"] - 1].value == "CR Rate":
                row[keys["set_value_by_hand"] - 1].value = 33
                row[keys["why_set_by_hand"] - 1].value = "Two items were re-scoped as plan work."
            if row[keys["kpi"] - 1].value == "Escaped Defect Rate":
                row[keys["set_value_by_hand"] - 1].value = 0      # no reason
            if row[keys["kpi"] - 1].value == "Velocity":
                row[keys["value"] - 1].value = 999                 # a computed cell
    wbk.save(tmp / "t.xlsx")

    r = run(["scripts/workbook.py", "review", "--tracker", str(tmp / "t.xlsx"),
             "--kif", str(EX / "northwind-q3" / "run.kif.json"), "--results", str(tmp / "sav.results.json"),
             "--reasons", str(EX / "northwind-q3" / "reasons.yaml"),
             "--out-kif", str(tmp / "patched.kif.json"), "--out-reasons", str(tmp / "r2.yaml"),
             "--out-manual", str(tmp / "manual.yaml"), "--by", "Tester"])
    check("review reads the edits back", r.returncode == 0, r.stderr)
    out = r.stdout
    check("a judgement edit is picked up", "understood" in out and "No -> Yes" in out, out)
    check("an hours edit is picked up", "hours_dev" in out, out)
    check("a defect edit is picked up", "rejected" in out, out)
    check("a hand-set value with a reason is kept", "set by hand" in out, out)
    check("a hand-set value with no reason is refused, and says what is missing",
          "needs a reason" in out, out)
    check("typing over a computed cell is reported, not silently dropped",
          "'Value' is computed" in out, out)

    patched = json.loads((tmp / "patched.kif.json").read_text())
    t = next(x for x in patched["tasks"] if x["key"] == "TKT-3175")
    check("the edit really landed in the extract", t["understood"] == "Yes" and t["hours_dev"] == 30.0,
          f"{t['understood']} / {t['hours_dev']}")

    r = run(["scripts/kpi_engine.py", "--kif", str(tmp / "patched.kif.json"),
             "--profile", str(EX / "northwind-q3" / "profile.yaml"),
             "--reasons", str(tmp / "r2.yaml"), "--manual", str(tmp / "manual.yaml"),
             "--out", str(tmp / "results2.json"), "--payloads", str(tmp / "pay2.json")])
    check("the engine reruns on the edited extract", r.returncode == 0, r.stderr)
    after = json.loads((tmp / "results2.json").read_text())
    before_tc = measure(sav, "Initial Scope", "Task Comprehension")["value"]
    after_tc = measure(after, "Initial Scope", "Task Comprehension")["value"]
    check("a judgement edit moves the KPI", after_tc > before_tc, f"{before_tc} -> {after_tc}")
    before_dr = measure(sav, "Initial Scope", "Defect Rate")["value"]
    after_dr = measure(after, "Initial Scope", "Defect Rate")["value"]
    check("rejecting a defect lowers the defect rate", after_dr < before_dr, f"{before_dr} -> {after_dr}")

    cr = measure(after, "Initial Scope", "CR Rate")
    check("a hand-set value is used", cr["value"] == 33.0, str(cr["value"]))
    check("...and the computed one is kept beside it", cr["computed_value"] == 40.0, str(cr.get("computed_value")))
    check("...and its note says a person set it, and why",
          "Recorded as 33%" in cr["note"] and "re-scoped" in cr["note"], cr["note"])
    edr = measure(after, "Initial Scope", "Escaped Defect Rate")
    check("a hand-set value with no reason never reaches the numbers", not edr["overridden"])

    pay = json.loads((tmp / "pay2.json").read_text())[0]
    sent = {k["name"]: k for k in pay["kpis"]}
    check("PMS gets the hand-set value, flagged, with the computed one alongside",
          sent["CR Rate"]["value"] == 33.0 and sent["CR Rate"]["overridden"]
          and sent["CR Rate"]["computed_value"] == 40.0, json.dumps(sent["CR Rate"])[:200])

    print("\nAdapter: asana bridge")
    legacy = {
        "projName": "Bridge test",
        "cfg": {"projectGid": "1", "pmsProjectId": 361},
        "projectDates": {"clientDate": "2026-09-11", "commitDate": "2026-09-11", "clientCheck": "Delivery"},
        "periodRows": [{"name": "Initial Scope", "pmsId": 723, "start": "2026-07-01", "end": "2026-08-12",
                        "releaseDate": "2026-08-12", "teamHours": 10}],
        "taskRows": [{"period": "Initial Scope", "key": "T-1", "title": "x", "type": "Task", "initial": "Yes",
                      "est": 10, "dev": 10, "qaC": 2, "created": "2026-07-01", "devdone": "2026-07-10",
                      "closed": "2026-07-20", "status": "Done", "und": "Yes", "cexp": "Yes", "commit": "Yes",
                      "reo": "No"}],
        "defectRows": [{"period": "Initial Scope", "key": "Bug 1", "title": "y", "kind": "Bug", "by": "QA",
                        "on": "2026-07-12", "phase": "Pre-release", "existing": "No", "rej": "No",
                        "fs": "Fixed / Closed"}],
        "review": [],
    }
    (tmp / "legacy.json").write_text(json.dumps(legacy))
    r = run(["adapters/asana/extract.py", "--profile", str(EX / "northwind-q3" / "profile.yaml"),
             "--project", "q3-release", "--from-extract", str(tmp / "legacy.json"),
             "--out", str(tmp / "bridge.kif.json")])
    check("asana bridge runs", r.returncode == 0, r.stderr)
    r = run(["scripts/validate_kif.py", "--kif", str(tmp / "bridge.kif.json")])
    check("asana bridge output is valid KIF", r.returncode == 0, r.stdout)
    bridged = json.loads((tmp / "bridge.kif.json").read_text())
    check("'Pre-release' becomes 'QA'", bridged["defects"][0]["phase"] == "QA")
    check("the handover date survives", bridged["periods"][0]["handover_date"] == "2026-08-12")

    # ---- Source of truth and scan bounds --------------------------------------------
    # These decide how long a run takes. The risk is not a crash: it is that a profile
    # saying "board only" quietly reads chat anyway, or that a missing plan is reported as
    # a failure to a team that deliberately does not have one.
    print("\nSource of truth and scan bounds")
    to = EX / "tracker-only" / "profile.yaml"
    check("the tracker-only example exists", to.exists())

    r = run(["scripts/profile_tool.py", "validate", "--profile", str(to)])
    check("tracker-only profile validates", r.returncode == 0, r.stderr + r.stdout)

    prof = yaml.safe_load(to.read_text())
    check("sources.mode is accepted by the schema", prof["sources"]["mode"] == "tracker-only")
    check("scan bounds the window", prof["scan"]["window"] == "period+grace")
    check("scan can bound the board", prof["scan"]["tracker_scope"] == "touched-since")

    r = run(["scripts/preflight.py", "--profile", str(to), "--project", "atlas", "--state", str(tmp / "pf-tracker.json")])
    out = r.stdout
    check("a missing plan is 'not needed', not a gap, on a tracker-only profile",
          "not needed - this profile treats the tracker as the single source of truth" in out, out[:400])
    check("no chat or mail listed is a complete configuration, not a gap",
          "none - open questions go to a person" in out and "At least one evidence channel" not in out, out[:400])
    check("readiness reports the reach of a run", "A run has a bounded reach" in out)

    unbounded = tmp / "unbounded.yaml"
    prof_u = yaml.safe_load(to.read_text())
    prof_u["scan"]["window"] = "all"
    unbounded.write_text(yaml.safe_dump(prof_u, sort_keys=False))
    r = run(["scripts/preflight.py", "--profile", str(unbounded), "--project", "atlas"])
    check("an unbounded run is flagged rather than accepted silently",
          "unbounded - every run reads all history" in r.stdout, r.stdout[:400])

    # scan must be settable per client, or a lead with two clients cannot have one cheap and
    # one thorough.
    from importlib import import_module
    sys.path.insert(0, str(HERE))
    profile_lib = import_module("profile_lib")
    check("scan can be overridden per account or project", "scan" in profile_lib.OVERRIDABLE)

    layered = tmp / "layered.yaml"
    prof_l = yaml.safe_load(to.read_text())
    prof_l["accounts"] = [{"id": "fast", "name": "Fast",
                           "overrides": {"scan": {"comments": "never"},
                                         "sources": {"mode": "multi-source"}}}]
    prof_l["projects"][0]["account"] = "fast"
    layered.write_text(yaml.safe_dump(prof_l, sort_keys=False))
    merged, _ = profile_lib.resolve(profile_lib.load(layered), "atlas")
    check("an account override reaches scan", merged["scan"]["comments"] == "never")
    check("...without losing the rest of the block", merged["scan"]["window"] == "period+grace")
    check("an account can change the source of truth", merged["sources"]["mode"] == "multi-source")


    # ----------------------------------------------------------------------------------
    print("\nOne command, and the sheet it produces")
    # A run used to be fourteen commands. The arithmetic was never the slow part - the round
    # trips were - so the whole deterministic chain has to work in one call, and has to end
    # by naming what is missing rather than sending somebody off to search for everything.
    one = tmp / "onepass"
    r = run(["scripts/run.py", "--profile", str(EX / "acme-jira" / "profile.yaml"),
             "--project", "acme-identity", "--out-dir", str(one), "--skip-preflight",
             "--reasons", str(tmp / "one.reasons.yaml"), "--manual", str(tmp / "one.manual.yaml")])
    check("one command runs the whole pipeline", r.returncode == 0, r.stderr[-400:])
    for artefact in ("run.kif.json", "results.json", "report.md", "payloads.json", "tracker.xlsx"):
        check(f"...and leaves {artefact}", (one / artefact).exists())
    out = r.stdout
    check("it says what the tracker could not answer", "could not answer" in out, out[-300:])
    check("...naming the items, not just the KPI", "ACME-101" in out, out[-600:])
    check("...and what would answer it", "Usually in" in out, out[-300:])

    # An item nobody committed to is not a missing fact - Delivery Commitment measures
    # promises kept, so there is nothing to look up. Saying otherwise sends somebody
    # searching chat for an hour to find a blank that was already correct.
    check("an item with no commitment is not reported as a gap",
          "no 'met_commitment'" not in out, out[-600:])
    for bad in ("1 reports", "1 tasks", "1 periods", "1 report have", "1 task have"):
        check(f"no '{bad}' in the shopping list", bad not in out)
    check("it reports its own timings", "= " in out and "s\n" in out)

    sys.path.insert(0, str(HERE))
    from workbook import FIRST as F, HDR as H, KEYROW as K  # noqa: E402
    from openpyxl import load_workbook as _lw3  # noqa: E402
    twb = _lw3(one / "tracker.xlsx")
    tr2 = twb["Task Register"]
    keys2 = {c.value: c.column for c in tr2[K] if c.value}
    import datetime as _d
    dates = [tr2.cell(row=rr, column=keys2["created"]).value
             for rr in range(F, F + 6)]
    check("dates are written as dates, so the column sorts as one",
          any(isinstance(v, (_d.date, _d.datetime)) for v in dates), str(dates[:3]))
    check("...with a date format on the cell",
          tr2.cell(row=F, column=keys2["created"]).number_format == "yyyy-mm-dd",
          tr2.cell(row=F, column=keys2["created"]).number_format)
    check("rows are a normal height, not three lines deep",
          (tr2.row_dimensions[F].height or 99) <= 20, str(tr2.row_dimensions[F].height))
    check("the base font is set, so nothing falls back to Calibri",
          tr2.cell(row=F, column=keys2["title"]).font.name == "Arial")

    # Colour has to mean something, and it has to mean the same thing the tool will accept.
    from workbook import COLS_TASKS, EDITABLE as WB_EDIT, IN as WB_IN  # noqa: E402
    yellow = {k for k, _h, _w, role, _dt in COLS_TASKS if role == WB_IN}
    views = {"client_expected", "team_committed"}
    check("every yellow column is one review will actually read back",
          yellow - views <= WB_EDIT["Task Register"],
          str(sorted(yellow - views - WB_EDIT["Task Register"])))
    fills = {k: tr2.cell(row=F, column=keys2[k]).fill.fgColor.rgb
             for k in ("title", "understood", "hours_total")}
    check("yellow means yours, white means read, grey means computed",
          fills["understood"].endswith("FFF2CC") and fills["title"].endswith("FFFFFF")
          and fills["hours_total"].endswith("F2F2F2"), str(fills))
    check("formatting carries on past the last row",
          tr2.cell(row=tr2.max_row, column=keys2["title"]).border.left.style == "thin")
    check("the columns with a fixed set of answers get a dropdown",
          len(tr2.data_validations.dataValidation) >= 5,
          str(len(tr2.data_validations.dataValidation)))

    # The dropdowns come off the KIF schema, so the sheet cannot offer a value the contract
    # would reject - or miss one the contract has just gained.
    from workbook import _choices  # noqa: E402
    ch = _choices("tasks")
    kif_schema = json.loads((ROOT / "schemas" / "kif.schema.json").read_text())
    enum = [v for v in kif_schema["properties"]["tasks"]["items"]["properties"]
            ["met_commitment"]["enum"] if v]
    check("dropdown values come from the schema, not a copy of it",
          ch["met_commitment"] == ",".join(enum), ch.get("met_commitment", ""))

    # The round trip is the whole promise of an editable sheet.
    tr2.cell(row=F, column=keys2["understood"]).value = "Yes"
    tr2.cell(row=F, column=keys2["reopened"]).value = "No"
    edited_key = tr2.cell(row=F, column=keys2["key"]).value
    twb.save(one / "tracker.xlsx")
    r = run(["scripts/run.py", "review", "--profile", str(EX / "acme-jira" / "profile.yaml"),
             "--project", "acme-identity", "--out-dir", str(one), "--by", "Tester",
             "--reasons", str(tmp / "one.reasons.yaml"), "--manual", str(tmp / "one.manual.yaml")])
    check("one command folds the sheet back in and recomputes", r.returncode == 0, r.stderr[-400:])
    check("...and says what it read back", "understood" in r.stdout, r.stdout[-300:])
    patched2 = json.loads((one / "run.kif.json").read_text())
    tsk = next(x for x in patched2["tasks"] if x["key"] == edited_key)
    check("...and the edit really landed in the extract", tsk["understood"] == "Yes")
    check("a link cell round trips as its target, not as the word on it",
          all("Open" not in str(x.get("link") or "") for x in patched2["tasks"]))

    pipeline_tests(tmp)
    question_tests(tmp)

    audit = subprocess.run([sys.executable, "-m", "unittest", "discover", "-s", str(ROOT / "tests"), "-v"],
                           capture_output=True, text=True)
    check("audit failure-path regressions", audit.returncode == 0, audit.stdout + audit.stderr)

    print(f"\n{passed} passed, {failed} failed\n")
    return 1 if failed else 0


def pipeline_tests(tmp: Path) -> None:
    """The run as people actually use it: read a board, judge it, keep the judgement, build
    the sheet, read the sheet back. Everything here runs offline against examples/northwind-board."""
    import shutil
    sys.path.insert(0, str(ROOT / "adapters" / "asana"))
    import board as B
    import sheet_google
    import sheet_model
    from ledger import Ledger

    print("\nTags people mistype")
    word, how, raw = B.find_tag("[Tablet][Exisiting] Bug 04: labels cropped", ["Existing", "Pre-existing"])
    check("'[Exisiting]' is read as Existing", word == "Existing" and raw == "Exisiting", f"{word} {raw}")
    check("...and the match says it was tolerant, so the sheet can say so too", how == "fuzzy", how)
    check("an exact tag is exact", B.find_tag("[Web][Existing] Bug 3", ["Existing"])[1] == "exact")
    check("a short tag gets no tolerance: [CRs] is not silently CR", B.find_tag("[CRs] Thing", ["CR"])[0] is None)
    check("a component tag is not mistaken for anything", B.find_tag("[Tablet] Bug 2", ["Existing"])[0] is None)

    print("\nAsana, read directly")
    import api as asana
    raw = {"project": {"gid": "77", "name": "Demo"}, "sections": [{"name": "To Do"}, {"name": "Closed"}],
           "tasks": [{"gid": "1", "name": "TKT-9 Thing", "memberships": [{"project": {"gid": "77"}, "section": {"name": "Closed"}}],
                      "created_at": "2026-07-01T00:00:00Z", "modified_at": "2026-07-09T00:00:00Z",
                      "custom_fields": [{"name": "Estimated Time", "number_value": 5}]}],
           "stories": {"1": [
               {"gid": "s1", "resource_subtype": "section_changed", "created_at": "2026-07-05T00:00:00Z",
                "old_section": {"name": "To Do"}, "new_section": {"name": "Closed"}},
               {"gid": "s2", "resource_subtype": "section_changed", "created_at": "2026-07-06T00:00:00Z",
                "new_section": {"name": "Another project's column"}},
               {"gid": "s3", "resource_subtype": "comment_added", "created_at": "2026-07-07T00:00:00Z", "text": "done",
                "created_by": {"name": "QA"}}]}}
    snap = asana.from_raw(raw, r"TKT-\d+")
    it = snap["items"][0]
    check("a card's key, column and estimate are read", (it["key"], it["section"], it["fields"].get("Estimated Time")) == ("TKT-9", "Closed", 5))
    check("a move in this board's columns becomes a status event", [e["to"] for e in it["events"]] == ["Closed"], str(it["events"]))
    check("...and a move in some other project's column does not", all(e["to"] != "Another project's column" for e in it["events"]))
    check("a comment keeps a link to itself", it["comments"][0]["url"].endswith("/1/s3/f"), it["comments"][0]["url"])
    check("not being signed in is an explanation with every way in, not a stack trace",
          "Do not paste a token into a chat" in asana.NO_TOKEN and "browser_snapshot.js" in asana.NO_TOKEN
          and "Sign in in your browser" in asana.NO_TOKEN, asana.NO_TOKEN[:300])

    print("\nOther trackers: Jira and GitHub, through the same pipeline")
    import importlib.util
    def reader(name):
        spec = importlib.util.spec_from_file_location(f"reader_{name}", ROOT / "adapters" / name / "api.py")
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
        return mod
    jira, gh = reader("jira"), reader("github")
    adf = {"type": "doc", "content": [{"type": "paragraph", "content": [
        {"type": "text", "text": "Could you confirm "}, {"type": "mention", "attrs": {"text": "@Dana"}},
        {"type": "text", "text": " what should happen?"}]}]}
    raw = {"site": "https://acme.atlassian.net", "fields": [{"id": "customfield_10016", "name": "Story Points"}], "issues": [
        {"id": "10001", "key": "ACME-7", "fields": {
            "summary": "Login audit trail", "description": adf, "status": {"name": "Done", "statusCategory": {"key": "done"}},
            "issuetype": {"name": "Story", "subtask": False}, "created": "2026-07-01T09:00:00.000+0000",
            "updated": "2026-07-20T09:00:00.000+0000", "resolutiondate": "2026-07-20T09:00:00.000+0000",
            "resolution": {"name": "Done"}, "reporter": {"displayName": "Rafi"}, "assignee": {"displayName": "Tania"},
            "labels": ["backend"], "parent": {"id": "9000", "key": "ACME-1"}, "timeoriginalestimate": 28800,
            "customfield_10016": 5, "comment": {"total": 1, "comments": [
                {"id": "55", "created": "2026-07-03T09:00:00.000+0000", "author": {"displayName": "Tania"}, "body": adf}]}},
         "changelog": {"total": 3, "histories": [
            {"created": "2026-07-02T09:00:00.000+0000", "author": {"displayName": "Tania"},
             "items": [{"field": "status", "fromString": "To Do", "toString": "In Progress"}]},
            {"created": "2026-07-10T09:00:00.000+0000", "author": {"displayName": "Tania"},
             "items": [{"field": "assignee", "fromString": "a", "toString": "b"}]},
            {"created": "2026-07-20T09:00:00.000+0000", "author": {"displayName": "QA"},
             "items": [{"field": "status", "fromString": "In Progress", "toString": "Done"},
                       {"field": "resolution", "fromString": None, "toString": "Done"}]}]}},
        {"id": "10002", "key": "ACME-9", "fields": {
            "summary": "Password reset email never arrives", "status": {"name": "Closed", "statusCategory": {"key": "done"}},
            "issuetype": {"name": "Bug", "subtask": False}, "created": "2026-07-15T09:00:00.000+0000",
            "updated": "2026-07-16T09:00:00.000+0000", "resolution": {"name": "Cannot Reproduce"}, "labels": []},
         "changelog": {"total": 0, "histories": []}}]}
    prof = {"tracker": {"adapter": "jira", "url": "https://acme.atlassian.net", "story_point_field": "Story Points"}}
    jb = jira.from_raw(raw, prof, {"tracker_ref": "ACME", "name": "Acme"})
    j1, j2 = jb["items"]
    check("Jira: a status change in the changelog is a move on the board; other changes are not",
          [(e["from"], e["to"]) for e in j1["events"] if e["kind"] == "section"] == [("To Do", "In Progress"), ("In Progress", "Done")])
    check("Jira: rich text comes through as words, mentions included", "confirm @Dana what should happen" in j1["comments"][0]["text"], j1["comments"][0]["text"])
    check("Jira: estimate in hours, story points, type and resolution arrive as fields",
          (j1["fields"].get("Original Estimate"), j1["fields"].get("Story Points"), j1["fields"].get("Type")) == (8.0, 5, "Story"), str(j1["fields"]))
    check("Jira: a story under an epic is a deliverable of its own, not a sub-task to skip", j1["parent"] is None and j1["fields"].get("Epic") == "ACME-1")
    check("Jira: every item links back to the issue", j1["url"] == "https://acme.atlassian.net/browse/ACME-7")
    import classify
    from ledger import Ledger as _L
    kj, _w = classify.to_kif(jb, {**prof, "conventions": {"defect_by": "issue-type", "defect_values": ["Bug"]},
                                  "workflow": {"delivered_when": {"values": ["Done"]}, "closed_when": {"values": ["Done", "Closed"]}}},
                             {"name": "Acme"}, {}, _L(tmp / "jira-ledger.json"), "2026-09-18")
    check("Jira: the issue type decides what is a defect", [d["key"] for d in kj["defects"]] == ["ACME-9"] and [t["key"] for t in kj["tasks"]] == ["ACME-7"])
    check("Jira: a 'Cannot Reproduce' resolution is a rejected report", kj["defects"][0]["rejected"] == "Yes", str(kj["defects"][0].get("basis")))
    check("Jira: delivery is read from the move into the delivered state", kj["tasks"][0]["delivered"] == "2026-07-20")

    node = {"id": "I_1", "number": 42, "title": "Crash on empty config", "body": "Steps...", "url": "https://github.com/acme/web/issues/42",
            "state": "CLOSED", "stateReason": "COMPLETED", "createdAt": "2026-07-01T00:00:00Z", "updatedAt": "2026-07-09T00:00:00Z",
            "closedAt": "2026-07-09T00:00:00Z", "author": {"login": "rafi"}, "assignees": {"nodes": [{"login": "tania"}]},
            "labels": {"nodes": [{"name": "bug"}, {"name": "exisiting"}]}, "milestone": {"title": "4.2"}, "issueType": {"name": "Bug"},
            "projectItems": {"nodes": [{"project": {"number": 7, "title": "Web"}, "fieldValues": {"nodes": [
                {"__typename": "ProjectV2ItemFieldSingleSelectValue", "name": "Done", "field": {"name": "Status"}},
                {"__typename": "ProjectV2ItemFieldNumberValue", "number": 3, "field": {"name": "Estimate"}}]}}]},
            "comments": {"totalCount": 1, "nodes": [{"createdAt": "2026-07-02T00:00:00Z", "author": {"login": "qa"}, "bodyText": "Confirmed.", "url": "https://github.com/acme/web/issues/42#c1"}]},
            "timelineItems": {"nodes": [
                {"__typename": "ProjectV2ItemStatusChangedEvent", "createdAt": "2026-07-03T00:00:00Z", "previousStatus": "Todo", "status": "In review", "project": {"number": 7}},
                {"__typename": "ProjectV2ItemStatusChangedEvent", "createdAt": "2026-07-09T00:00:00Z", "previousStatus": "In review", "status": "Done", "project": {"number": 7}},
                {"__typename": "ClosedEvent", "createdAt": "2026-07-09T00:00:00Z", "stateReason": "COMPLETED"}]}}
    g1 = gh.item_from_issue(node, "acme/web", False, 7, "Status")
    check("GitHub: with a Project, status is the project's Status field and its changes are the history",
          g1["section"] == "Done" and [e["to"] for e in g1["events"] if e["kind"] == "section"] == ["In review", "Done"], str(g1["events"]))
    check("GitHub: labels arrive as tags, a project's number fields as fields", g1["tags"] == ["bug", "exisiting"] and g1["fields"].get("Estimate") == 3)
    plain = dict(node, projectItems={"nodes": []}, stateReason="NOT_PLANNED",
                 timelineItems={"nodes": [{"__typename": "ClosedEvent", "createdAt": "2026-07-09T00:00:00Z", "stateReason": "NOT_PLANNED"}]})
    g2 = gh.item_from_issue(plain, "acme/web", True, None, "Status")
    check("GitHub: without a Project, status is Open, Closed or Not planned", g2["section"] == "Not planned" and g2["events"][0]["to"] == "Not planned")
    check("GitHub: several repositories keep their issue numbers apart", g2["key"] == "web#42" and g1["key"] == "#42")
    gb = {"tracker": "github", "adapter": "github", "capabilities": gh.CAPABILITIES, "url": "", "items": [g1, g2], "project_name": "Web"}
    kg, _w = classify.to_kif(gb, {"conventions": {"defect_by": "label", "defect_values": ["bug"]},
                                  "workflow": {"closed_when": {"values": ["Done", "Closed"]}}}, {"name": "Web"}, {},
                             _L(tmp / "gh-ledger.json"), "2026-09-18")
    d1 = next(d for d in kg["defects"] if d["key"] == "#42")
    check("GitHub: a label decides what is a defect, and a mistyped 'exisiting' label is still read",
          d1["pre_existing"] == "Yes" and "exisiting" in d1["check"], d1["check"])
    check("GitHub: closed as not planned is a rejected report", next(d for d in kg["defects"] if d["key"] == "web#42")["rejected"] == "Yes")

    print("\nSigning in: every route, the best one suggested")
    import connect, os
    keep = {k: os.environ.pop(k, None) for k in ("ASANA_TOKEN", "ASANA_PAT", "JIRA_EMAIL", "JIRA_TOKEN", "JIRA_PAT", "GITHUB_TOKEN", "GH_TOKEN")}
    os.environ["KPI_COPILOT_HOME"] = str(tmp / "empty-home")
    try:
        text = connect.advise(["asana", "jira"])
        check("each service offers a browser sign-in, a token and a no-credential route",
              all(w in text for w in ("Sign in in your browser", "Personal access token", "API token", "already signed in")), text[:300])
        check("the suggestion is whatever needs no more setup on this machine", "<- suggested" in text.split("\n")[1] and "token" in text.split("\n")[1].lower(), text.split("\n")[1])
        check("a browser sign-in that still needs the company's one-time setup says so, and how to do it",
              "needs the one-time company setup first" in text and "http://localhost:8765/callback" in text)
        (tmp / "empty-home").mkdir(exist_ok=True)
        (tmp / "empty-home" / "asana_client.json").write_text(json.dumps({"client_id": "x", "client_secret": "y"}))
        check("once that setup exists, the browser sign-in becomes the suggestion", connect.recommend("asana")[0]["id"] == "browser")
        check("for a schedule, only routes that work with nobody there are offered",
              {r["id"] for r in connect.recommend("asana", unattended=True)} == {"token"})
        check("GitHub prefers the login the machine already has", connect.ROUTES["github"][0]["id"] == "cli")
        os.environ["ASANA_TOKEN"] = "test-token-not-real"
        check("a token in the environment is found, and counts as connected",
              connect.status("asana") ["via"] == "token" and connect.credential("asana") == {"bearer": "test-token-not-real"})
        os.environ.pop("ASANA_TOKEN")
        check("what a profile needs is worked out from it: tracker, and Google only if it is used",
              connect.needed({"tracker": {"adapter": "jira"}, "output": {"workbook": "xlsx"}}) == ["jira"]
              and connect.needed({"tracker": {"adapter": "github"}, "output": {"workbook_location": "https://drive.google.com/drive/folders/1ExampleDriveFolderId000000000000"}}) == ["github", "google"])
    finally:
        for k, v in keep.items():
            if v is not None:
                os.environ[k] = v
        os.environ.pop("KPI_COPILOT_HOME", None)
    r = run(["scripts/kpi.py", "auth", "asana", "--route", "token"], env={"KPI_COPILOT_HOME": str(tmp / "no-credentials")})
    check("a token is never taken through a chat: with no terminal, it refuses and says what to do instead",
          r.returncode == 1 and "must not pass through a chat" in r.stderr and "Sign in in your browser" in r.stderr, r.stderr[:300])
    r = run(["scripts/kpi.py", "run", "--profile", str(EX / "northwind-board" / "profile.yaml"), "--project", "northwind-q3",
             "--adapter", "trello"], env={"KPI_COPILOT_HOME": str(tmp / "no-credentials")})
    check("a tracker with no reader names the ones there are, and the way to start today",
          r.returncode != 0 and "asana" in (r.stderr + r.stdout) and "github" in (r.stderr + r.stdout) and "CSV" in (r.stderr + r.stdout),
          (r.stderr + r.stdout)[:300])

    print("\nOne command, from a board to a sheet")
    ws = tmp / "nw"
    shutil.copytree(EX / "northwind-board", ws)
    base = ["scripts/kpi.py", "run", "--profile", str(ws / "profile.yaml"), "--project", "northwind-q3",
            "--today", "2026-09-18"]
    env = {"KPI_COPILOT_HOME": str(tmp / "no-credentials")}
    r = run(base + ["--board", str(ws / "board.json")], env=env)
    check("the run finishes with a sheet on the first pass, nothing judged yet", r.returncode == 0, r.stderr[-600:])
    proj = ws / "northwind-q3"
    kif = json.loads((proj / "runs" / "2026-09-18" / "run.kif.json").read_text())
    rv = run(["scripts/validate_kif.py", "--kif", str(proj / "runs" / "2026-09-18" / "run.kif.json")])
    check("what it built is valid KIF", rv.returncode == 0, rv.stdout)
    bug4 = next(d for d in kif["defects"] if d["key"] == "Bug 04")
    check("the mistyped [Exisiting] bug is treated as already in the product", bug4["pre_existing"] == "Yes")
    check("...and its row says how that was decided", "read [Exisiting] as Existing" in bug4["check"], bug4["check"])
    bug8 = next(d for d in kif["defects"] if d["key"] == "Bug 08")
    check("a report from the client after handover belongs to the period handed over",
          bug8["period"] == "Initial Scope" and bug8["phase"] == "Post-release", f"{bug8['period']} {bug8['phase']}")
    doc = next(t for t in kif["tasks"] if t["key"].startswith("PLAN:"))
    check("a plan item with no card is not counted as late - unknown is not missed",
          doc["met_commitment"] is None and doc["met_client_date"] is None)
    check("a CR is recognised from the estimates even when its tag is mistyped",
          next(t for t in kif["tasks"] if t["key"] == "NW-161")["type"] == "CR")
    check("hours come from the plan, and the row says so",
          next(t for t in kif["tasks"] if t["key"] == "NW-101")["hours_source"].startswith("Project plan"))
    check("a handover date is read from the timeline sheet through its column mapping",
          kif["periods"][0]["handover_date"] == "2026-08-12", str(kif["periods"][0].get("handover_date")))
    rework = next(t for t in kif["tasks"] if t["key"] == "NW-104")
    check("closed and then reopened is rework", rework["reopened"] == "Yes")
    historical = classify.Classifier.__new__(classify.Classifier)
    historical.wf = {"closed_when": {"values": ["Ready for QA"]}}
    check("a reopened card keeps its historical close even while it is back in progress",
          historical._closed({"section": "Testing Failed", "events": [
              {"kind": "section", "from": "In Progress", "to": "Ready for QA",
               "at": "2026-09-10T09:00:00Z"},
              {"kind": "section", "from": "Ready for QA", "to": "Testing Failed",
               "at": "2026-09-17T09:00:00Z"},
          ]}) == "2026-09-10")
    first = next(t for t in kif["tasks"] if t["key"] == "NW-103")
    check("a QA failure during the first round is not", first["reopened"] == "No" and "first time" in (first["rework_evidence"] or ""))

    queue = json.loads((proj / "judge" / "queue.json").read_text())
    asked = {(i["item"]["title"][:24], a["field"]) for i in queue["items"] for a in i["asks"]}
    check("only what the rules were unsure of is put to the assistant", 3 <= len(queue["items"]) <= 8, str(len(queue["items"])))
    check("...including the tolerant tag match, to be confirmed",
          any(f == "pre_existing" and "Exisiting" in t for t, f in asked), str(asked))
    check("an excluded card is never worth a question", not any("Milestone" in t or "Checklist" in t for t, _ in asked))
    check("the queue carries the definitions it is to be judged by", "pre_existing" in queue["rubric"] and "reasons" in queue["rubric"])
    check("every question comes with the context to answer it, so nobody opens the board",
          all(i["item"].get("title") and i["item"].get("url") for i in queue["items"]))
    check("a missed KPI with no reason is asked for", any(n["kpi"] == "Defect Rate" for n in queue["notes"]))
    check("the run says what to do next, and tells the assistant not to go searching",
          "NEXT" in r.stdout and "Do not open the board or search anywhere else" in r.stdout, r.stdout[-500:])
    check("what nobody can know from outside is put to a person", "scope:deliverydocumentation" in r.stdout)

    print("\nThe assistant answers once, and it is kept")
    ids = {i["item"]["title"][:30]: i["item_id"] for i in queue["items"]}
    def iid(part): return next(v for k, v in ids.items() if part in k)
    answers = {"answers": [
        {"item_id": iid("Bug 04"), "field": "pre_existing", "value": "Yes", "why": "The tag is a misspelling of Existing."},
        {"item_id": iid("NW-150"), "field": "nature", "value": "Excluded", "why": "A clean-up nobody planned or asked for."},
        {"item_id": iid("Bug 09"), "field": "pre_existing", "value": "Maybe", "why": "unsure"},
        {"item_id": iid("NW-101"), "field": "understood", "value": "No", "why": ""}],
        "reasons": [{"period": "Initial Scope", "kpi": "Rework Rate", "why": "The filter reset after a tablet restart, which round one never tried."}]}
    (proj / "judge" / "answers.json").write_text(json.dumps(answers))
    r = run(["scripts/kpi.py", "judge", "--profile", str(ws / "profile.yaml"), "--project", "northwind-q3",
             "--today", "2026-09-18"], env=env)
    check("answers are folded in and the numbers recomputed in the same command", r.returncode == 0 and "Applied 2 judgements" in r.stdout, r.stdout[:300] + r.stderr[-300:])
    check("an answer outside the allowed values is refused, by name", "'Maybe' is not one of Yes, No" in r.stdout)
    check("an answer with no why is refused: a judgement that cannot say why is a guess", "no 'why' given" in r.stdout)
    q2 = json.loads((proj / "judge" / "queue.json").read_text())
    again = {(i["item_id"], a["field"]) for i in q2["items"] for a in i["asks"]}
    check("what was answered is not asked again", (iid("Bug 04"), "pre_existing") not in again and (iid("NW-150"), "nature") not in again)
    check("what was refused still is", (iid("NW-101"), "understood") in again)
    led = Ledger(proj / "ledger.json")
    check("the ledger records who decided, and why", led.get(iid("Bug 04"), "pre_existing")["by"] == "ai"
          and "misspelling" in led.get(iid("Bug 04"), "pre_existing")["why"])

    snap = json.loads((proj / "cache" / "board.json").read_text())
    for it in snap["items"]:
        if it["id"] == iid("Bug 04"):
            it["title"] = "[Tablet] Bug 04: Button labels cropped on small screens"
    (tmp / "moved.json").write_text(json.dumps(snap))
    run(base + ["--board", str(tmp / "moved.json")], env=env)
    q3 = json.loads((proj / "judge" / "queue.json").read_text())
    check("when a card changes, an assistant's answer about it no longer stands", True if any(
        i["item_id"] == iid("Bug 04") for i in q3["items"]) or
        next(d for d in json.loads((proj / "runs" / "2026-09-18" / "run.kif.json").read_text())["defects"]
             if d["key"] == "Bug 04")["pre_existing"] == "No" else False)
    run(base + ["--board", str(ws / "board.json")], env=env)

    print("\nThe sheet: looks made, not assembled - and is alive")
    from openpyxl import load_workbook
    book = next(proj.glob("KPI Tracker - *.xlsx"))
    wb = load_workbook(book)
    check("tabs are the ones a lead expects, in order",
          wb.sheetnames == ["Dashboard", "Period Overview", "KPI Summary", "Task Register", "Defect Register",
                            "Periods", "Open Questions", "Config", "Run Log", "PMS Push Log", "Read Me"], str(wb.sheetnames))
    tr, sm = wb["Task Register"], wb["KPI Summary"]
    fonts = {c.font.name for row in tr.iter_rows(min_row=1, max_row=12) for c in row if c.value is not None}
    check("every written cell is Arial; nothing falls back to Calibri", fonts == {"Arial"}, str(fonts))
    check("headers are the navy band with white bold text",
          tr["B3"].fill.fgColor.rgb.endswith("1F3864") and tr["B3"].font.b and tr["B3"].font.color.rgb.endswith("FFFFFF"))
    heads = {c.value: c for c in tr[3]}
    check("yellow is yours, grey is worked out, white was read",
          tr.cell(4, heads["Item Type"].column).fill.fgColor.rgb.endswith("FFF2CC")
          and tr.cell(4, heads["Status"].column).fill.fgColor.rgb.endswith("F2F2F2")
          and tr.cell(4, heads["Title"].column).fill.fgColor.rgb.endswith("FFFFFF"))
    check("dates are dates", tr.cell(4, heads["Delivered"].column).number_format == "yyyy-mm-dd"
          and hasattr(tr.cell(4, heads["Delivered"].column).value, "year"))
    check("the numbers are live formulas over the registers", str(sm["E5"].value).startswith("=COUNTIFS(") and str(sm["G5"].value).startswith("=IF("))
    check("the note for PMS follows both editable fields without a leading separator", sm["N5"].value == '=L5&IF(TRIM(M5)="","",IF(TRIM(L5)="",""," || ")&TRIM(M5))')
    check("the engine's own figure sits beside each live one", isinstance(sm["O5"].value, (int, float)) and "Run again before pushing" in str(sm["P5"].value))
    every = [str(c.value) for w in wb.worksheets for row in w.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    check("no formula relies on how a spreadsheet treats an empty cell in a criterion", not any('"<>"' in f for f in every))
    check("only functions Excel and Google both have", not any(fn in f for f in every for fn in ("LET(", "TEXTJOIN(", "FILTER(", "MINIFS("))
          or all("DUMMYFUNCTION" in f for f in every if "FILTER(" in f))
    check("Excel dashboard uses native bars without undefined-function error markers",
          any("REPT(" in f for f in every) and not any("DUMMYFUNCTION" in f or "SPARKLINE" in f for f in every))
    check("the row id a rerun needs is there but out of sight", tr.column_dimensions[tr.cell(3, heads["Row ID"].column).column_letter].hidden)
    check("a card with no ticket key is not labelled with a sixteen-digit id",
          not any(str(tr.cell(r, heads["Ticket"].column).value or "").isdigit() for r in range(4, 25)))
    try:
        import formulas  # noqa: F401  (optional: pip install formulas)
        res = json.loads((proj / "runs" / "2026-09-18" / "results.json").read_text())
        sol = formulas.ExcelModel().loads(str(book)).finish().calculate()
        live = {str(k).split("!")[-1].strip("'"): v.value[0][0] for k, v in sol.items() if "KPI SUMMARY" in str(k).upper()}
        bad = []
        for rr in range(4, sm.max_row + 1):
            per, kpi = sm.cell(rr, 1).value, sm.cell(rr, 3).value
            if not per:
                continue
            m = next(m for p in res["periods"] if p["period"] == per for m in p["measures"] if m["name"] == kpi)
            got = live.get(f"G{rr}")
            got = None if got in ("", None) else round(float(got), 2)
            if got != (None if m["value"] is None else round(m["value"], 2)):
                bad.append(f"{per} {kpi}: sheet {got}, engine {m['value']}")
        check("every live formula gives the engine's number", not bad, "; ".join(bad))
    except ImportError:
        check("formula verification dependency is installed", False, "Install requirements-test.txt; formulas must not be skipped.")

    print("\nA person's edit survives the rerun")
    ds = wb["Defect Register"]
    dh = {c.value: c.column for c in ds[3]}
    for rr in range(4, 20):
        if ds.cell(rr, dh["Ticket"]).value == "Bug 02":
            ds.cell(rr, dh["Pre-existing?"]).value = "Yes"
        if ds.cell(rr, dh["Ticket"]).value == "Bug 04":
            ds.cell(rr, dh["Pre-existing?"]).value = "No"
    for rr in range(4, sm.max_row + 1):
        if sm.cell(rr, 1).value == "Initial Scope" and sm.cell(rr, 3).value == "CR Rate":
            sm.cell(rr, 13).value = "Both additions were approved before work began."
    cf = wb["Config"]
    for rr in range(1, 40):
        if cf.cell(rr, 2).value == "Count observations as defects?":
            cf.cell(rr, 3).value = "Yes"
    import datetime as _dt
    wb["Periods"].cell(6, 8).value = _dt.date(2026, 9, 16)
    wb["Open Questions"].cell(4, 5).value = _dt.date(2026, 8, 12)
    wb.save(book)
    r = run(base + ["--offline"], env=env)
    check("the run reads the sheet back before rebuilding it", "Read back from the sheet" in r.stdout, r.stdout[:600])
    led = Ledger(proj / "ledger.json")
    b2 = next(i["id"] for i in json.loads((ws / "board.json").read_text())["items"] if "Bug 02" in i["title"])
    check("a Yes/No typed in the sheet becomes a person's judgement", (led.get(b2, "pre_existing") or {}).get("by") == "human")
    check("...and outranks the assistant's", led.get(iid("Bug 04"), "pre_existing")["value"] == "No"
          and led.get(iid("Bug 04"), "pre_existing")["by"] == "human")
    import yaml
    facts = yaml.safe_load((proj / "facts" / "periods.yaml").read_text())
    check("a date typed on the Periods tab lands in the project's facts", str(facts["periods"][1]["handover_date"]) == "2026-09-16")
    reasons = yaml.safe_load((proj / "facts" / "reasons.yaml").read_text())
    check("a reason typed on KPI Summary is kept as the person's", reasons["Initial Scope"]["CR Rate"].startswith("Both additions")
          and reasons["_authors"]["Initial Scope|CR Rate"] == "human")
    check("an answer typed on Open Questions is kept, and the question stops being asked",
          "scope:deliverydocumentation" not in r.stdout.split("NEXT")[-1])
    check("a counting rule changed in the sheet is reported and NOT applied", "NOT applied" in r.stdout and "policy.count_observations" in r.stdout)
    check("what moved since the last run is said out loud", "Moved since the last run" in r.stdout)
    (proj / "judge" / "answers.json").write_text(json.dumps({"answers": [
        {"item_id": iid("Bug 04"), "field": "pre_existing", "value": "Yes", "why": "Trying to overrule a person."}]}))
    r = run(["scripts/kpi.py", "judge", "--profile", str(ws / "profile.yaml"), "--project", "northwind-q3",
             "--today", "2026-09-18", "--no-rerun"], env=env)
    check("an assistant cannot overrule what a person decided", "a person already answered this" in r.stdout, r.stdout)

    print("\nThe same sheet, for Google")
    k = json.loads((proj / "runs" / "2026-09-18" / "run.kif.json").read_text())
    res = json.loads((proj / "runs" / "2026-09-18" / "results.json").read_text())
    tabs = sheet_model.build(k, res, {"profile": {}})
    register = next(t for t in tabs if t.name == "Task Register")
    fresh = sheet_google.tab_requests(register, 111, 4, None)
    again = sheet_google.tab_requests(register, 111, 4, {"properties": {"gridProperties": {"rowCount": 10, "columnCount": 5}},
                                                       "conditionalFormats": [{}, {}]})
    kinds = lambda reqs: [next(iter(x)) for x in reqs]                                    # noqa: E731
    check("a tab that is not there yet is added", kinds(fresh)[0] == "addSheet")
    check("a tab that is there is cleared and rebuilt in place, old colour rules first",
          kinds(again)[:4] == ["deleteConditionalFormatRule", "deleteConditionalFormatRule", "unmergeCells", "updateCells"],
          str(kinds(again)[:5]))
    check("formats go down as runs, not once per cell", 20 < kinds(fresh).count("repeatCell") < 400, str(kinds(fresh).count("repeatCell")))
    check("dropdowns and colour rules come across", "setDataValidation" in kinds(fresh) and "addConditionalFormatRule" in kinds(fresh))
    rules = [c["formula"] for t in tabs for c in t.cond]
    check("no colour rule looks at another tab, which Google refuses", not any("!" in f for f in rules), str([f for f in rules if "!" in f]))
    cells = [v for req in sheet_google.tab_requests(next(t for t in tabs if t.name == "Dashboard"), 5, 1, None) if "updateCells" in req
             for row in req["updateCells"].get("rows", []) for v in row["values"]]
    check("in Google the bars are real SPARKLINEs", any("SPARKLINE" in (c.get("userEnteredValue") or {}).get("formulaValue", "") for c in cells))
    check("a Drive link, a folder link and a bare id all resolve",
          sheet_google.G.file_id("https://docs.google.com/spreadsheets/d/1AbcDefGhiJklMnoPqrStuVwxYz0123456789/edit#gid=0") == "1AbcDefGhiJklMnoPqrStuVwxYz0123456789"
          and sheet_google.G.file_id("https://drive.google.com/drive/folders/1ExampleDriveFolderId000000000000") == "1ExampleDriveFolderId000000000000"
          and sheet_google.G.file_id("runs/x.xlsx") is None)
    prof = (ws / "profile.yaml").read_text().replace("  workbook: xlsx", "  workbook: google-sheets\n  workbook_location: https://drive.google.com/drive/folders/1ExampleDriveFolderId000000000000")
    (ws / "profile.yaml").write_text(prof)
    r = run(base + ["--offline"], env=env)
    r = run(base + ["--board", str(ws / "board.json")], env=env)
    check("asked for a Google Sheet with Google not connected: the local sheet is still written",
          r.returncode == 0 and next(proj.glob("KPI Tracker - *.xlsx")).exists())
    check("...and offers connection or local review without replacing unrelated spreadsheet tabs",
          "Google Sheet not updated" in r.stdout and "kpi.py auth google" in r.stdout
          and "choose local workbook output" in r.stdout and "Replace spreadsheet" not in r.stdout, r.stdout[-900:])

    print("\nSources: those on the list, and no others")
    import sources as S
    st = S.pull({"sources": {"plan": {"kind": "pdf", "ref": "https://drive.google.com/file/d/1ExamplePlanFileId000000000000000/view"}}},
                tmp / "src-proj", tmp, offline=True)
    check("a source it cannot reach is not guessed at: it says where to drop a copy", st[0]["state"] == "missing" and "inbox" in st[0]["note"], str(st))
    (tmp / "src-proj" / "inbox").mkdir(parents=True, exist_ok=True)
    (tmp / "src-proj" / "inbox" / "plan.pdf").write_bytes(b"%PDF-1.4 demo")
    st = S.pull({"sources": {"plan": {"kind": "pdf", "ref": "https://drive.google.com/file/d/1ExamplePlanFileId000000000000000/view"}}},
                tmp / "src-proj", tmp, offline=True)
    check("a file dropped in the inbox is picked up", st[0]["state"] == "fresh" and st[0]["path"].endswith("plan.pdf"))
    said = S.staleness(st, {"plan": {"items": [{"title": "x"}], "source": {"fingerprint": "older", "as_of": "2026-07-01"}}})
    check("a digest written from an older version of its source is called out", said and "changed after" in said[0], str(said))
    check("an empty digest says what to read and where to write it", "Read it once" in S.staleness(st, {"plan": {}})[0])
    check("the run log names what was deliberately not read",
          any("chat, mail" in str(c.value) for row in load_workbook(next(proj.glob("KPI Tracker - *.xlsx")))["Run Log"].iter_rows() for c in row if c.value))


def question_tests(tmp: Path) -> None:
    """A question names what it is about, and the sheet says where each answer stands."""
    import types
    import yaml
    import kpi
    import note_policy
    import sheet_readback as R
    from ledger import Ledger

    print("\nQuestions: which items, and where each answer stands")
    kif = json.loads((EX / "northwind-q3" / "run.kif.json").read_text())
    for t in kif["tasks"]:
        if t.get("key") == "TKT-3479":
            t["hours_qa"] = None
            t["closed"] = t.get("closed") or t.get("delivered")
    (tmp / "q.kif.json").write_text(json.dumps(kif))
    cfg = yaml.safe_load((EX / "northwind-q3" / "profile.yaml").read_text())
    cfg.setdefault("sources", {})["missing_estimate"] = "skip"
    cfg["sources"]["hours_basis"] = "dev+qa"
    (tmp / "q.yaml").write_text(yaml.safe_dump(cfg))
    res = compute(tmp / "q.kif.json", tmp / "q.yaml", None, tmp / "q.results.json")
    vel = next(m for p in res["periods"] for m in p["measures"]
               if m["name"] == "Velocity" and any("estimate" in x for x in m["review_items"]))
    detail = next(x for x in vel["review_items"] if "estimate" in x)
    items = note_policy.items_for(vel, detail)
    mine = next((i for i in items if i["key"] == "TKT-3479"), {})
    check("a missing-estimate question names every item, its link and what is missing",
          "no QA estimate" in mine.get("detail", "") and "link" in mine and len(items) == len({i["key"] for i in items}), str(items))
    check("...and the question's id does not depend on the items listed",
          note_policy.question("P", "Velocity", detail, items)["id"] == note_policy.question("P", "Velocity", detail)["id"])

    cfg["sources"]["missing_estimate"] = "partial"
    (tmp / "qp.yaml").write_text(yaml.safe_dump(cfg))
    part = compute(tmp / "q.kif.json", tmp / "qp.yaml", None, tmp / "qp.results.json")
    pv = next(m for p in part["periods"] for m in p["measures"] if m["name"] == "Velocity" and p["period"] == vel.get("period", p["period"])
              and "TKT-3479" in (m.get("counted_keys") or []))
    check("missing_estimate: partial counts the part of an estimate that is recorded, and says so",
          not any("estimate" in x for x in pv["review_items"]) and "part of" in pv["note"], pv["note"])

    rk = json.loads((EX / "northwind-q3" / "run.kif.json").read_text())
    for t in rk["tasks"]:
        if t.get("reopened") == "No":
            t["reopened"] = None
    (tmp / "rk.kif.json").write_text(json.dumps(rk))
    cfg = yaml.safe_load((EX / "northwind-q3" / "profile.yaml").read_text())
    cfg.setdefault("workflow", {}).setdefault("reopened_when", {})["no_history"] = "not-reopened"
    (tmp / "rk.yaml").write_text(yaml.safe_dump(cfg))
    before = measure(compute(tmp / "rk.kif.json", EX / "northwind-q3" / "profile.yaml", None, tmp / "rk0.json"),
                     "Initial Scope", "Rework Rate")
    after = measure(compute(tmp / "rk.kif.json", tmp / "rk.yaml", None, tmp / "rk1.json"), "Initial Scope", "Rework Rate")
    check("no_history: not-reopened counts an unrecorded task as not reopened, and the note says so",
          (after["denominator"] or 0) > (before["denominator"] or 0) and "counted as not reopened" in after["note"],
          f"{before['denominator']} -> {after['denominator']}: {after['note']}")

    grp = json.loads((EX / "northwind-q3" / "run.kif.json").read_text())
    done = [t for t in grp["tasks"] if t.get("period") == "Initial Scope" and t.get("type") == "Task"
            and t.get("rework_closed", t.get("closed"))][:2]
    for t in done:
        t.update(reopened=None, reopen_count=None, effort_group="grp-1", _item=t.get("_item") or t["key"])
    grp["tasks"].append({"period": "Initial Scope", "key": "GRP-1", "title": "A grouped feature", "type": "Excluded",
                         "effort_only": True, "_item": "grp-1", "reopened": "Yes", "reopen_count": 1,
                         "delivered": done[0].get("delivered"), "closed": done[0].get("closed"),
                         "rework_closed": done[0].get("rework_closed", done[0].get("closed")), "planned": None})
    (tmp / "grp.kif.json").write_text(json.dumps(grp))
    cfg = yaml.safe_load((EX / "northwind-q3" / "profile.yaml").read_text())
    cfg.setdefault("workflow", {}).setdefault("reopened_when", {}).update(no_history="not-reopened")
    (tmp / "grp0.yaml").write_text(yaml.safe_dump(cfg))
    base = measure(compute(tmp / "grp.kif.json", tmp / "grp0.yaml", None, tmp / "grp0.json"), "Initial Scope", "Rework Rate")
    check("tickets of a reopened group are never assumed not reopened", "counted as not reopened" not in base["note"]
          or base["denominator"] < len([t for t in grp["tasks"] if t.get("period") == "Initial Scope"]), base["note"])
    cfg["workflow"]["reopened_when"]["group_history"] = "count-group"
    (tmp / "grp1.yaml").write_text(yaml.safe_dump(cfg))
    g1 = measure(compute(tmp / "grp.kif.json", tmp / "grp1.yaml", None, tmp / "grp1.json"), "Initial Scope", "Rework Rate")
    check("group_history: count-group counts the group once, from its card, and the note says so",
          g1["numerator"] == (base["numerator"] or 0) + 1 and "counted once each from the group card" in g1["note"]
          and "GRP-1" in g1["counted_keys"], f"{base['numerator']}/{base['denominator']} -> {g1['numerator']}/{g1['denominator']}: {g1['note']}")

    check("a date written in a sentence is read", R._dates_in("Handed over to the client on 09/25. Thread", 2026) == ["2026-09-25"])
    led = Ledger(tmp / "q.ledger.json")
    facts = {"periods": {"periods": [{"name": "Cycle 2"}]}}
    applied = R._file_answer("handover:Cycle 2", "Handed over to the client on 09/25", led, facts)
    check("...so a handover answered in words is applied", facts["periods"]["periods"][0].get("handover_date", "").endswith("-09-25")
          and applied, str(facts))
    check("...but two dates in one answer are left for a person to settle",
          R._file_answer("handover:Cycle 2", "Built 09/24, handed over 09/25", led, {"periods": {"periods": [{"name": "Cycle 2"}]}}) is None)

    ws = types.SimpleNamespace(ledger=led, today="2026-09-26")
    current = [{"id": "review:Cycle 2|Velocity:abc", "about": "Cycle 2 · Velocity", "question": "Review this", "blocking": False},
               {"id": "handover:Cycle 3", "about": "Cycle 3", "question": "When?", "blocking": True}]
    led.set_answer("review:Cycle 2|Velocity:abc", "Add these hours.", "human", "answered in the sheet")
    led.set_answer("handover:Cycle 2", "2026-09-25", "human", "answered in the sheet")
    led.settle("handover:Cycle 2", "applied", "handover date set to 09/25", "tool")
    led.note_question("handover:Cycle 2", "Cycle 2", "When did it reach the client?", "2026-09-20")
    rows = {r["id"]: r for r in kpi.question_rows(ws, current)}
    check("an unanswered question is Open, an unacted answer waits for the assistant",
          rows["handover:Cycle 3"]["status"] == "Open" and rows["review:Cycle 2|Velocity:abc"]["status"] == "Answered - to apply"
          and rows["review:Cycle 2|Velocity:abc"]["blocking"])
    check("an answer the tool used is Applied, and a question no longer raised stays in the history",
          rows["handover:Cycle 2"]["status"] == "Applied" and rows["handover:Cycle 2"]["asked_on"] == "2026-09-20"
          and rows["handover:Cycle 2"]["done"] == "handover date set to 09/25")
    led.settle("review:Cycle 2|Velocity:abc", "resolved", "counted the recorded hours", "assistant")
    rows = {r["id"]: r for r in kpi.question_rows(types.SimpleNamespace(ledger=led, today="2026-09-27"), current)}
    check("an answer acted on is Resolved, with what was done, and keeps the day it was first asked",
          rows["review:Cycle 2|Velocity:abc"]["status"] == "Resolved"
          and rows["review:Cycle 2|Velocity:abc"]["asked_on"] == "2026-09-26"
          and rows["review:Cycle 2|Velocity:abc"]["done"] == "counted the recorded hours")


if __name__ == "__main__":
    raise SystemExit(main())
