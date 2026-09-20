#!/usr/bin/env python3
"""
The workbook model, written as an .xlsx - and read back.

Writing is a straight translation of sheet_model's description into openpyxl. Three details
are worth knowing:

  * Every cell gets the base font. A workbook that falls back to Calibri in the cells nobody
    styled looks assembled rather than made.
  * Excel gets native formulas for the dashboard bars. The Google writer uses the model's
    SPARKLINE variant. Undefined-function wrappers caused error markers in desktop Excel.
  * Formulas are written without cached results, as openpyxl must. Excel, Numbers, ONLYOFFICE
    and Google all calculate on open. A viewer that does not (a file preview) shows the grey
    cells empty - the numbers are also printed by the run and saved in results.json.

Reading back is how a person's edits survive a rerun. `snapshot()` records what was written
into every yellow cell; `read_grid()` reads what is there now; sheet_readback.py works out
the difference and files each edit where it belongs.
"""

from __future__ import annotations

import datetime as dt
import os
import tempfile
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import sheet_model as M

FONT = "Arial"
_THIN = Side(style="thin", color=M.LINE)
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style(spec: dict) -> tuple[Font, PatternFill | None, Alignment, Border | None]:
    font = Font(name=FONT, size=spec.get("size") or 10, bold=bool(spec.get("b")), italic=bool(spec.get("i")),
                color=spec.get("color") or "000000", underline="single" if spec.get("u") else None)
    fill = PatternFill("solid", fgColor=spec["fill"]) if spec.get("fill") else None
    align = Alignment(horizontal=spec.get("h"), vertical=spec.get("v"), wrap_text=bool(spec.get("wrap")))
    return font, fill, align, (_BOX if spec.get("box") else None)


def write(tabs: list[M.Tab], out: Path) -> Path:
    wb = load_workbook(out) if out.exists() else Workbook()
    if not out.exists():
        wb.remove(wb.active)
    for tab in tabs:
        if tab.name[:31] in wb.sheetnames:
            del wb[tab.name[:31]]
    cache = {name: _style(spec) for name, spec in M.STYLES.items()}
    for tab in tabs:
        ws = wb.create_sheet(tab.name[:31])
        ws.sheet_properties.tabColor = tab.color
        ws.sheet_view.showGridLines = tab.gridlines
        for (r, c), cell in tab.cells.items():
            x = ws.cell(row=r, column=c)
            if cell.get("f"):
                x.value = cell["f"]
            elif "v" in cell:
                x.value = cell["v"]
                if isinstance(cell["v"], str):
                    x.data_type = "s"  # tracker text beginning '=' is text, never an executable formula
            font, fill, align, box = cache.get(cell.get("style") or "text", cache["text"])
            x.font, x.alignment = font, align
            if fill:
                x.fill = fill
            if box:
                x.border = box
            if cell.get("fmt"):
                x.number_format = cell["fmt"]
            if cell.get("link"):
                x.hyperlink = cell["link"]
        for c, w in tab.widths.items():
            ws.column_dimensions[get_column_letter(c)].width = w
        for c in tab.hidden_cols:
            ws.column_dimensions[get_column_letter(c)].hidden = True
        for r, h in tab.heights.items():
            ws.row_dimensions[r].height = h
        for r1, c1, r2, c2 in tab.merges:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        if tab.filter_range:
            r1, c1, r2, c2 = tab.filter_range
            ws.auto_filter.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
        if any(tab.freeze):
            ws.freeze_panes = ws.cell(row=tab.freeze[0] + 1, column=tab.freeze[1] + 1)
        for v in tab.validations:
            r1, c1, r2, c2 = v["range"]
            formula = f"={v['source']}" if v.get("source") else '"' + ",".join(v["values"]) + '"'
            dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
            dv.add(f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}")
            ws.add_data_validation(dv)
        for rule in tab.cond:
            r1, c1, r2, c2 = rule["range"]
            ws.conditional_formatting.add(
                f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}",
                FormulaRule(formula=[rule["formula"].lstrip("=")], stopIfTrue=False,
                            fill=(PatternFill(start_color=rule["fill"], end_color=rule["fill"], fill_type="solid")
                                  if rule.get("fill") else None),
                            font=Font(color=rule["color"]) if rule.get("color") else None))
    wb.active = wb.sheetnames.index(tabs[0].name[:31])
    out.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(suffix=".xlsx", dir=out.parent)
    os.close(fd)
    try:
        wb.save(name)
        os.replace(name, out)
    finally:
        if os.path.exists(name):
            os.unlink(name)
    return out


# --------------------------------------------------------------------------------------
# reading back
# --------------------------------------------------------------------------------------

def plain_value(v: Any) -> Any:
    """One comparable form for whatever a cell holds, in either kind of sheet."""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def read_grid(path: Path) -> Callable[[str, int, int], Any]:
    """grid(tab, row, col) -> what is typed there now. A formula reads as None: nobody's
    edit lives in a formula cell."""
    wb = load_workbook(path)

    def grid(tab: str, r: int, c: int) -> Any:
        if tab[:31] not in wb.sheetnames:
            return None
        v = wb[tab[:31]].cell(row=r, column=c).value
        if isinstance(v, str) and v.startswith("="):
            return None
        return plain_value(v)

    grid.rows = lambda tab: wb[tab[:31]].max_row if tab[:31] in wb.sheetnames else 0   # type: ignore[attr-defined]
    return grid
