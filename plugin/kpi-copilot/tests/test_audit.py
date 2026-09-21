"""Failure-path regressions from the product audit. No credentials or live writes."""
from __future__ import annotations
import contextlib
import copy
import importlib.util
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import board, classify, judge, kpi, kpi_engine, kpi_registry, ledger
import pms_push, profile_lib, profile_tool, sheet_google, sheet_model, sheet_readback, sheet_xlsx, sources
import yaml
from openpyxl import Workbook, load_workbook


def adapter(name):
    spec = importlib.util.spec_from_file_location('audit_' + name, ROOT / 'adapters' / name / 'api.py')
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class AuditTests(unittest.TestCase):
    def grouped_example(self):
        def item(key, title, parent=None, count=0, history=True):
            return {"id": key, "key": key, "title": title, "url": f"https://tracker.example/items/{key}",
                    "parent": parent, "subtasks": count, "section": "Done" if history else None,
                    "created_at": "2025-03-01", "events": [
                        {"kind": "section", "at": "2025-03-03", "from": "Doing", "to": "QA"},
                        {"kind": "section", "at": "2025-03-04", "from": "QA", "to": "Done"}
                    ] if history else [], "comments": [], "fields": {}}
        snap = {"capabilities": ["status_history", "comments"], "items": [
            item("BASE", "Original feature", count=2),
            item("BASE-A", "API component", "BASE", history=False),
            item("BASE-B", "Mobile component", "BASE", history=False),
            item("EXTRA", "[QA] Approved service"),
            item("EXTRA-A", "[Existing] Bug 1: repair old validation"),
            item("EXTRA-B", "[Existing] Bug 2: repair old rendering"),
        ]}
        profile = {"tracker": {"options": {"include_subtasks": True}},
                   "conventions": {"exclude_patterns": [r"^\[QA\]"], "defect_pattern": r"Bug \d+",
                       "additional_request_label": "Additional Request",
                       "grouping": {"split_source_children": True, "inherit_parent_delivery": True,
                                    "linked_members": {"EXTRA": ["EXTRA-A", "EXTRA-B"]}}},
                   "workflow": {"delivered_when": {"values": ["QA"]},
                                "closed_when": {"values": ["Done"]}, "reopened_when": {"values": ["Doing"]}},
                   "sources": {"hours_basis": "dev+qa"}}
        facts = {"periods": {"periods": [{"name": "Cycle", "handover_date": "2025-03-05"}]},
                 "plan": {"items": [{"board_key": "BASE", "title": "Original feature", "dev_hours": 12, "qa_hours": 4}]},
                 "estimates": {"items": [{"board_key": "EXTRA", "title": "[QA] Approved service", "dev_hours": 8, "qa_hours": 4}]}}
        return snap, profile, facts

    def test_group_members_count_once_and_keep_aggregate_hours_once(self):
        snap, profile, facts = self.grouped_example()
        kif, work = classify.to_kif(snap, profile, {"name": "Example"}, facts, None, "2025-03-08")
        engine = kpi_engine.Engine(kif, profile, kpi._registry(self.ws()))
        self.assertEqual([t['key'] for t in engine.deliverables_in('Cycle')], ['BASE-A', 'BASE-B', 'EXTRA-A', 'EXTRA-B'])
        self.assertEqual(engine.velocity(kif['periods'][0]).value, 28)
        self.assertEqual(len(kif['defects']), 2)
        self.assertEqual(len({d['_row'] for d in kif['defects']}), 2)
        self.assertEqual(engine.counted_defects('Cycle')[0], [])
        self.assertFalse(any(q['id'].startswith('scope:') for q in work['questions']))
        members = [t for t in kif['tasks'] if t.get('effort_group') == 'BASE']
        self.assertTrue(all(t['delivered'] == '2025-03-03' and t['hours_dev'] is None for t in members))
        self.assertTrue(all(t['reopened'] is None for t in members))

    def test_grouped_workbook_formulas_match_engine_and_preserve_defect_identity(self):
        snap, profile, facts = self.grouped_example()
        kif, _ = classify.to_kif(snap, profile, {"name": "Example"}, facts, None, "2025-03-08")
        self.assert_grouped_formulas(kif, profile)

    def assert_grouped_formulas(self, kif, profile):
        import formulas
        registry = kpi._registry(self.ws())
        results = {'periods': [p.as_dict() for p in kpi_engine.Engine(kif, profile, registry).run()]}
        tabs = sheet_model.build(kif, results, {'profile': profile, 'registry': registry, 'as_of': '2025-03-08'})
        path = self.base / 'groups.xlsx'
        sheet_xlsx.write(tabs, path)
        sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
        live = {str(k).split('!')[-1].strip("'"): v.value[0][0] for k,v in sol.items() if 'KPI SUMMARY' in str(k).upper()}
        sheet = load_workbook(path)['KPI Summary']
        for row in range(4, sheet.max_row + 1):
            name = sheet.cell(row,3).value
            if not name: continue
            expected = next(m['value'] for m in results['periods'][0]['measures'] if m['name']==name)
            actual = live[f'G{row}']
            self.assertEqual(None if actual in ('',None) else round(float(actual),2), expected, name)

    def test_partial_group_delivery_does_not_invent_a_per_ticket_estimate(self):
        snap, profile, facts = self.grouped_example()
        profile['conventions']['grouping']['inherit_parent_delivery'] = False
        snap['items'][1]['events'] = copy.deepcopy(snap['items'][0]['events'])
        snap['items'][1]['section'] = 'Done'
        kif, _ = classify.to_kif(snap, profile, {}, facts, None, '2025-03-08')
        engine = kpi_engine.Engine(kif, profile, kpi._registry(self.ws()))
        self.assertIsNone(engine.velocity(kif['periods'][0]).value)
        self.assert_grouped_formulas(kif, profile)

    def test_pending_delivery_check_cannot_turn_a_blank_handoff_into_on_time(self):
        snap, profile, facts = self.grouped_example()
        for item in snap['items']:
            item['events'], item['section'] = [], None
        facts['periods']['periods'][0].update(client_date='2025-03-30', client_check='Delivery')
        kif, _ = classify.to_kif(snap, profile, {}, facts, None, '2025-03-08')
        result = kpi_engine.Engine(kif, profile, kpi._registry(self.ws())).client_expectation(kif['periods'][0])
        self.assertIsNone(result.value)
        self.assert_grouped_formulas(kif, profile)

    def test_old_card_history_does_not_credit_a_later_planned_period(self):
        item = {"id": "REUSED", "key": "REUSED", "title": "Installer repair",
                "created_at": "2025-04-01", "section": "To Do", "fields": {}, "comments": [],
                "events": [
                    {"kind": "section", "at": "2025-04-03", "from": "Doing", "to": "QA"},
                    {"kind": "section", "at": "2025-04-04", "from": "QA", "to": "Done"},
                ]}
        snap = {"capabilities": ["status_history"], "items": [item]}
        profile = {"workflow": {"delivered_when": {"values": ["QA", "Done"]},
                                 "closed_when": {"values": ["Done"]}}}
        facts = {"periods": {"periods": [{"name": "Current phase", "start": "2026-09-07"}]},
                 "plan": {"items": [{"board_key": "REUSED", "title": "Installer repair",
                                      "period": "Current phase", "dev_hours": 40, "qa_hours": 12}]}}
        kif, _ = classify.to_kif(snap, profile, {}, facts, None, "2026-09-21")
        row = next(t for t in kif["tasks"] if t["key"] == "REUSED")
        self.assertIsNone(row["delivered"])
        self.assertIsNone(row["closed"])

    def test_explicit_source_delivery_can_precede_period_start(self):
        item = {"id": "CARRY", "key": "CARRY", "title": "Carried item",
                "created_at": "2025-04-01", "section": "Done", "fields": {}, "comments": [],
                "events": [{"kind": "section", "at": "2025-04-04", "from": "QA", "to": "Done"}]}
        snap = {"capabilities": ["status_history"], "items": [item]}
        profile = {"workflow": {"delivered_when": {"values": ["Done"]},
                                 "closed_when": {"values": ["Done"]}}}
        facts = {"periods": {"periods": [{"name": "Current phase", "start": "2026-09-07"}]},
                 "plan": {"items": [{"board_key": "CARRY", "title": "Carried item",
                                      "period": "Current phase", "dev_hours": 1,
                                      "delivered": "2026-09-05"}]}}
        kif, _ = classify.to_kif(snap, profile, {}, facts, None, "2026-09-21")
        row = next(t for t in kif["tasks"] if t["key"] == "CARRY")
        self.assertEqual(row["delivered"], "2026-09-05")
        self.assertIsNone(row["closed"])

    def test_explicit_member_answers_survive_missing_individual_history(self):
        snap, profile, facts = self.grouped_example()
        stored = ledger.Ledger(self.base / 'member-answers.json')
        stored.set('BASE-A', 'understood', 'No', 'human', 'The client clarified the required behavior.')
        stored.set('BASE-A', 'reopened', 'Yes', 'human', 'The individual item was returned from QA.')
        kif, _ = classify.to_kif(snap, profile, {}, facts, stored, '2025-03-08')
        member = next(t for t in kif['tasks'] if t['key'] == 'BASE-A')
        self.assertEqual((member['understood'], member['reopened']), ('No', 'Yes'))

    def test_plain_english_review_answer_resolves_unassessed_comprehension_rows(self):
        snap, profile, facts = self.grouped_example()
        stored = ledger.Ledger(self.base / 'batch-review.json')
        stored.set_answer('review:Cycle|Task Comprehension:example',
                          "It's fine. Count those items as understood.", 'human')
        kif, _ = classify.to_kif(snap, profile, {}, facts, stored, '2025-03-08')
        members = [t for t in kif['tasks'] if t.get('effort_group') == 'BASE']
        self.assertTrue(members)
        self.assertTrue(all(t['understood'] == 'Yes' for t in members))
        self.assertTrue(all(t['basis']['understood']['by'] == 'human' for t in members))

    def test_plain_english_period_answer_resolves_delivery_outcome_and_date(self):
        snap, profile, facts = self.grouped_example()
        stored = ledger.Ledger(self.base / 'batch-delivery-review.json')
        stored.set_answer('review:Cycle|Client Expectation:example',
                          'Count as meet expectation. Expected date count as Sep 30.', 'human')
        kif, _ = classify.to_kif(snap, profile, {}, facts, stored, '2025-09-08')
        delivered = [t for t in kif['tasks'] if t.get('type') != 'Excluded' and t.get('delivered')]
        self.assertTrue(delivered)
        self.assertTrue(all(t['client_date'] == '2025-09-30' for t in delivered))
        self.assertTrue(all(t['met_client_date'] == 'Yes' for t in delivered))
        engine = kpi_engine.Engine(kif, profile, kpi._registry(self.ws()))
        self.assertEqual(engine.client_expectation(kif['periods'][0]).value, 100)

    def test_period_delivery_answer_replaces_pending_rows(self):
        snap, profile, facts = self.grouped_example()
        for item in snap['items']:
            item['section'] = 'Doing'
        stored = ledger.Ledger(self.base / 'pending-delivery-review.json')
        stored.set_answer('review:Cycle|Client Expectation:example',
                          'Count as meet expectation. Expected date count as Sep 30.', 'human')
        kif, _ = classify.to_kif(snap, profile, {}, facts, stored, '2025-09-08')
        included = [t for t in kif['tasks'] if t.get('type') != 'Excluded']
        self.assertTrue(included)
        self.assertTrue(all(t['met_client_date'] == 'Yes' for t in included))

    def test_rework_counts_each_close_to_reopen_cycle(self):
        item = {"id": "A", "key": "A", "title": "Feature", "created_at": "2025-03-01",
                "section": "Doing", "fields": {}, "comments": [], "events": [
                    {"kind": "section", "at": "2025-03-02", "from": "Doing", "to": "QA"},
                    {"kind": "section", "at": "2025-03-03", "from": "QA", "to": "Done"},
                    {"kind": "section", "at": "2025-03-04", "from": "Done", "to": "Doing"},
                    {"kind": "section", "at": "2025-03-05", "from": "Doing", "to": "Done"},
                    {"kind": "section", "at": "2025-03-06", "from": "Done", "to": "Doing"},
                ]}
        snap = {"capabilities": ["status_history"], "items": [item]}
        profile = {"workflow": {"delivered_when": {"values": ["QA"]},
                                "closed_when": {"values": ["Done"]},
                                "reopened_when": {"closed_values": ["Done"], "values": ["Doing"]}}}
        facts = {"periods": {"periods": [{"name": "Cycle"}]}}
        c = classify.Classifier(snap, profile, {}, facts, None, '2025-03-08')
        row = c.task_row(item, classify.Proposal('Task', 'planned work', 1), None)
        self.assertEqual((row['reopened'], row['reopen_count']), ('Yes', 2))
        kif = {"project": {}, "periods": [{"name": "Cycle"}], "tasks": [row], "defects": [],
               "generated": {"capabilities": ["status_history"]}}
        measure = kpi_engine.Engine(kif, profile, kpi._registry(self.ws())).rework_rate(kif['periods'][0])
        self.assertEqual((measure.numerator, measure.denominator, measure.value), (2, 1, 200))

    def test_group_member_defect_edits_do_not_change_its_delivery_classification(self):
        snap, profile, facts = self.grouped_example()
        kif, _ = classify.to_kif(snap, profile, {}, facts, None, '2025-03-08')
        registry = kpi._registry(self.ws())
        results = {'periods': [p.as_dict() for p in kpi_engine.Engine(kif, profile, registry).run()]}
        tabs = sheet_model.build(kif, results, {'profile': profile, 'registry': registry})
        state = sheet_readback.snapshot(tabs, {})
        defect_tab = next(t for t in tabs if t.name == 'Defect Register')
        col = next(i for i, key, _ in state['spec']['Defect Register']['fields'] if key == 'kind')
        defect_tab.cells[(state['spec']['Defect Register']['first'], col)]['v'] = 'Observation'
        grids = {t.name: sheet_readback._model_grid(t) for t in tabs}
        def grid(tab, row, col):
            return grids[tab](tab, row, col)
        grid.rows = lambda tab: grids[tab].rows(tab)
        stored = ledger.Ledger(self.base / 'dual-row-ledger.json')
        sheet_readback.fold(state, grid, stored, facts, {}, {i['id']: i for i in snap['items']})
        refreshed, _ = classify.to_kif(snap, profile, {}, facts, stored, '2025-03-08')
        self.assertEqual(next(t['type'] for t in refreshed['tasks'] if t['key'] == 'EXTRA-A'), 'CR')
        self.assertEqual(next(d['kind'] for d in refreshed['defects'] if d['key'] == 'EXTRA-A'), 'Observation')
        self.assertEqual(next(d['kind'] for d in refreshed['defects'] if d['key'] == 'EXTRA-B'), 'Bug')

    def test_group_missing_or_overlapping_members_are_not_silently_undercounted(self):
        snap, profile, facts = self.grouped_example()
        profile['conventions']['grouping']['linked_members']['EXTRA'].append('UNREAD')
        with self.assertRaisesRegex(ValueError, 'missing configured members'):
            classify.to_kif(snap, profile, {}, facts, None, '2025-03-08')
        profile['conventions']['grouping']['linked_members']['EXTRA'] = ['BASE-A', 'EXTRA-A']
        with self.assertRaisesRegex(ValueError, 'more than one'):
            classify.to_kif(snap, profile, {}, facts, None, '2025-03-08')

    def test_approved_source_overrides_generic_qa_exclusion(self):
        snap, profile, facts = self.grouped_example()
        profile['conventions']['grouping'] = {}
        kif, work = classify.to_kif(snap, profile, {}, facts, None, '2025-03-08')
        self.assertEqual(next(t['type'] for t in kif['tasks'] if t['key']=='EXTRA'), 'CR')
        self.assertFalse(any(q['id'].startswith('scope:') for q in work['questions']))

    def test_all_notes_require_review_and_changed_evidence_invalidates_it(self):
        snap, profile, facts = self.grouped_example()
        kif, work = classify.to_kif(snap, profile, {'name':'Example'}, facts, None, '2025-03-08')
        results = {'periods':[p.as_dict() for p in kpi_engine.Engine(kif,profile,kpi._registry(self.ws())).run()]}
        q = judge.build_queue(work, results, facts, self.base, ['Cycle'], kif, profile)
        self.assertEqual(len(q['notes']), 9)
        judge.write_queue(q, self.base)
        answers = self.base/'judge'/'answers.json'
        answers.write_text(json.dumps({'reasons':[{'period':n['period'], 'kpi':n['kpi'], 'why':'',
            'review_signature':n['review_signature']} for n in q['notes']]}))
        output = judge.apply_answers(answers, snap, ledger.Ledger(self.base/'review-ledger.json'), facts, ['Cycle'])
        self.assertEqual(output['refused'], [])
        self.assertEqual(judge.build_queue(work,results,facts,self.base,['Cycle'],kif,profile)['notes'], [])
        results['periods'][0]['measures'][0]['note_parts'].append('A newly recorded delivery changes the result.')
        self.assertEqual(len(judge.build_queue(work,results,facts,self.base,['Cycle'],kif,profile)['notes']), 1)

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        shutil.copytree(ROOT / 'examples' / 'northwind-board', self.base / 'fixture')
        self.profile = self.base / 'fixture' / 'profile.yaml'
        self.profile_data = profile_lib.load(self.profile)
        self.snap = json.loads((self.profile.parent / 'board.json').read_text())
        self.args = SimpleNamespace(profile=str(self.profile), project='northwind-q3', today='2026-09-18',
                                    date=None, offline=False, by='', no_publish=False)

    def ws(self):
        return kpi.Workspace(self.args)

    def test_rework_boundary_is_independent_of_qa_effort_completion(self):
        snap, profile, facts = self.grouped_example()
        snap['items'] = [snap['items'][0]]
        snap['items'][0]['subtasks'] = 0
        snap['items'][0]['events'] = [
            {'kind':'section','at':'2025-03-03','from':'Doing','to':'QA'},
            {'kind':'section','at':'2025-03-04','from':'QA','to':'Doing'}]
        snap['items'][0]['section'] = 'Doing'
        facts['estimates']['items'] = []
        profile['conventions']['grouping'] = {}
        profile['workflow']['reopened_when']['closed_values'] = ['QA','Done']
        kif, _ = classify.to_kif(snap, profile, {'name':'Example'}, facts, None, '2025-03-08')
        task = kif['tasks'][0]
        self.assertIsNone(task['closed'])
        self.assertEqual(task['rework_closed'], '2025-03-03')
        self.assertEqual(task['reopened'], 'Yes')
        engine = kpi_engine.Engine(kif, profile, kpi._registry(self.ws()))
        self.assertEqual(engine.velocity(kif['periods'][0]).value, 12)
        self.assertEqual(engine.rework_rate(kif['periods'][0]).value, 100)
        self.assert_grouped_formulas(kif, profile)

    def test_shared_effort_can_wait_for_client_handover(self):
        snap, profile, facts = self.grouped_example()
        profile['sources']['team_hours_when'] = 'handover'
        facts['periods']['periods'][0].update(handover_date=None, team_hours=9)
        kif, _ = classify.to_kif(snap, profile, {'name':'Example'}, facts, None, '2025-03-08')
        engine = kpi_engine.Engine(kif, profile, kpi._registry(self.ws()))
        self.assertEqual(engine.velocity(kif['periods'][0]).value, 28)
        self.assert_grouped_formulas(kif, profile)
        kif['periods'][0]['handover_date'] = '2025-03-08'
        self.assertEqual(engine.velocity(kif['periods'][0]).value, 37)

    def test_configured_ancestor_wins_over_overlapping_dates(self):
        snap, profile, facts = self.grouped_example()
        profile['conventions']['grouping'] = {}
        profile['periods'] = {'by_ancestor':[{'title':r'Milestone\s*1', 'period':'First'}]}
        snap['items'][0]['title'] = 'Bug reporting - Milestone 1'
        facts['periods']['periods'] = [{'name':name,'start':'2025-03-01','end':'2025-03-20'} for name in ['First','Second']]
        classifier = classify.Classifier(snap, profile, {'name':'Example'}, facts, None, '2025-03-08')
        result = classifier.period_of(snap['items'][1], '2025-03-04', None, report=True)
        self.assertEqual(result['value'], 'First')
        self.assertGreater(result['confidence'], .9)

    def test_review_diagnostics_do_not_remove_real_delivery_results(self):
        import note_policy
        good, review = note_policy.split('Three items were delivered. Individual histories are unavailable. || '
                                         'Five items lack individual discussion history. || '
                                         'Delivery cannot yet be confirmed. || The build was two days late.')
        self.assertEqual(good, ['Three items were delivered.', 'The build was two days late.'])
        self.assertEqual(len(review), 3)

    def test_edited_summary_survives_and_stale_summary_is_not_published(self):
        import note_policy
        snap, profile, facts = self.grouped_example()
        kif, _ = classify.to_kif(snap, profile, {'name':'Example'}, facts, None, '2025-03-08')
        engine = kpi_engine.Engine(kif, profile, kpi._registry(self.ws()))
        first = engine.run()[0].measures[0]
        tag = 'Cycle|' + first.name
        engine.reasons.update({'_summaries':{tag:'The delivery represents 28 estimated hours.'},
                               '_summary_bases':{tag:note_policy.basis(first.as_dict())}})
        result = engine.run()
        self.assertFalse(result[0].measures[0].publish_blocked)
        self.assertIn('represents 28',result[0].measures[0].note)
        kif['periods'][0]['team_hours'] = 4
        result = engine.run()
        self.assertTrue(result[0].measures[0].publish_blocked)
        payload = kpi_engine.to_pms_payloads(result, kif)[0]
        self.assertNotIn(first.name,[k['name'] for k in payload['kpis']])
        self.assertTrue(next(k for k in payload['skipped'] if k['name']==first.name)['preserve_existing'])

    def test_both_note_fields_are_read_back_including_intentional_blank(self):
        spec = {'kind':'summary','first':4,'last':4,'period':1,'kpi':3,'summary':12,'why':13,
                'manual_value':17,'manual_why':18,'note_bases':{'Cycle|Velocity':'basis'}}
        state = {'spec':{'KPI Summary':spec},'values':{'KPI Summary':{'rows':{'Cycle|Velocity':
                 {'summary':'Old summary','why':'Old context','value':None,'manual_why':None}}}}}
        cells = {1:'Cycle',3:'Velocity',12:'User summary',13:''}
        facts = {}
        sheet_readback.fold(state, lambda tab,r,c:cells.get(c), ledger.Ledger(self.base/'notes.json'),facts,{}, {})
        self.assertEqual(facts['reasons']['_summaries']['Cycle|Velocity'],'User summary')
        self.assertEqual(facts['reasons']['Cycle']['Velocity'],'')
        self.assertEqual(facts['reasons']['_summary_bases']['Cycle|Velocity'],'basis')

    def test_nested_project_registry_uses_global_ids_and_explicit_thresholds(self):
        default = json.loads(kpi_registry.DEFAULT.read_text())
        rows = [{'id':900,'threshold':None,'kpi':{'id':17,'name':'Defect Rate','threshold':20,
                 'isMinimumThreshold':False,'minValue':0,'maxValue':100}}]
        result = kpi_registry.merge(rows, default)
        metric = next(k for k in result['kpis'] if k['key']=='defect_rate')
        self.assertEqual(metric['pms_id'],17)
        self.assertEqual(metric['threshold'],20)
        self.assertEqual(kpi_registry.project_thresholds(rows)[0],{})
        rows[0]['threshold'] = 15
        self.assertEqual(kpi_registry.project_thresholds(rows)[0]['defect_rate']['threshold'],15)

    def test_pms_partial_update_clears_unknown_but_preserves_stale_edited_note(self):
        import os
        self.profile_data['output']['mode']='assisted-push'
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        payload=self.base/'payload.json'
        payload.write_text(json.dumps([{'projectId':101,'periodId':7,'name':'Cycle','description':'Cycle',
            'kpis':[{'name':'Velocity','kpiId':1,'value':10,'note':'Ten estimated hours.'}],
            'skipped':[{'name':'Defect Rate','kpiId':5,'reason':'Review estimates'},
                       {'name':'CR Rate','kpiId':9,'reason':'Review edited note','preserve_existing':True}]}]))
        before={'7':{'Defect Rate':{'value':10,'note':'Old'},'CR Rate':{'value':20,'note':'Keep'}}}
        after={'7':{'Velocity':{'value':10,'note':'Ten estimated hours.'},
                    'Defect Rate':{'value':None,'note':None},'CR Rate':before['7']['CR Rate']}}
        sent=[]
        def api(base,path,**kw):
            if kw.get('method')=='PUT': sent.append(kw); return {}
            if path.endswith('/periods'): return {'data':[{'id':7,'projectId':101,'name':'Cycle'}]}
            return {'data':{'id':7,'updatedAt':'2025-03-08T00:00:00Z'}}
        with patch.dict(os.environ,{'PMS_TOKEN':'test-only'}), patch.object(pms_push,'read_current',side_effect=[before,after]), patch.object(pms_push,'_api',side_effect=api):
            result=pms_push.main(['--profile',str(self.profile),'--payloads',str(payload),'--apply'])
        self.assertEqual(result,0)
        self.assertNotIn('kpis',sent[0]['body'])
        self.assertEqual(sent[0]['body']['periodKpis'][-1],{'kpiId':5,'value':None,'note':None})
        self.assertNotIn(9,[k['kpiId'] for k in sent[0]['body']['periodKpis']])
        self.assertEqual(sent[0]['extra']['Last-Modified'],'2025-03-08T00:00:00Z')

    def test_host_transport_rejects_credentials_before_writing_request(self):
        import os, time, host_transport
        (self.base/'session.json').write_text(json.dumps({'services':['pms'],'expires_at':time.time()+60}))
        with patch.dict(os.environ,{'KPI_HOST_BRIDGE':str(self.base)}):
            with self.assertRaises(host_transport.TransportError):
                host_transport.call('pms','GET','https://pms.example/api/projects',headers={'Cookie':'secret'})
        self.assertEqual(list(self.base.glob('*.request.json')),[])

    def test_optional_connected_host_javascript_guards(self):
        import subprocess
        node=shutil.which('node')
        if not node:
            self.skipTest('Node is optional for non-host command-line runs')
        result=subprocess.run([node,str(ROOT/'tests/test_host_transport.mjs')],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)

    def test_configured_google_sheet_is_authority_even_with_local_baseline(self):
        ws=self.ws()
        ws.profile['output'].update(workbook='google-sheets',workbook_file='fictional-remote-sheet-1234567890')
        state={'destination':{'kind':'local','path':str(self.base/'stale.xlsx')},'values':{},'spec':{}}
        with patch.object(sheet_readback,'load_state',return_value=state), patch.object(kpi.G,'how_signed_in',return_value='host'), \
             patch.object(sheet_google,'read_grid',return_value='fresh-grid') as remote, \
             patch.object(sheet_readback,'fold',return_value=['latest notes']) as fold:
            self.assertEqual(kpi.read_back(ws,{},self.args),['latest notes'])
        remote.assert_called_once_with('fictional-remote-sheet-1234567890',[])
        self.assertEqual(fold.call_args.args[1],'fresh-grid')

    def test_unreadable_google_never_falls_back_to_local_for_push(self):
        ws=self.ws()
        state={'destination':{'kind':'google','id':'remote'},'values':{},'spec':{}}
        with patch.object(sheet_readback,'load_state',return_value=state), patch.object(kpi.G,'how_signed_in',return_value='host'), \
             patch.object(sheet_google,'read_grid',side_effect=RuntimeError('unavailable')), patch.object(sheet_xlsx,'read_grid') as local:
            with self.assertRaisesRegex(SystemExit,'Nothing will overwrite'):
                kpi.read_back(ws,{},self.args)
        local.assert_not_called()

    def test_authoritative_source_delivery_survives_a_feedback_status(self):
        snap, profile, facts=self.grouped_example()
        snap['items']=[snap['items'][0]]
        snap['items'][0].update(events=[],section='Feedback',subtasks=0)
        profile['conventions']['grouping']={}
        facts['estimates']['items']=[]
        facts['plan']['items'][0].update(delivered='2025-03-06',delivery_evidence='Recorded handover')
        kif,_=classify.to_kif(snap,profile,{'name':'Example'},facts,None,'2025-03-08')
        self.assertEqual(kif['tasks'][0]['delivered'],'2025-03-06')
        self.assertIsNone(kif['tasks'][0]['closed'])
        self.assertEqual(kif['tasks'][0]['delivery_evidence'],'Recorded handover')

    def test_exact_source_match_cannot_be_reused_by_a_similar_card(self):
        snap, profile, facts=self.grouped_example()
        facts['plan']['items']=[{'title':'Original feature','dev_hours':12,'qa_hours':4}]
        snap['items'][1]['title']='[QA] Original feature'
        classifier=classify.Classifier(snap,profile,{'name':'Example'},facts,None,'2025-03-08')
        self.assertEqual(classifier.match_scope(snap['items'][0])[0]['title'],'Original feature')
        self.assertIsNone(classifier.match_scope(snap['items'][1])[0])

    def test_google_creates_all_tabs_before_writing_cross_tab_formulas(self):
        dashboard, config = sheet_model.Tab('Dashboard'), sheet_model.Tab('Config')
        dashboard.put(1, 1, f="='Config'!A1")
        config.put(1, 1, 'Demo')
        sent = []
        class Session:
            identity = 'test'
            def call(self, method, url, body=None, params=None):
                if method == 'GET':
                    return {'properties': {'locale': 'en_US'}, 'sheets': []}
                sent.extend(body['requests'])
                return {}
        sheet_google.publish([dashboard, config], {}, 'Demo', known_id='fictional-sheet', session=Session())
        created = [i for i, r in enumerate(sent) if 'addSheet' in r]
        writes = [i for i, r in enumerate(sent) if 'updateCells' in r]
        self.assertEqual(len(created), 2)
        self.assertLess(max(created), min(writes))

    def test_asana_custom_key_matches_plan_and_overrides_title_reference(self):
        api = adapter('asana')
        task = {'gid': '1', 'name': 'Related to DEMO-9',
                'custom_fields': [{'name': 'Ticket', 'display_value': 'DEMO-42'}]}
        raw = {'project': {'gid': '77'}, 'tasks': [task], 'stories': {'1': []}}
        snap = api.from_raw(raw, r'DEMO-\d+', key_field='Ticket')
        self.assertEqual(snap['items'][0]['key'], 'DEMO-42')
        self.assertEqual(api.item_from_task(task, '77', r'DEMO-\d+')['key'], 'DEMO-9')
        c = classify.Classifier(snap, {}, {}, {'plan': {'items': [
            {'board_key': 'DEMO-42', 'title': 'A different planning title', 'dev_hours': 4}
        ]}}, None, '2026-09-18')
        row, score = c.match_scope(snap['items'][0])
        self.assertEqual((row['dev_hours'], score), (4, 1.0))

    def test_asana_incomplete_history_is_refused_and_disabled_comments_not_claimed(self):
        api = adapter('asana')
        raw = {'project': {'gid': '77'}, 'tasks': [{'gid': '1', 'name': 'Demo'}]}
        with self.assertRaisesRegex(board.ReaderError, 'missing task history'):
            api.from_raw(raw)
        raw['stories'] = {'1': []}
        self.assertNotIn('comments', api.from_raw(raw, comments='never')['capabilities'])

    def test_wrong_project_export_preserves_previous_board(self):
        ws = self.ws()
        ws.profile['tracker'] = {'adapter': 'asana'}
        ws.project['tracker_ref'] = '77'
        cache = ws.dir / 'cache' / 'board.json'
        board.save(cache, self.snap)
        before = cache.read_bytes()
        raw = self.base / 'other-board.json'
        raw.write_text(json.dumps({'project': {'gid': '88'}, 'tasks': [], 'stories': {}}))
        args = SimpleNamespace(board=None, adapter=None, offline=False, from_raw=str(raw))
        with self.assertRaisesRegex(SystemExit, 'different tracker project'):
            kpi.get_board(ws, args, lambda *_: None)
        self.assertEqual(cache.read_bytes(), before)

    def test_source_mapping_filters_unapproved_rows_and_translates_periods(self):
        self.assertEqual(str(sources._cell(10.0)), '10')
        path = self.base / 'estimates.csv'
        path.write_text('Item,Hours,Approval,Batch\nLogin,5,Approved,10\nSearch,8,Draft,99\n')
        spec = {'columns': {'title': 'Item', 'dev_hours': 'Hours', 'approval': 'Approval', 'period': 'Batch'},
                'where': {'approval': ['Approved']}, 'values': {'period': {10: 'Cycle A'}}}
        rows, problem = sources.table(path, spec)
        self.assertFalse(problem)
        self.assertEqual([(r['title'], r['period']) for r in rows], [('Login', 'Cycle A')])
        path.write_text(path.read_text() + 'Export,3,Approved,20\n')
        rows, problem = sources.table(path, spec)
        self.assertEqual(rows, [])
        self.assertIn('unmapped value', problem)
        spec['where']['missing_column'] = 'Yes'
        self.assertIn('needs a column mapping', sources.table(path, spec)[1])

    def test_explicit_included_key_overrides_admin_pattern_without_including_others(self):
        profile = {'conventions': {'exclude_patterns': [r'^\[QA\]']},
                   'custom_instructions': {'rule_overrides': [
                       {'rule': 'include_key', 'value': 'DEMO-42 = CR', 'why': 'Approved testing deliverable'}]}}
        c = classify.Classifier({}, profile, {}, {}, None, '2026-09-18')
        self.assertEqual(c.nature({'title': '[QA] Release testing', 'key': 'DEMO-42'})[0]['value'], 'CR')
        self.assertEqual(c.nature({'title': '[QA] Routine checklist', 'key': 'DEMO-43'})[0]['value'], 'Excluded')

    def test_assignee_allowlist_excludes_other_and_unassigned_cards_before_counting(self):
        profile = {'conventions': {'defect_by': 'issue-type', 'defect_values': ['Bug'],
                                   'assignee_include': ['Alex Example', 'Sam Example']}}
        facts = {'periods': {'periods': [{'name': 'Cycle A'}]}}
        items = [
            {'id':'1','key':'DEMO-1','title':'Included work','assignee':'alex example'},
            {'id':'2','key':'DEMO-2','title':'Other work','assignee':'Client Person'},
            {'id':'3','key':'DEMO-3','title':'Unassigned work','assignee':None},
            {'id':'4','key':'DEMO-4','title':'Outside report','assignee':'Client Person',
             'fields':{'Type':'Bug'}},
            {'id':'5','key':'DEMO-5','title':'Included report','assignee':'Sam Example',
             'fields':{'Type':'Bug'}},
        ]
        stored = ledger.Ledger(self.base / 'assignee-ledger.json')
        stored.set('2','nature','Task','human','Previously included by hand')
        stored.set('2','set:type','Task','human','Previously edited in the sheet')
        work = classify.Classifier({'items':items}, profile, {}, facts, stored, '2026-09-18').run()
        self.assertEqual([r['key'] for r in work['defects']], ['DEMO-5'])
        by_key = {r['key']:r for r in work['tasks']}
        self.assertEqual(by_key['DEMO-1']['type'], 'Task')
        for key in ('DEMO-2','DEMO-3','DEMO-4'):
            self.assertEqual(by_key[key]['type'], 'Excluded')
            self.assertIn('configured project team', by_key[key]['exclude_reason'])
        _, run_work = classify.to_kif({'items':items}, profile, {}, facts, stored, '2026-09-18')
        self.assertEqual(run_work['counts']['excluded_assignees'], 3)
        tabs = {t.name:t for t in sheet_model.build(
            {'project':{},'periods':[{'name':'Cycle A'}],'tasks':work['tasks'],'defects':work['defects']},
            {'periods':[]}, {'profile':profile})}
        self.assertTrue(any(c.get('v') == 'Alex Example, Sam Example' for c in tabs['Config'].cells.values()))

    def test_assignee_scope_does_not_count_grouped_outside_defects_or_shared_hours(self):
        snap, profile, facts = self.grouped_example()
        profile['conventions']['assignee_include'] = ['Alex Example']
        for item in snap['items']:
            item['assignee'] = 'Alex Example'
        snap['items'][-1]['assignee'] = 'Outside Example'
        kif, _ = classify.to_kif(snap, profile, {}, facts, None, '2025-03-08')
        self.assertNotIn('EXTRA-B', [d['key'] for d in kif['defects']])
        outside = next(t for t in kif['tasks'] if t['key'] == 'EXTRA-B')
        self.assertEqual(outside['type'], 'Excluded')
        parent = next(t for t in kif['tasks'] if t['key'] == 'EXTRA')
        self.assertFalse(parent['effort_only'])
        engine = kpi_engine.Engine(kif, profile, kpi._registry(self.ws()))
        self.assertIsNone(engine.velocity(kif['periods'][0]).value)
        self.assert_grouped_formulas(kif, profile)

    def test_unassigned_group_parent_keeps_budget_for_included_members(self):
        snap, profile, facts = self.grouped_example()
        profile['conventions']['assignee_include'] = ['Alex Example']
        for item in snap['items']:
            item['assignee'] = None if item['key'] in ('BASE', 'EXTRA') else 'Alex Example'
        kif, _ = classify.to_kif(snap, profile, {}, facts, None, '2025-03-08')
        engine = kpi_engine.Engine(kif, profile, kpi._registry(self.ws()))
        self.assertEqual(engine.velocity(kif['periods'][0]).value, 28)
        self.assertEqual(len(engine.deliverables_in('Cycle')), 4)

    def test_judge_context_keeps_early_clarification_and_recent_replies(self):
        comments = [{'at': '2026-07-01', 'text': 'Please confirm which account type is intended.', 'url': 'demo:1'}]
        comments += [{'at': '2026-09-01', 'text': 'Progress update', 'url': f'demo:{i}'} for i in range(2, 12)]
        item = {'id': 'a', 'title': 'Demo', 'comments': comments, 'subtasks': 3}
        c = classify.Classifier({'items': [item]}, {}, {}, {}, None, '2026-09-18')
        c.queue = [{'item_id': 'a', 'field': 'understood', '_context': 'comments'}]
        c._attach_context()
        ctx = c.queue[0]['item']
        self.assertEqual(ctx['comments'][0]['url'], 'demo:1')
        self.assertEqual(ctx['comments'][-1]['url'], 'demo:11')
        self.assertEqual(ctx['comments_omitted'], 4)
        self.assertEqual(ctx['subtasks'], 3)

    def test_approved_estimated_bug_fix_is_an_addition_with_its_hours(self):
        c = classify.Classifier({}, {'conventions': {'defect_pattern': r'Bug \d+'}}, {},
                                {'estimates': {'items': [{'board_key': 'DEMO-2', 'dev_hours': 5}]}},
                                None, '2026-09-18')
        nature, row = c.nature({'title': 'Bug 2: approved additional fix', 'key': 'DEMO-2'})
        self.assertEqual(nature['value'], 'CR')
        self.assertEqual(row['dev_hours'], 5)

    def test_rejection_discussion_requires_judgement_and_retains_early_evidence(self):
        for text in ('Closed as a non-issue.', 'This is not a non-issue; the bug needs fixing.'):
            item = {'id': 'a', 'title': 'Bug 4: Missing result', 'created_at': '2026-09-01',
                    'comments': [{'at': '2026-09-02', 'text': text, 'url': 'demo:decision'}] +
                                [{'at': '2026-09-15', 'text': 'Update'} for _ in range(8)]}
            c = classify.Classifier({'items': [item]}, {}, {},
                                    {'periods': {'periods': [{'name': 'Cycle A'}]}}, None, '2026-09-18')
            c.defect_row(item, classify.Proposal('Bug', 'Report', 1))
            c._attach_context()
            question = next(q for q in c.queue if q['field'] == 'rejected')
            self.assertLess(question['confidence'], classify.SURE)
            self.assertEqual(question['item']['comments'][0]['url'], 'demo:decision')

    def test_scope_matching_never_fuzzes_conflicting_identifiers(self):
        self.assertEqual(classify._similar('Bug 9: Login validation fails', 'Bug 19: Login validation fails'), 0)
        c = classify.Classifier({}, {}, {}, {'plan': {'items': [
            {'board_key': 'DEMO-1', 'title': 'Identical title', 'dev_hours': 5}]}}, None, '2026-09-18')
        self.assertIsNone(c.match_scope({'id': '2', 'key': 'DEMO-2', 'title': 'Identical title'})[0])

    def test_separate_estimates_on_one_delivery_card_are_counted_once_each(self):
        item = {'id': 'a', 'key': 'DEMO-2', 'title': 'Combined delivery', 'created_at': '2026-09-01'}
        facts = {'periods': {'periods': [{'name': 'Cycle A', 'start': '2026-09-01', 'end': '2026-09-30'}]},
                 'estimates': {'items': [
                     {'board_key': 'DEMO-2', 'title': 'Login', 'dev_hours': 3, 'period': 'Cycle A'},
                     {'board_key': 'DEMO-2', 'title': 'Search', 'dev_hours': 7, 'period': 'Cycle A'}]}}
        c = classify.Classifier({'items': [item]}, {}, {}, facts, None, '2026-09-18')
        work = c.run()
        self.assertEqual([(r['title'], r['hours_dev']) for r in work['tasks']], [('Login', 3), ('Search', 7)])
        self.assertEqual(len({r['_row'] for r in work['tasks']}), 2)
        self.assertFalse(any(q['id'].startswith('scope:') for q in c.questions))
        stored = ledger.Ledger(self.base / 'estimate-ledger.json')
        rows = {r['_row']: {'key': r['key'], 'title': r['title'], 'hours_dev': r['hours_dev'],
                           'item': r['_row'], 'understood': r['understood']} for r in work['tasks']}
        first, second = work['tasks']
        fields = [(1, 'key', 'in'), (2, 'title', 'in'), (3, 'hours_dev', 'in'),
                  (4, 'item', 'meta'), (5, 'understood', 'in')]
        state = {'spec': {'Task Register': {'kind': 'rows', 'key': 'key', 'first': 1, 'fields': fields}},
                 'values': {'Task Register': {'rows': rows}}}
        cells = [[first['key'], first['title'], 9, first['_row'], 'No'],
                 [second['key'], second['title'], 7, second['_row'], None]]
        def grid(tab, r, col):
            return cells[r - 1][col - 1]
        grid.rows = lambda tab: 2
        sheet_readback.fold(state, grid, stored, facts, {}, {'a': item})
        facts['estimates']['items'].reverse()
        refreshed = classify.Classifier({'items': [item]}, {}, {}, facts, stored, '2026-09-18').run()['tasks']
        self.assertEqual({r['title']: r['hours_dev'] for r in refreshed}, {'Login': 9, 'Search': 7})
        self.assertEqual(next(r for r in refreshed if r['title'] == 'Login')['understood'], 'No')
        self.assertIsNone(next(r for r in refreshed if r['title'] == 'Search')['understood'])
        facts['estimates']['items'].append(copy.deepcopy(facts['estimates']['items'][0]))
        with self.assertRaisesRegex(ValueError, 'distinct titles'):
            classify.Classifier({'items': [item]}, {}, {}, facts, stored, '2026-09-18').run()

    def test_missing_explicit_source_key_stays_in_scope_and_asks_for_delivery(self):
        facts = {'plan': {'items': [{'board_key': 'DEMO-404', 'title': 'Missing card', 'dev_hours': 4}]}}
        c = classify.Classifier({'items': []}, {}, {}, facts, None, '2026-09-18')
        work = c.run()
        self.assertEqual(len(work['tasks']), 1)
        self.assertEqual(work['tasks'][0]['hours_dev'], 4)
        self.assertIsNone(work['tasks'][0]['delivered'])
        self.assertTrue(any(q['id'].startswith('scope:') for q in c.questions))

    def test_asana_expands_descendants_once_even_when_already_on_project(self):
        api = adapter('asana')
        rows = {'a': [{'gid': 'b', 'parent': {'gid': 'a'}, 'num_subtasks': 1}],
                'b': [{'gid': 'c', 'parent': {'gid': 'b'}, 'num_subtasks': 0}]}
        calls = []
        class Client:
            def pages(self, path, params):
                calls.append(path)
                return rows[path.split('/')[2]]
        got = api.expand_subtasks(Client(), [{'gid': 'a', 'num_subtasks': 1},
                                             {'gid': 'b', 'num_subtasks': 1}], workers=1)
        self.assertEqual({t['gid'] for t in got}, {'a', 'b', 'c'})
        self.assertEqual(sorted(calls), ['/tasks/a/subtasks', '/tasks/b/subtasks'])

    def test_asana_linked_tasks_are_explicit_deduplicated_and_bounded(self):
        api = adapter('asana')
        calls = []
        class Client:
            def get(self, path, params):
                calls.append(path)
                return {'data': {'gid': path.rsplit('/', 1)[1], 'name': 'External deliverable'}}
        got = api.add_linked_tasks(Client(), [{'gid': '1'}], ['1', '2', '2'])
        self.assertEqual(calls, ['/tasks/2'])
        self.assertEqual({t['gid'] for t in got}, {'1', '2'})
        with self.assertRaises(api.AsanaError):
            api.add_linked_tasks(Client(), [], ['../projects/other'])
        task = {'gid': '2', 'memberships': [{'project': {'gid': '88'}, 'section': {'name': 'Closed'}}]}
        self.assertIsNone(api.item_from_task(task, '77', None)['section'])
        self.assertEqual(api.item_from_task(task, '77', None, linked=True)['section'], 'Closed')
        task['memberships'].append({'project': {'gid': '99'}, 'section': {'name': 'In Progress'}})
        self.assertIsNone(api.item_from_task(task, '77', None, linked=True)['section'])

    def test_source_override_is_scoped_and_requires_a_reason(self):
        path = self.base / 'estimates.csv'
        path.write_text('Item,Approval,Batch\nLogin,Pending,10\nSearch,Pending,20\n')
        spec = {'columns': {'title': 'Item', 'approval': 'Approval', 'period': 'Batch'},
                'where': {'approval': ['Approved']},
                'overrides': [{'match': {'period': 10}, 'set': {'approval': 'Approved', 'board_key': 'DEMO-2'},
                               'why': 'Lead confirmed this batch'}]}
        rows, problem = sources.table(path, spec)
        self.assertFalse(problem)
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]['title'], rows[0]['board_key']), ('Login', 'DEMO-2'))
        self.assertEqual(rows[0]['mapping_reason'], 'Lead confirmed this batch')
        spec['overrides'][0].pop('why')
        self.assertIn('needs match fields and a reason', sources.table(path, spec)[1])

    def test_configured_registry_is_used_by_engine_and_sheet(self):
        reg = json.loads(kpi_registry.DEFAULT.read_text())
        reg['source'] = 'pms'
        reg['projects'] = {'101': {'defect_rate': {'threshold': 87}}}
        self.profile_data['organization']['kpi_registry'] = 'targets/custom.json'
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        target = self.profile.parent / 'targets' / 'custom.json'
        target.parent.mkdir()
        target.write_text(json.dumps(reg))
        ws = self.ws()
        kif, _ = classify.to_kif(self.snap, ws.profile, ws.project, ws.facts, ws.ledger, ws.today)
        ws.run_dir.mkdir(parents=True)
        path = ws.run_dir / 'run.kif.json'
        path.write_text(json.dumps(kif))
        results = kpi.compute(ws, path)
        m = next(m for m in results['periods'][0]['measures'] if m['key'] == 'defect_rate')
        self.assertEqual(m['threshold'], 87)
        self.assertEqual(kpi._registry(ws), reg)
        tabs = sheet_model.build(kif, results, {'profile': ws.profile, 'registry': reg})
        config = next(t for t in tabs if t.name == 'Config')
        row = next(r for (r,c), cell in config.cells.items() if c == 2 and cell.get('v') == 'Defect Rate')
        self.assertEqual(config.cells[row,6]['v'], 87)
        self.assertTrue(config.cells[row,7]['v'])

    def test_profile_override_type_fails_before_workspace_created(self):
        self.profile_data['projects'][0]['overrides'] = {'policy': {'count_observations': 'false'}}
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        with self.assertRaises(SystemExit):
            self.ws()
        self.assertFalse((self.profile.parent / 'northwind-q3' / 'ledger.json').exists())

    def test_custom_registry_missing_does_not_silently_fallback(self):
        with self.assertRaises(SystemExit):
            kpi_registry.resolve_path(self.profile, {'organization': {'kpi_registry': 'missing.json'}})

    def test_corrupt_memory_and_baseline_are_preserved(self):
        for filename, reader in [('ledger.json', ledger.Ledger), ('sheet_state.json', sheet_readback.load_state)]:
            p = self.base / filename
            p.write_text('{broken')
            with self.assertRaises(SystemExit):
                reader(p)
            self.assertEqual(p.read_text(), '{broken')

    def test_failed_readback_stops_before_overwrite(self):
        ws = self.ws()
        book = ws.dir / 'review.xlsx'
        book.write_bytes(b'not a zip')
        state = {'destination': {'kind': 'xlsx', 'path': str(book)}, 'values': {'Config': {}}}
        sheet_readback.save_state(ws.dir / 'sheet_state.json', state)
        with self.assertRaisesRegex(SystemExit, 'Nothing will overwrite'):
            kpi.read_back(ws, {}, self.args)
        self.assertEqual(book.read_bytes(), b'not a zip')

    def test_remote_offline_preview_keeps_original_baseline_and_file(self):
        ws = self.ws()
        book = ws.dir / 'review.xlsx'
        book.write_bytes(b'keep local edits too')
        state = {'destination': {'kind': 'google', 'id': 'sheet-123', 'path': str(book)}, 'values': {'Config': {}}}
        path = ws.dir / 'sheet_state.json'
        sheet_readback.save_state(path, state)
        before = path.read_bytes()
        self.args.offline = True
        self.assertTrue(kpi.read_back(ws, {}, self.args))
        tab = sheet_model.Tab('Preview')
        tab.put(1, 1, 'preview')
        with patch.object(sheet_model, 'build', return_value=[tab]):
            dest, _ = kpi.publish(ws, {}, {}, {}, self.args)
        self.assertIn('Preview - ', dest['path'])
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(book.read_bytes(), b'keep local edits too')

    def test_failed_google_publish_preserves_destination(self):
        ws = self.ws()
        ws.profile['output'] = {'workbook': 'google-sheets'}
        state = {'destination': {'kind': 'google', 'id': 'sheet-123'}, 'values': {}}
        ws.previous_sheet_state = state
        path = ws.dir / 'sheet_state.json'
        sheet_readback.save_state(path, state)
        before = path.read_bytes()
        tab = sheet_model.Tab('Preview'); tab.put(1, 1, 'value')
        with patch.object(sheet_model, 'build', return_value=[tab]), \
             patch.object(kpi.G, 'how_signed_in', return_value='test'), \
             patch.object(sheet_google, 'publish', side_effect=kpi.G.GoogleError('unavailable')):
            _, notes = kpi.publish(ws, {}, {}, {}, self.args)
        self.assertIn('not updated', notes[0])
        self.assertEqual(path.read_bytes(), before)

    def test_local_writer_preserves_custom_tabs_and_literal_text(self):
        path = self.base / 'review.xlsx'
        wb = Workbook(); wb.active.title = 'My notes'; wb.active['A1'] = 'Keep this'; wb.save(path)
        tab = sheet_model.Tab('Task Register'); tab.put(1, 1, '=HYPERLINK("bad","bad")')
        sheet_xlsx.write([tab], path)
        wb = load_workbook(path)
        self.assertEqual(wb['My notes']['A1'].value, 'Keep this')
        self.assertEqual(wb['Task Register']['A1'].data_type, 's')

    def test_google_readback_is_not_cut_off_at_2000_rows(self):
        class Session:
            def call(self, *args, **kw):
                self.ranges = kw['params']['ranges']
                return {'valueRanges': [{'values': [['x']] * 2101}]}
        s = Session(); grid = sheet_google.read_grid('id', ["Lead's tasks"], session=s)
        self.assertEqual(grid.rows("Lead's tasks"), 2101)
        self.assertEqual(s.ranges, ["'Lead''s tasks'"])

    def test_missing_review_tab_is_not_treated_as_cleared_cells(self):
        ws = self.ws()
        grid = lambda *args: None
        grid.rows = lambda tab: 0
        with self.assertRaisesRegex(ValueError, 'missing or empty'):
            sheet_readback.fold({'values': {'Config': {}}, 'spec': {'Config': {'kind':'cells','cells':{}}}},
                                grid, ws.ledger, ws.facts, {}, {})

    def test_changed_comment_invalidates_judgement(self):
        item = copy.deepcopy(self.snap['items'][0]); item['comments'] = [{'text':'no clarification', 'at':'2026-09-01'}]
        first = board.fingerprint(item)
        item['comments'][0]['text'] = 'client explained the requirement'
        self.assertNotEqual(first, board.fingerprint(item))

    def test_undecidable_answer_becomes_persistent_person_question(self):
        ws = self.ws()
        c = classify.Classifier(self.snap, ws.profile, ws.project, ws.facts, ws.ledger, ws.today)
        item = self.snap['items'][0]
        path = self.base / 'answers.json'
        path.write_text(json.dumps({'answers':[{'item_id':item['id'], 'field':'understood', 'value':None,
                                                'why':'Client clarification is not recorded.'}]}))
        judge.apply_answers(path, self.snap, ws.ledger, ws.facts, c.period_names)
        c.settle(item, 'understood', classify.Proposal('Yes','uncertain',0.5), 'Was it understood?', ['Yes','No'])
        self.assertFalse(c.queue)
        self.assertEqual(len(c.questions), 1)
        sheet_readback._file_answer(c.questions[0]['id'], 'No', ws.ledger, ws.facts)
        self.assertEqual(ws.ledger.get(item['id'], 'understood')['value'], 'No')
        self.assertIsNone(ws.ledger.get(item['id'], 'deferred:understood'))

    def test_plain_english_understood_answer_is_kept_as_the_yes_choice(self):
        ws = self.ws()
        item = self.snap['items'][0]
        sheet_readback._file_answer(
            f"judge:{item['id']}:understood", "It's fine, count them as understood.", ws.ledger, ws.facts
        )
        self.assertEqual(ws.ledger.get(item['id'], 'understood')['value'], 'Yes')

    def test_ambiguous_plain_english_judgement_is_still_refused(self):
        ws = self.ws()
        item = self.snap['items'][0]
        with self.assertRaisesRegex(ValueError, 'choose one of Yes, No'):
            sheet_readback._file_answer(
                f"judge:{item['id']}:understood", "Please review this again.", ws.ledger, ws.facts
            )

    def test_stale_digest_is_withheld_but_not_deleted(self):
        facts = {'plan': {'source': {'fingerprint':'old'}, 'items':[{'title':'Scope'}]}}
        active, problems = sources.usable_facts([{'role':'plan','state':'fresh','path':'plan.pdf','fingerprint':'new'}], facts)
        self.assertEqual(active['plan'], {})
        self.assertTrue(facts['plan']['items'])
        self.assertTrue(problems)

    def test_changed_source_does_not_reuse_other_source_cache(self):
        ws = self.ws()
        source = self.base / 'old.pdf'; source.write_bytes(b'old source')
        sources.pull({'sources':{'plan':{'kind':'pdf','ref':str(source)}}},ws.dir,self.base,offline=True)
        result = sources.pull({'sources':{'plan':{'kind':'pdf','ref':'new-missing.pdf'}}},ws.dir,self.base,offline=True)
        self.assertEqual(result[0]['state'],'missing')
        self.assertIsNone(result[0]['path'])

    def test_mapping_never_falls_back_to_another_tab(self):
        path = self.base / 'data.xlsx'; wb = Workbook(); wb.active.title='Other project'
        wb.active.append(['Item','Hours']); wb.active.append(['Secret scope',8]); wb.save(path)
        rows, problem = sources.table(path, {'tab':'Expected project','columns':{'title':'Item','dev_hours':'Hours'}})
        self.assertEqual(rows,[]); self.assertIn('missing',problem)

    def test_empty_mapped_table_removes_old_rows(self):
        path = self.base / 'data.csv'; path.write_text('Item,Hours\n')
        facts = {'estimates':{'items':[{'title':'removed'}]}}
        sources.facts_from_mapping({'role':'estimates','path':str(path),'fingerprint':'new'},
                                   {'map':{'columns':{'title':'Item','dev_hours':'Hours'}}},facts)
        self.assertEqual(facts['estimates']['items'],[])

    def test_velocity_missing_is_not_zero_and_zero_is_valid(self):
        ws = self.ws()
        kif, _ = classify.to_kif(self.snap, ws.profile, ws.project, ws.facts, ws.ledger, ws.today)
        t = next(t for t in kif['tasks'] if t.get('delivered') and t['type'] != 'Excluded')
        kif['tasks']=[t]; kif['project']['velocity_unit']='Story Points'
        per = next(p for p in kif['periods'] if p['name']==t['period'])
        reg = kpi._registry(ws)
        t['story_points']=None
        self.assertIsNone(kpi_engine.Engine(copy.deepcopy(kif),ws.profile,reg).velocity(per).value)
        t['story_points']=0
        self.assertEqual(kpi_engine.Engine(copy.deepcopy(kif),ws.profile,reg).velocity(per).value,0)

    def test_long_period_names_remain_valid_join_keys(self):
        ws=self.ws(); facts=copy.deepcopy(ws.facts)
        old=facts['periods']['periods'][0]['name']; new=old+' with a long descriptive name'
        facts['periods']['periods'][0]['name']=new
        for role in ('plan','estimates'):
            for item in facts.get(role,{}).get('items',[]):
                if item.get('period')==old: item['period']=new
        kif,_=classify.to_kif(self.snap,ws.profile,ws.project,facts,ws.ledger,ws.today)
        self.assertIn(new,[p['name'] for p in kif['periods']])
        self.assertTrue(all(t['period'] in [p['name'] for p in kif['periods']] for t in kif['tasks']))

    def test_formula_ranges_cover_large_boards(self):
        ws=self.ws(); kif,_=classify.to_kif(self.snap,ws.profile,ws.project,ws.facts,ws.ledger,ws.today)
        task=copy.deepcopy(kif['tasks'][0]); kif['tasks']=[dict(task,key=f'T-{i}') for i in range(1100)]
        engine=kpi_engine.Engine(kif,ws.profile,kpi._registry(ws))
        results={'periods':[r.as_dict() for r in engine.run()]}
        tabs=sheet_model.build(kif,results,{'profile':ws.profile})
        sm=next(t for t in tabs if t.name=='KPI Summary')
        self.assertIn(f'${1100 + sheet_model.SPARE + 4}', sm.cells[4,5]['f'])

    def test_jira_uses_fully_paginated_history(self):
        jira=adapter('jira')
        issue={'id':'1','key':'DEMO-1','fields':{},'changelog':{'histories':[]},'_histories':[
            {'created':'2026-09-01','items':[{'field':'status','fromString':'Open','toString':'Closed'}]}]}
        issue['changelog']['histories']=[{'created':'2026-08-01','items':[]}]
        it=jira.item_from_issue(issue,'https://example.atlassian.net',None,'Points','Hours')
        self.assertEqual(it['events'][0]['to'],'Closed')

    def test_github_paginates_comments_and_history(self):
        gh=adapter('github')
        node={'id':'I_1','comments':{'nodes':[{'bodyText':'first'}],'pageInfo':{'hasNextPage':True,'endCursor':'c'}},
              'timelineItems':{'nodes':[],'pageInfo':{'hasNextPage':True,'endCursor':'t'}}}
        class Client:
            def query(self,q,v):
                field='comments' if v['after']=='c' else 'timelineItems'
                return {'node':{field:{'nodes':[{'bodyText':'last'}] if field=='comments' else [{'__typename':'ReopenedEvent'}],
                                      'pageInfo':{'hasNextPage':False}}}}
        got=gh.complete_issue(Client(),node,gh._FULL)
        self.assertEqual(len(got['comments']['nodes']),2)
        self.assertEqual(got['timelineItems']['nodes'][0]['__typename'],'ReopenedEvent')

    def test_pms_wrong_project_is_refused_before_network(self):
        self.profile_data['output']['mode']='assisted-push'
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        payload=self.base/'payload.json'; payload.write_text(json.dumps([{'projectId':999,'periodId':1,'name':'Cycle','kpis':[]}]))
        with patch.object(pms_push,'_api') as api:
            result=pms_push.main(['--profile',str(self.profile),'--payloads',str(payload),'--apply'])
        self.assertEqual(result,2); api.assert_not_called()

    def test_pms_array_response_and_string_period_ids(self):
        raw=[{'name':'Velocity','periods':[{'periodId':7,'periodKpiValue':{'value':0,'note':'zero'}}]}]
        with patch.object(pms_push,'_api',return_value=raw):
            self.assertEqual(pms_push.read_current('https://pms.example.com',101,None)['7']['Velocity']['value'],0)

    def test_pms_null_period_placeholder_is_empty_not_malformed(self):
        placeholder={'id':None,'name':None,'description':None,'updatedAt':None,
                     'periodKpiValue':{'value':None,'note':None,'updatedAt':None}}
        raw={'data':[{'kpi':{'name':'Velocity'},'periods':[placeholder]}]}
        with patch.object(pms_push,'_api',return_value=raw):
            self.assertEqual(pms_push.read_current('https://pms.example.com',101,None),{})
            placeholder['periodKpiValue']['value']=0
            with self.assertRaisesRegex(ValueError,'identity'):
                pms_push.read_current('https://pms.example.com',101,None)

    def test_pms_create_uses_period_kpis_and_verifies(self):
        import os
        self.profile_data['output']['mode']='assisted-push'
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        payload=self.base/'payload.json'
        payload.write_text(json.dumps([{'projectId':101,'name':'Cycle','description':'Delivery',
            'kpis':[{'name':'Velocity','kpiId':1,'value':10,'note':'Ten estimated hours.'}]}]))
        sent=[]
        def api(base,path,**kw):
            if kw.get('method')=='POST':
                sent.append(kw['body']); return {'data':{'id':7}}
            return {'data':[]}
        after={'7':{'Velocity':{'value':10,'note':'Ten estimated hours.'}}}
        with patch.dict(os.environ,{'PMS_TOKEN':'test-only'}), patch.object(pms_push,'read_current',side_effect=[{},after]), patch.object(pms_push,'_api',side_effect=api):
            result=pms_push.main(['--profile',str(self.profile),'--payloads',str(payload),'--apply','--create-periods'])
        self.assertEqual(result,0)
        self.assertEqual(set(sent[0]),{'name','description','periodKpis'})
        self.assertEqual(sent[0]['periodKpis'][0]['value'],10)

    def test_minimal_profile_and_local_targets(self):
        path=self.base/'minimal.yaml'
        profile_tool.init(path,None)
        self.assertEqual(profile_tool.validate(path),0)
        self.profile_data['targets']={'defect_rate':{'value':20,'why':'Agreed phase target'}}
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        ws=self.ws()
        kif,_=classify.to_kif(self.snap,ws.profile,ws.project,ws.facts,ws.ledger,ws.today)
        ws.run_dir.mkdir(parents=True); path=ws.run_dir/'run.kif.json'; path.write_text(json.dumps(kif))
        result=kpi.compute(ws,path)
        measure=next(m for m in result['periods'][0]['measures'] if m['key']=='defect_rate')
        self.assertEqual(measure['threshold'],20)
        self.assertIn('Local review target',measure['threshold_source'])

    def test_registry_uses_maximum_when_minimum_is_null(self):
        targets,_=kpi_registry.project_thresholds([{'name':'Defect Rate','threshold':None,'minValue':None,'maxValue':15}])
        self.assertEqual(targets['defect_rate']['threshold'],15)

    def test_google_expands_grid_before_dimension_changes(self):
        tab=sheet_model.Tab('Example'); tab.put(1,30,'Wide')
        reqs=sheet_google.tab_requests(tab,1,0,{'properties':{'gridProperties':{'rowCount':5,'columnCount':8}}})
        resize=next(i for i,x in enumerate(reqs) if 'updateSheetProperties' in x)
        dimensions=next(i for i,x in enumerate(reqs) if 'updateDimensionProperties' in x)
        self.assertLess(resize,dimensions)

    def test_wrong_pms_period_owner_never_writes(self):
        import os
        self.profile_data['output']['mode']='assisted-push'
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        payload=self.base/'payload.json'; payload.write_text(json.dumps([{'projectId':101,'periodId':7,'name':'Cycle','kpis':[]}]))
        calls=[]
        def api(base,path,**kw):
            calls.append(kw.get('method','GET'))
            return {'projectId':999}
        with patch.dict(os.environ,{'PMS_TOKEN':'test-only'}), patch.object(pms_push,'read_current',return_value={}), patch.object(pms_push,'_api',side_effect=api):
            result=pms_push.main(['--profile',str(self.profile),'--payloads',str(payload),'--apply'])
        self.assertNotEqual(result,0)
        self.assertNotIn('PUT',calls)

    def test_incomplete_pms_payload_never_writes(self):
        self.profile_data['output']['mode']='assisted-push'
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        payload=self.base/'payload.json'; payload.write_text(json.dumps([{'projectId':101,'periodId':7,'name':'Cycle','kpis':[], 'skipped':[{'name':'Velocity'}]}]))
        with patch.object(pms_push,'_api') as api:
            result=pms_push.main(['--profile',str(self.profile),'--payloads',str(payload),'--apply'])
        self.assertEqual(result,2); api.assert_not_called()

    def test_legacy_push_log_without_timestamp(self):
        ws=self.ws(); ws.run_dir.mkdir(parents=True)
        (ws.run_dir/'push_log.json').write_text(json.dumps([{'period':'Cycle','result':'verified'}]))
        rows,pushed=kpi.push_history(ws)
        self.assertEqual(len(rows),1); self.assertIn('Cycle',pushed)

    def test_run_label_cannot_escape_workspace(self):
        self.args.date='../outside'
        with self.assertRaisesRegex(SystemExit,'single folder'):
            self.ws()

    def test_unresolved_plan_workbook_matches_engine_including_future_dates(self):
        import formulas
        with contextlib.redirect_stdout(io.StringIO()):
            code=kpi.main(['run','--profile',str(self.profile),'--project','northwind-q3',
                          '--board',str(self.profile.parent/'board.json'),'--offline','--today','2026-09-01'])
        self.assertEqual(code,0)
        directory=self.profile.parent/'northwind-q3'
        book=next(directory.glob('KPI Tracker*.xlsx'))
        result=json.loads((directory/'runs/2026-09-01/results.json').read_text())
        sol=formulas.ExcelModel().loads(str(book)).finish().calculate()
        live={str(key).split('!')[-1].strip("'"):value.value[0][0] for key,value in sol.items() if 'KPI SUMMARY' in str(key).upper()}
        sheet=load_workbook(book)['KPI Summary']
        for row in range(4,sheet.max_row+1):
            period,name=sheet.cell(row,1).value,sheet.cell(row,3).value
            if not period: continue
            expected=next(m['value'] for p in result['periods'] if p['period']==period for m in p['measures'] if m['name']==name)
            actual=live[f'G{row}']; actual=None if actual in ('',None) else round(float(actual),2)
            self.assertEqual(actual,None if expected is None else round(expected,2),f'{period}: {name}')
        self.assertFalse(any('TODAY()' in str(c.value) for sh in load_workbook(book) for row in sh for c in row))

    def test_registry_name_aliases_keep_live_formulas_and_dashboard_labels(self):
        registry = json.loads(kpi_registry.DEFAULT.read_text())
        aliases = {'Defect Rejection Rate': 'Rejection Rate', 'CR Rate': 'Change Request Rate'}
        for metric in registry['kpis']:
            metric['name'] = aliases.get(metric['name'], metric['name'])
        (self.profile.parent / 'aliases.json').write_text(json.dumps(registry))
        self.profile_data['organization']['kpi_registry'] = 'aliases.json'
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        self.test_unresolved_plan_workbook_matches_engine_including_future_dates()
        book = load_workbook(next((self.profile.parent / 'northwind-q3').glob('KPI Tracker*.xlsx')))
        names = [c.value for row in book['Dashboard'] for c in row if c.column == 1]
        for name in aliases.values():
            self.assertIn(name, names)
            row = next(row for row in book['Config'] if row[1].value == name)
            self.assertTrue(row[7].value)

    def test_rerun_without_period_facts_preserves_missing_period_question(self):
        (self.profile.parent/'northwind-q3/facts/periods.yaml').unlink(missing_ok=True)
        args=['run','--profile',str(self.profile),'--project','northwind-q3','--board',str(self.profile.parent/'board.json'),'--offline']
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(kpi.main(args),0)
            self.assertEqual(kpi.main(args),0)
        nxt=json.loads((self.profile.parent/'northwind-q3/next.json').read_text())
        self.assertTrue(any(q['id']=='periods' for q in nxt['questions']))

    def test_first_google_run_cannot_claim_published_when_only_local_file_exists(self):
        self.profile_data['output']['workbook'] = 'google-sheets'
        self.profile_data['output']['workbook_file'] = 'fictional-sheet-id'
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            self.assertEqual(kpi.main(['run', '--profile', str(self.profile), '--project', 'northwind-q3',
                                      '--board', str(self.profile.parent/'board.json'), '--offline']), 0)
        nxt = json.loads((self.profile.parent/'northwind-q3/next.json').read_text())
        self.assertTrue(nxt['sheet_pending'])
        self.assertIn('configured Google Sheet has not been updated', output.getvalue())
        self.assertNotIn('Nothing. The sheet is up to date.', output.getvalue())

    def test_pms_readback_mismatch_is_a_failure(self):
        import os
        self.profile_data['output']['mode']='assisted-push'
        self.profile.write_text(yaml.safe_dump(self.profile_data))
        payload=self.base/'payload.json'; payload.write_text(json.dumps([{'projectId':101,'periodId':7,'name':'Cycle',
            'kpis':[{'name':'Velocity','kpiId':1,'value':10,'note':'10 points'}]}]))
        log=self.base/'push.json'
        responses=[{'data':[{'id':7,'projectId':101,'name':'Cycle'}]},
                   {'data':{'id':7,'updatedAt':'2026-01-01T00:00:00Z'}}, {}]
        with patch.dict(os.environ,{'PMS_TOKEN':'test-only'}), patch.object(pms_push,'read_current',side_effect=[{},None]), patch.object(pms_push,'_api',side_effect=responses):
            result=pms_push.main(['--profile',str(self.profile),'--payloads',str(payload),'--apply','--log',str(log)])
        self.assertNotEqual(result,0)
        self.assertIn('mismatch',json.loads(log.read_text())['results'][0]['result'])

    def test_changed_inputs_and_unresolved_review_refuse_submission(self):
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(kpi.main(['run','--profile',str(self.profile),'--project','northwind-q3',
                '--board',str(self.profile.parent/'board.json'),'--offline']),0)
        args=['push','--profile',str(self.profile),'--project','northwind-q3','--apply']
        with patch.object(kpi,'_module') as writer:
            with self.assertRaisesRegex(SystemExit,'still needs review'): kpi.main(args)
            self.profile_data['owner']['name']='Changed owner'
            self.profile.write_text(yaml.safe_dump(self.profile_data))
            with self.assertRaisesRegex(SystemExit,'Inputs changed'): kpi.main(args)
            writer.assert_not_called()


if __name__ == '__main__':
    unittest.main()
