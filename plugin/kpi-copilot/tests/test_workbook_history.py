import copy
import datetime as dt
import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import workbook_history as H
import sheet_model as M
import sheet_google as G
import sheet_readback as R
import sheet_xlsx as X
from ledger import Ledger


def sample(start, name, value=80):
    p = {'name': name, 'start': start, 'end': (dt.date.fromisoformat(start)+dt.timedelta(days=6)).isoformat()}
    kif = {'project': {'name': 'Example', 'velocity_unit': 'Story Points'}, 'periods': [p],
           'tasks': [{'_item':'same-ticket', 'period':name, 'key':'EX-1', 'type':'Task', 'title':'Example work', 'understood':'Yes'}], 'defects':[]}
    res = {'periods':[{'period':name,'measures':[{'name':n, 'value':value, 'threshold':90, 'direction':'higher-is-better'} for n in M.KPI_ORDER]}]}
    return kif, res, {'profile':{}, 'as_of':p['end'], 'reasons':{name:{'Defect Rate':'A documented cause.'}}}


class HistoryTests(unittest.TestCase):
    def test_new_period_keeps_old_and_refresh_upserts_by_dates(self):
        history = {}
        a = sample('2026-01-01','January 1 - January 7')
        b = sample('2026-01-08','January 8 - January 14')
        H.upsert(history,*a); H.upsert(history,*b)
        refresh = sample('2026-01-01','January 1 - January 7, 2026',91)
        H.upsert(history,*refresh)
        kif,res,ctx = H.combine(history,*refresh)
        self.assertEqual(len(kif['periods']),2)
        self.assertEqual(res['periods'][0]['measures'][0]['value'],91)
        self.assertEqual(res['periods'][1]['measures'][0]['value'],80)
        self.assertEqual(ctx['reasons'][b[0]['periods'][0]['name']]['Defect Rate'],'A documented cause.')
        self.assertEqual(len({r.get('_row') or r['_item'] for r in kif['tasks']}),2)
        self.assertEqual(len(refresh[0]['periods']),1)  # publishing never expands PMS payload scope

    def test_250_periods_have_no_dashboard_or_lookup_limit(self):
        h = {}
        for i in range(250):
            s = sample((dt.date(2020,1,1)+dt.timedelta(days=7*i)).isoformat(),f'Week {i+1}')
            H.upsert(h,*s)
        kif,res,ctx = H.combine(h,*s)
        tabs={t.name:t for t in M.build(kif,res,ctx)}
        self.assertEqual(tabs['Dashboard'].cells[4,2]['v'],'Week 250')
        self.assertEqual(tabs['Period Overview'].cells[5,1]['v'],'Week 250')
        self.assertEqual(tabs['Period Overview'].cells[254,1]['v'],'Week 1')
        self.assertIn('$2254',tabs['Dashboard'].cells[8,8]['f'])
        self.assertTrue(any('$294' in (v.get('source') or '') for v in tabs['Task Register'].validations))
        self.assertLess(tabs['Dashboard'].size[0],30)
        self.assertEqual(tabs['Period Overview'].filter_range,(4,1,254,11))

    def test_archive_edit_does_not_change_active_ticket_ledger(self):
        h={}; a=sample('2026-01-01','First'); b=sample('2026-01-08','Second')
        H.upsert(h,*a); H.upsert(h,*b)
        kif,res,ctx=H.combine(h,*b)
        tabs=M.build(kif,res,ctx)
        state=R.snapshot(tabs,{'kind':'xlsx'})
        with tempfile.TemporaryDirectory() as tmp:
            book=Path(tmp)/'book.xlsx'; X.write(tabs,book)
            from openpyxl import load_workbook
            wb=load_workbook(book); wb['Task Register']['P4']='No'; wb['Dashboard']['B4']='First'; wb.save(book)
            facts={}; ledger=Ledger(Path(tmp)/'ledger.json')
            R.fold(state,X.read_grid(book),ledger,facts,{}, {})
            self.assertEqual(facts['view']['period'],'First')
            archive_id=kif['tasks'][0]['_row']
            self.assertEqual(facts['history_edits'][archive_id]['understood'],'No')
            self.assertFalse(ledger.data['items'])
            ctx.update(facts)
            out,_,_=H.combine(h,b[0],b[1],ctx)
            self.assertEqual(out['tasks'][0]['understood'],'No')
            self.assertEqual(out['tasks'][1]['understood'],'Yes')
            refreshed = copy.deepcopy(a[0])
            H.apply_edits(refreshed, facts)
            self.assertEqual(refreshed['tasks'][0]['understood'],'No')

    def test_names_cannot_silently_combine_different_years(self):
        h={};a=sample('2025-01-01','January 1 - January 7');b=sample('2026-01-01','January 1 - January 7')
        H.upsert(h,*a);H.upsert(h,*b)
        with self.assertRaisesRegex(ValueError,'across years'):H.combine(h,*b)

    def test_later_run_cannot_recalculate_archived_values_under_new_policy(self):
        h={};a=sample('2026-01-01','First',80);b=sample('2026-01-08','Second',95)
        a[1]['periods'][0]['measures'][0].update(numerator=8,denominator=10)
        H.upsert(h,*a);H.upsert(h,*b)
        tabs={t.name:t for t in M.build(*H.combine(h,*b))}
        summary=tabs['KPI Summary']
        self.assertEqual([summary.cells[4,c]['v'] for c in (5,6,7,8)], [8,10,80,90])
        self.assertTrue(all(not summary.cells[4,c].get('f') for c in (5,6,7,8)))
        self.assertTrue(summary.cells[13,7].get('f'))
        self.assertIn('<>history:*',summary.cells[21,6]['f'])

    def test_connector_snapshot_rejects_other_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'cells.json';path.write_text(json.dumps({'spreadsheetId':'wrong','sheets':[]}))
            with self.assertRaisesRegex(ValueError,'different spreadsheet'):G.connector_grid(path,'expected')

    def test_historical_corrections_refresh_date_results_but_cannot_bypass_scope(self):
        kif, _, _ = sample('2026-01-01', 'First')
        row = kif['tasks'][0]
        row.update(delivered='2026-01-07', commit_date='2026-01-06', met_commitment='No',
                   exclude_reason='assignee is outside the configured project team', type='Excluded')
        rid = H.row_id(H.period_key(kif['periods'][0]), 'tasks', row, 0)
        H.apply_edits(kif, {'history_edits':{rid:{'type':'Task','delivered':'2026-01-05'}}}, {}, '2026-01-08')
        self.assertEqual(row['met_commitment'], 'Yes')
        self.assertEqual(row['type'], 'Excluded')

    def test_connector_uses_precise_values_and_preserves_dates(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/'cells.json'
            doc = {'spreadsheetId':'example','sheets':[{'properties':{'title':'Review'},'data':[{'rowData':[{'values':[
                {'formattedValue':'1,234.57','effectiveValue':{'numberValue':1234.567}},
                {'formattedValue':'01/02/2026','effectiveValue':{'numberValue':46024},
                 'effectiveFormat':{'numberFormat':{'type':'DATE'}}},
                {'formattedValue':'0.00%','effectiveValue':{'numberValue':0}}
            ]}]}]}]}
            path.write_text(json.dumps(doc))
            grid = G.connector_grid(path,'example')
            self.assertEqual(grid('Review',1,1),1234.567)
            self.assertEqual(grid('Review',1,2),'01/02/2026')
            self.assertEqual(grid('Review',1,3),0)
            doc['sheets'][0]['data'][0]['rowData'][0]['values'][0] = {
                'effectiveValue':{'errorValue':{'type':'REF'}}}
            path.write_text(json.dumps(doc))
            with self.assertRaisesRegex(ValueError,'Formula error'):
                G.connector_grid(path,'example')

    def test_filters_merges_and_formats_exist_in_both_writers(self):
        tabs=M.build(*sample('2026-01-01','First'))
        dash=tabs[0]
        requests=G.tab_requests(dash,1,0,None)
        self.assertTrue(any('mergeCells' in r for r in requests))
        self.assertEqual(G._format({},'0.00"%"')['numberFormat']['type'],'NUMBER')
        with tempfile.TemporaryDirectory() as tmp:
            book=Path(tmp)/'book.xlsx';X.write(tabs,book)
            from openpyxl import load_workbook
            wb=load_workbook(book)
            self.assertEqual(wb.active.title,'Dashboard')
            self.assertTrue(wb['Task Register'].auto_filter.ref)
            self.assertIn('B4:E4',str(wb['Dashboard'].merged_cells))


if __name__=='__main__':unittest.main()
